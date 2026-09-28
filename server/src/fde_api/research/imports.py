from __future__ import annotations

import csv
from dataclasses import dataclass, replace
from defusedxml import ElementTree as SafeElementTree
from defusedxml.common import DefusedXmlException
from hashlib import sha256
from io import BytesIO, StringIO
import json
from pathlib import Path
from pathlib import PurePosixPath
import posixpath
import re
from secrets import token_hex, token_urlsafe
import threading
from typing import Any, BinaryIO
from zipfile import BadZipFile, ZipFile

from flask import current_app
from itsdangerous import BadSignature, SignatureExpired, URLSafeTimedSerializer
import openpyxl
from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet._reader import WorkSheetParser
from redis import Redis
from redis.exceptions import RedisError
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError, SQLAlchemyError

from fde_api.auth.models import User
from fde_api.extensions import db
from fde_api.research.models import (
    ProjectResearchImportBatch,
    ProjectResearchSubject,
    SUBJECT_TYPES,
)


IMPORT_COLUMNS = {"名称": "name", "说明": "description", "父级": "parent_name"}
_IMPORT_SUBJECT_TYPES = frozenset(SUBJECT_TYPES) - {"project"}
MAX_IMPORT_ROWS = 5_000
MAX_IMPORT_BYTES = 10 * 1024 * 1024
PREVIEW_TTL_SECONDS = 30 * 60
_COMMIT_LOCK_TTL_SECONDS = 60
_TOKEN_SALT = "fde-research-subject-import-v1"
_PREVIEW_KEY_PREFIX = "fde:research-import:preview:"
_LOCK_KEY_PREFIX = "fde:research-import:lock:"
_FENCE_KEY_PREFIX = "fde:research-import:fence:"
_SUPPORTED_EXTENSIONS = frozenset({".csv", ".xlsx"})
_CSV_FIELD_SIZE_LOCK = threading.Lock()
_MAX_XLSX_MEMBERS = 512
_MAX_XLSX_MEMBER_BYTES = 16 * 1024 * 1024
_MAX_XLSX_TOTAL_BYTES = 32 * 1024 * 1024
_MAX_XLSX_COMPRESSION_RATIO = 100
_MAX_XLSX_USED_COLUMNS = 16_384
_MAX_XLSX_CELL_RECORDS = 25_100
_MAX_XLSX_SOURCE_ROW = 100_000
_CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
_RELATIONSHIPS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_VML_NS = "urn:schemas-microsoft-com:vml"
_OFFICE_NS = "urn:schemas-microsoft-com:office:office"
_EXCEL_NS = "urn:schemas-microsoft-com:office:excel"
_COMMENT_SHAPE_STYLE = re.compile(
    r"position:absolute; margin-left:59\.25pt;margin-top:1\.5pt;"
    r"width:(?:0|[1-9][0-9]*)(?:\.[0-9]+)?px;"
    r"height:(?:0|[1-9][0-9]*)(?:\.[0-9]+)?px;"
    r"z-index:1;visibility:hidden"
)
_ALLOWED_CONTENT_TYPES = frozenset(
    {
        "application/vnd.openxmlformats-package.relationships+xml",
        "application/xml",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml",
        "application/vnd.openxmlformats-officedocument.theme+xml",
        "application/vnd.openxmlformats-package.core-properties+xml",
        "application/vnd.openxmlformats-officedocument.extended-properties+xml",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.table+xml",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.comments+xml",
        "application/vnd.openxmlformats-officedocument.vmlDrawing",
    }
)
_PART_TYPES = {
    "docProps/core.xml": "application/vnd.openxmlformats-package.core-properties+xml",
    "docProps/app.xml": "application/vnd.openxmlformats-officedocument.extended-properties+xml",
    "xl/workbook.xml": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
    "xl/styles.xml": "application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml",
    "xl/sharedStrings.xml": "application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml",
    "xl/calcChain.xml": "application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml",
}
_ROOT_RELATIONSHIP_TYPES = frozenset(
    {
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument",
        "http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties",
    }
)
_WORKBOOK_RELATIONSHIP_TYPES = frozenset(
    {
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain",
    }
)
_WORKSHEET_RELATIONSHIP_TYPES = frozenset(
    {
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/table",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments",
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/vmlDrawing",
    }
)


@dataclass(frozen=True, slots=True)
class ImportRowError:
    row: int
    column: str
    code: str
    message: str

    def to_dict(self) -> dict[str, Any]:
        return {
            "row": self.row,
            "column": self.column,
            "code": self.code,
            "message": self.message,
        }


@dataclass(frozen=True, slots=True)
class ImportRow:
    row: int
    name: object
    description: object
    parent_name: object

    def to_dict(self) -> dict[str, Any]:
        return {
            "row": self.row,
            "name": self.name,
            "description": self.description,
            "parent_name": self.parent_name,
        }


@dataclass(frozen=True, slots=True)
class SparseSourceRow:
    row: int
    values: dict[int, object]


@dataclass(frozen=True, slots=True)
class ImportPreview:
    subject_type: str
    rows: tuple[ImportRow, ...]
    errors: tuple[ImportRowError, ...]
    preview_token: str | None = None
    header_errors: tuple[ImportRowError, ...] = ()

    def to_dict(self) -> dict[str, Any]:
        return {
            "preview_token": self.preview_token,
            "subject_type": self.subject_type,
            "row_count": len(self.rows),
            "rows": [row.to_dict() for row in self.rows],
            "errors": [error.to_dict() for error in self.errors],
            "expires_in_seconds": PREVIEW_TTL_SECONDS,
        }


@dataclass(frozen=True, slots=True)
class ImportResult:
    subjects: tuple[dict[str, Any], ...]
    project_version: int
    idempotent_replay: bool = False

    def to_dict(self) -> dict[str, Any]:
        return {
            "subjects": list(self.subjects),
            "imported_count": len(self.subjects),
            "project_version": self.project_version,
            "idempotent_replay": self.idempotent_replay,
        }


@dataclass(frozen=True, slots=True)
class CommitLease:
    state_key: str
    lock_key: str
    owner: str
    fence: int
    ttl: int


class ImportServiceError(Exception):
    def __init__(
        self,
        code: str,
        message: str,
        status: int,
        details: dict[str, Any] | None = None,
    ):
        super().__init__(message)
        self.code = code
        self.message = message
        self.status = status
        self.details = details


def preview_subject_import(
    stream: BinaryIO, filename: str, subject_type: str
) -> ImportPreview:
    """Parse one bounded homogeneous CSV/XLSX import without persisting file bytes."""
    if subject_type not in _IMPORT_SUBJECT_TYPES:
        raise _invalid_request()
    extension = Path(filename or "").suffix.lower()
    if extension not in _SUPPORTED_EXTENSIONS:
        raise ImportServiceError(
            "unsupported_import_format",
            "Only CSV and XLSX subject imports are supported.",
            415,
        )
    contents = _read_bounded(stream)
    if extension == ".csv":
        table = _csv_table(contents)
    else:
        _validate_xlsx_package(contents)
        table = _xlsx_table(contents)
    return _build_preview(table, subject_type)


def store_subject_import_preview(
    *, actor: User, project_id: str, preview: ImportPreview
) -> ImportPreview:
    """Authorize and bind a parsed preview to one project for thirty minutes."""
    authorize_subject_import(actor, project_id)
    preview_id = token_hex(16)
    token = _serializer().dumps({"preview_id": preview_id, "project_id": project_id})
    state = {
        "status": "preview",
        "project_id": project_id,
        "preview_id": preview_id,
        "subject_type": preview.subject_type,
        "rows": [row.to_dict() for row in preview.rows],
        "errors": [error.to_dict() for error in preview.errors],
        "header_errors": [error.to_dict() for error in preview.header_errors],
    }
    try:
        stored = _redis().set(
            _preview_key(preview_id),
            json.dumps(state, ensure_ascii=False, separators=(",", ":")),
            ex=PREVIEW_TTL_SECONDS,
            nx=True,
        )
    except RedisError:
        raise _storage_unavailable() from None
    if not stored:
        raise _storage_unavailable()
    return replace(preview, preview_token=token)


def commit_subject_import(
    project_id: str,
    preview_token: str,
    expected_version: int,
    actor: User,
) -> ImportResult:
    """Commit a project-bound preview once, replaying its result within the TTL."""
    from fde_api.research.subject_service import ResearchSubjectServiceError

    if type(expected_version) is not int or expected_version < 1:
        raise _invalid_request()
    claims = _token_claims(preview_token)
    if claims["project_id"] != project_id:
        raise ImportServiceError("project_not_found", "Project was not found.", 404)
    preview_id = claims["preview_id"]
    receipt = _authorized_import_receipt(actor, project_id, preview_id)
    if receipt is not None:
        return receipt
    state_key = _preview_key(preview_id)
    state = _load_state(state_key)
    if state.get("project_id") != project_id or state.get("preview_id") != preview_id:
        raise _preview_not_found()

    if state.get("status") == "committed":
        return _result_from_database_receipt(actor, project_id, preview_id)

    try:
        lease = _acquire_commit_lease(
            state_key, preview_id, expected_version
        )
    except ImportServiceError as error:
        if error.code != "research_import_commit_in_progress":
            raise
        latest = _load_state(state_key)
        if latest.get("status") == "committed":
            return _result_from_database_receipt(actor, project_id, preview_id)
        raise

    try:
        try:
            result = _commit_state(
                actor=actor,
                project_id=project_id,
                expected_version=expected_version,
                state=_load_state(state_key),
            )
        except (ImportServiceError, ResearchSubjectServiceError):
            _reset_commit_state(lease)
            raise
        except (IntegrityError, SQLAlchemyError):
            _reset_commit_state(lease)
            raise ImportServiceError(
                "research_import_failed",
                "Unable to import research subjects at this time.",
                503,
            ) from None
        if not _mark_commit_state(lease, result.to_dict()):
            # The database receipt remains authoritative if this owner lost its fence.
            raise _storage_unavailable()
        return result
    finally:
        _release_commit_lease(lease)


def _commit_state(
    *,
    actor: User,
    project_id: str,
    expected_version: int,
    state: dict[str, Any],
) -> ImportResult:
    from fde_api.research.form_service import build_subject_form_context
    from fde_api.research.subject_service import (
        ResearchSubjectServiceError,
        allocate_tracking_codes,
        create_subject_in_transaction,
        load_subject_project,
        serialize_subject,
    )

    rows = tuple(_row_from_dict(item) for item in state.get("rows", []))
    subject_type = state.get("subject_type")
    preview_id = state["preview_id"]
    session = db.session()
    try:
        with session.begin():
            project = load_subject_project(
                session, actor, project_id, for_update=True
            )
            recovered = _load_import_receipt(session, project.id, preview_id)
            if recovered is not None:
                return recovered
            if project.version != expected_version:
                raise ImportServiceError(
                    "stale_version",
                    "The project has been updated. Refresh and try again.",
                    409,
                )
            header_errors = [
                _error_from_dict(error) for error in state.get("header_errors", [])
            ]
            if header_errors:
                raise _invalid_import(header_errors)
            parents = _load_parent_candidates(
                session, project.id, subject_type, rows
            )
            tracking_codes = allocate_tracking_codes(
                session, project.id, subject_type, len(rows)
            )
            form_context = build_subject_form_context(session, project)
            subjects: list[ProjectResearchSubject] = []
            for position, (row, tracking_code) in enumerate(
                zip(rows, tracking_codes, strict=True)
            ):
                errors = _validate_import_row(row, subject_type)
                if errors:
                    raise _invalid_import(errors)
                parent_id = _resolve_parent(row, parents)
                subject = create_subject_in_transaction(
                    session=session,
                    actor=actor,
                    project=project,
                    input={
                        "subject_type": subject_type,
                        "subject_key": _subject_key(preview_id, subject_type, row.row),
                        "name": row.name,
                        "description": row.description,
                        "sort_order": position,
                        "parent_subject_id": parent_id,
                    },
                    tracking_code=tracking_code,
                    form_context=form_context,
                )
                subjects.append(subject)
            if not subjects:
                raise _invalid_import(
                    [
                        ImportRowError(
                            1,
                            "名称",
                            "empty_import",
                            "The import contains no subject rows.",
                        )
                    ]
                )
            project.version += 1
            result = ImportResult(
                tuple(serialize_subject(subject) for subject in subjects),
                project.version,
            )
            session.add(
                ProjectResearchImportBatch(
                    id=preview_id,
                    project_id=project.id,
                    subject_type=subject_type,
                    expected_project_version=expected_version,
                    committed_project_version=project.version,
                    result_json=result.to_dict(),
                )
            )
            session.flush()
            return result
    except ResearchSubjectServiceError:
        raise
    finally:
        session.close()


def _load_import_receipt(session, project_id: str, preview_id: str) -> ImportResult | None:
    receipt = session.scalar(
        select(ProjectResearchImportBatch).where(
            ProjectResearchImportBatch.id == preview_id,
            ProjectResearchImportBatch.project_id == project_id,
        )
    )
    if receipt is None:
        return None
    raw = receipt.result_json
    subjects = raw.get("subjects") if isinstance(raw, dict) else None
    if not isinstance(subjects, list):
        raise ImportServiceError(
            "research_import_conflict",
            "The subject import receipt is invalid.",
            409,
        )
    return ImportResult(
        tuple(subjects),
        receipt.committed_project_version,
        idempotent_replay=True,
    )


def _authorized_import_receipt(
    actor: User, project_id: str, preview_id: str
) -> ImportResult | None:
    from fde_api.research.subject_service import load_subject_project

    session = db.session()
    try:
        project = load_subject_project(session, actor, project_id)
        return _load_import_receipt(session, project.id, preview_id)
    finally:
        session.close()


def _result_from_database_receipt(
    actor: User, project_id: str, preview_id: str
) -> ImportResult:
    result = _authorized_import_receipt(actor, project_id, preview_id)
    if result is None:
        raise _preview_not_found()
    return result


def authorize_subject_import(actor: User, project_id: str) -> None:
    """Reject callers without the shared project subject-write capability."""
    from fde_api.research.subject_service import (
        ResearchSubjectServiceError,
        load_subject_project,
    )

    session = db.session()
    try:
        load_subject_project(session, actor, project_id)
    except ResearchSubjectServiceError as error:
        raise ImportServiceError(error.code, error.message, error.status) from None
    finally:
        session.close()


def _read_bounded(stream: BinaryIO) -> bytes:
    remaining = MAX_IMPORT_BYTES + 1
    chunks: list[bytes] = []
    while remaining:
        chunk = stream.read(min(64 * 1024, remaining))
        if not chunk:
            break
        if not isinstance(chunk, bytes):
            raise _invalid_file()
        chunks.append(chunk)
        remaining -= len(chunk)
    contents = b"".join(chunks)
    if len(contents) > MAX_IMPORT_BYTES:
        raise ImportServiceError(
            "research_import_too_large",
            "The subject import exceeds the 10 MiB limit.",
            413,
        )
    if not contents:
        raise _invalid_file()
    return contents


def _csv_table(contents: bytes) -> list[tuple[int, dict[int, object]]]:
    text = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            text = contents.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ImportServiceError(
            "unsupported_csv_encoding",
            "CSV files must use UTF-8 or GB18030 encoding.",
            422,
        )
    with _CSV_FIELD_SIZE_LOCK:
        previous_limit = csv.field_size_limit()
        try:
            csv.field_size_limit(MAX_IMPORT_BYTES)
            return _bounded_rows(csv.reader(StringIO(text, newline="")))
        except csv.Error:
            raise _invalid_file() from None
        finally:
            csv.field_size_limit(previous_limit)


def _xlsx_table(contents: bytes) -> list[tuple[int, dict[int, object]]]:
    workbook = None
    try:
        workbook = load_workbook(
            BytesIO(contents), read_only=True, data_only=True, keep_vba=False
        )
        if not workbook.worksheets:
            raise _invalid_file()
        worksheet = workbook.worksheets[0]
        # OOXML dimensions are untrusted hints. Recalculate from actual cell records.
        worksheet.reset_dimensions()
        _xlsx_used_bounds(contents, worksheet.title)
        return _bounded_rows(_sparse_worksheet_rows(workbook, worksheet))
    except ImportServiceError:
        raise
    except (
        BadZipFile,
        InvalidFileException,
        DefusedXmlException,
        SafeElementTree.ParseError,
        KeyError,
        OSError,
        ValueError,
    ):
        raise _invalid_file() from None
    finally:
        if workbook is not None:
            workbook.close()


def _validate_xlsx_package(contents: bytes) -> None:
    """Validate ZIP metadata, package declarations and XML before openpyxl sees bytes."""
    try:
        with ZipFile(BytesIO(contents)) as archive:
            infos = archive.infolist()
            _validate_xlsx_container(infos)
            names = {info.filename for info in infos if not info.is_dir()}
            if "[Content_Types].xml" not in names:
                raise _invalid_file()
            content_root = SafeElementTree.fromstring(
                archive.read("[Content_Types].xml")
            )
            declared_parts: dict[str, str] = {}
            defaults: dict[str, str] = {}
            for child in content_root:
                if child.tag == f"{{{_CONTENT_TYPES_NS}}}Default":
                    extension = child.attrib.get("Extension", "").lower()
                    content_type = child.attrib.get("ContentType", "")
                    if extension in defaults:
                        raise _invalid_file()
                    defaults[extension] = content_type
                elif child.tag == f"{{{_CONTENT_TYPES_NS}}}Override":
                    part_name = child.attrib.get("PartName", "").lstrip("/")
                    content_type = child.attrib.get("ContentType", "")
                    if part_name in declared_parts:
                        raise _invalid_file()
                    if (
                        content_type not in _ALLOWED_CONTENT_TYPES
                        or content_type != _expected_part_type(part_name)
                    ):
                        raise _unsupported_active_workbook()
                    declared_parts[part_name] = content_type
            for extension, content_type in defaults.items():
                if content_type not in _ALLOWED_CONTENT_TYPES:
                    raise _unsupported_active_workbook()
                if extension not in {"xml", "rels", "vml"}:
                    raise _unsupported_active_workbook()
            for name in names - {"[Content_Types].xml"}:
                extension = name.rsplit(".", 1)[-1].lower() if "." in name else ""
                expected_type = _expected_part_type(name)
                actual_type = declared_parts.get(name, defaults.get(extension))
                if expected_type is None or actual_type != expected_type:
                    raise _unsupported_active_workbook()
                if extension in {"xml", "rels", "vml"}:
                    root = SafeElementTree.fromstring(archive.read(name))
                    if extension == "rels":
                        _validate_relationships(root, name, names)
                    elif extension == "vml":
                        _validate_comment_vml(root)
    except ImportServiceError:
        raise
    except (
        BadZipFile,
        DefusedXmlException,
        SafeElementTree.ParseError,
        KeyError,
        OSError,
    ):
        raise _invalid_file() from None


def _validate_xlsx_container(infos) -> None:
    if not infos or len(infos) > _MAX_XLSX_MEMBERS:
        raise _workbook_too_complex()
    total = 0
    seen: set[str] = set()
    for info in infos:
        name = info.filename
        path = PurePosixPath(name)
        if (
            not name
            or "\x00" in name
            or "\\" in name
            or path.is_absolute()
            or ".." in path.parts
            or name in seen
            or info.flag_bits & 0x1
        ):
            raise _invalid_file()
        seen.add(name)
        if info.is_dir():
            continue
        if info.file_size > _MAX_XLSX_MEMBER_BYTES:
            raise _workbook_too_complex()
        total += info.file_size
        if total > _MAX_XLSX_TOTAL_BYTES:
            raise _workbook_too_complex()
        if info.file_size and (
            info.compress_size == 0
            or info.file_size / info.compress_size > _MAX_XLSX_COMPRESSION_RATIO
        ):
            raise _workbook_too_complex()


def _validate_relationships(root, relationship_part: str, names: set[str]) -> None:
    if root.tag != f"{{{_RELATIONSHIPS_NS}}}Relationships":
        raise _invalid_file()
    allowed_types = _relationship_types_for_part(relationship_part)
    for relationship in root:
        if relationship.tag != f"{{{_RELATIONSHIPS_NS}}}Relationship":
            raise _invalid_file()
        relationship_type = relationship.attrib.get("Type", "")
        target = relationship.attrib.get("Target", "")
        if (
            relationship.attrib.get("TargetMode", "").lower() == "external"
            or relationship_type not in allowed_types
            or not target
            or "\x00" in target
            or "\\" in target
            or "://" in target
            or ".." in PurePosixPath(target).parts
        ):
            raise _unsupported_active_workbook()
        resolved_target = _resolve_relationship_target(relationship_part, target)
        if (
            resolved_target not in names
            or not _relationship_target_matches(
                relationship_part, relationship_type, resolved_target
            )
        ):
            raise _unsupported_active_workbook()


def _relationship_types_for_part(relationship_part: str) -> frozenset[str]:
    if relationship_part == "_rels/.rels":
        return _ROOT_RELATIONSHIP_TYPES
    if relationship_part == "xl/_rels/workbook.xml.rels":
        return _WORKBOOK_RELATIONSHIP_TYPES
    if re.fullmatch(
        r"xl/worksheets/_rels/sheet[1-9][0-9]*\.xml\.rels",
        relationship_part,
    ):
        return _WORKSHEET_RELATIONSHIP_TYPES
    raise _unsupported_active_workbook()


def _resolve_relationship_target(relationship_part: str, target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    if relationship_part == "_rels/.rels":
        base = ""
    else:
        source_directory, rels_directory, filename = relationship_part.rsplit("/", 2)
        if rels_directory != "_rels" or not filename.endswith(".rels"):
            raise _unsupported_active_workbook()
        base = source_directory
    return posixpath.normpath(posixpath.join(base, target))


def _relationship_target_matches(
    relationship_part: str, relationship_type: str, target: str
) -> bool:
    suffix = relationship_type.rsplit("/", 1)[-1]
    if relationship_part == "_rels/.rels":
        return {
            "officeDocument": "xl/workbook.xml",
            "core-properties": "docProps/core.xml",
            "extended-properties": "docProps/app.xml",
        }.get(suffix) == target
    if relationship_part == "xl/_rels/workbook.xml.rels":
        patterns = {
            "worksheet": r"xl/worksheets/sheet[1-9][0-9]*\.xml",
            "styles": r"xl/styles\.xml",
            "sharedStrings": r"xl/sharedStrings\.xml",
            "theme": r"xl/theme/theme[1-9][0-9]*\.xml",
            "calcChain": r"xl/calcChain\.xml",
        }
    else:
        patterns = {
            "table": r"xl/tables/table[1-9][0-9]*\.xml",
            "comments": r"xl/comments/comment[1-9][0-9]*\.xml",
            "vmlDrawing": r"xl/drawings/commentsDrawing[1-9][0-9]*\.vml",
        }
    pattern = patterns.get(suffix)
    return pattern is not None and re.fullmatch(pattern, target) is not None


def _validate_comment_vml(root) -> None:
    shape_layout_tag = f"{{{_OFFICE_NS}}}shapelayout"
    id_map_tag = f"{{{_OFFICE_NS}}}idmap"
    shape_type_tag = f"{{{_VML_NS}}}shapetype"
    shape_tag = f"{{{_VML_NS}}}shape"
    stroke_tag = f"{{{_VML_NS}}}stroke"
    path_tag = f"{{{_VML_NS}}}path"
    fill_tag = f"{{{_VML_NS}}}fill"
    shadow_tag = f"{{{_VML_NS}}}shadow"
    textbox_tag = f"{{{_VML_NS}}}textbox"
    client_data_tag = f"{{{_EXCEL_NS}}}ClientData"
    move_tag = f"{{{_EXCEL_NS}}}MoveWithCells"
    size_tag = f"{{{_EXCEL_NS}}}SizeWithCells"
    auto_fill_tag = f"{{{_EXCEL_NS}}}AutoFill"
    row_tag = f"{{{_EXCEL_NS}}}Row"
    column_tag = f"{{{_EXCEL_NS}}}Column"
    vml_ext = f"{{{_VML_NS}}}ext"
    office_spt = f"{{{_OFFICE_NS}}}spt"
    office_connect_type = f"{{{_OFFICE_NS}}}connecttype"
    office_inset_mode = f"{{{_OFFICE_NS}}}insetmode"

    root_children = list(root)
    if (
        root.tag != "xml"
        or root.attrib
        or not _only_vml_whitespace(root.text)
        or len(root_children) < 3
        or [child.tag for child in root_children[:2]]
        != [shape_layout_tag, shape_type_tag]
        or any(child.tag != shape_tag for child in root_children[2:])
    ):
        raise _unsupported_active_workbook()

    shape_layout, shape_type, *shapes = root_children
    layout_children = _require_comment_vml_element(
        shape_layout,
        shape_layout_tag,
        {vml_ext: "edit"},
        [id_map_tag],
    )
    _require_comment_vml_element(
        layout_children[0],
        id_map_tag,
        {vml_ext: "edit", "data": "1"},
        [],
    )

    shape_type_children = _require_comment_vml_element(
        shape_type,
        shape_type_tag,
        {
            "id": "_x0000_t202",
            "coordsize": "21600,21600",
            office_spt: "202",
            "path": "m,l,21600r21600,l21600,xe",
        },
        [stroke_tag, path_tag],
    )
    _require_comment_vml_element(
        shape_type_children[0], stroke_tag, {"joinstyle": "miter"}, []
    )
    _require_comment_vml_element(
        shape_type_children[1],
        path_tag,
        {"gradientshapeok": "t", office_connect_type: "rect"},
        [],
    )

    shape_ids: set[str] = set()
    for shape in shapes:
        expected_attribute_names = {
            "type",
            "style",
            "fillcolor",
            office_inset_mode,
            "id",
        }
        shape_id = shape.attrib.get("id", "")
        if (
            shape.tag != shape_tag
            or set(shape.attrib) != expected_attribute_names
            or shape.attrib.get("type") != "#_x0000_t202"
            or _COMMENT_SHAPE_STYLE.fullmatch(shape.attrib.get("style", "")) is None
            or shape.attrib.get("fillcolor") != "#ffffe1"
            or shape.attrib.get(office_inset_mode) != "auto"
            or re.fullmatch(r"_x0000_s[0-9]{4,}", shape_id) is None
            or shape_id in shape_ids
            or not _only_vml_whitespace(shape.text)
            or not _only_vml_whitespace(shape.tail)
        ):
            raise _unsupported_active_workbook()
        shape_ids.add(shape_id)

        shape_children = list(shape)
        expected_shape_children = [
            fill_tag,
            shadow_tag,
            path_tag,
            textbox_tag,
            client_data_tag,
        ]
        if [child.tag for child in shape_children] != expected_shape_children:
            raise _unsupported_active_workbook()
        fill, shadow, path, textbox, client_data = shape_children
        _require_comment_vml_element(fill, fill_tag, {"color2": "#ffffe1"}, [])
        _require_comment_vml_element(
            shadow, shadow_tag, {"color": "black", "obscured": "t"}, []
        )
        _require_comment_vml_element(
            path, path_tag, {office_connect_type: "none"}, []
        )
        textbox_children = _require_comment_vml_element(
            textbox,
            textbox_tag,
            {"style": "mso-direction-alt:auto"},
            ["div"],
        )
        _require_comment_vml_element(
            textbox_children[0], "div", {"style": "text-align:left"}, []
        )

        client_children = _require_comment_vml_element(
            client_data,
            client_data_tag,
            {"ObjectType": "Note"},
            [move_tag, size_tag, auto_fill_tag, row_tag, column_tag],
        )
        _require_comment_vml_element(client_children[0], move_tag, {}, [])
        _require_comment_vml_element(client_children[1], size_tag, {}, [])
        _require_comment_vml_element(
            client_children[2], auto_fill_tag, {}, [], text="False"
        )
        _require_comment_vml_coordinate(client_children[3], row_tag, 1_048_575)
        _require_comment_vml_coordinate(client_children[4], column_tag, 16_383)


def _require_comment_vml_element(
    element,
    tag: str,
    attributes: dict[str, str],
    child_tags: list[str],
    *,
    text: str | None = None,
):
    children = list(element)
    valid_text = element.text == text if text is not None else _only_vml_whitespace(
        element.text
    )
    if (
        element.tag != tag
        or element.attrib != attributes
        or [child.tag for child in children] != child_tags
        or not valid_text
        or not _only_vml_whitespace(element.tail)
    ):
        raise _unsupported_active_workbook()
    return children


def _require_comment_vml_coordinate(element, tag: str, maximum: int) -> None:
    _require_comment_vml_element(element, tag, {}, [], text=element.text)
    if element.text is None or re.fullmatch(r"0|[1-9][0-9]*", element.text) is None:
        raise _unsupported_active_workbook()
    if int(element.text) > maximum:
        raise _unsupported_active_workbook()


def _only_vml_whitespace(value: str | None) -> bool:
    return value is None or not value.strip()


def _expected_part_type(name: str) -> str | None:
    if name in {"_rels/.rels", "xl/_rels/workbook.xml.rels"} or re.fullmatch(
        r"xl/worksheets/_rels/sheet[1-9][0-9]*\.xml\.rels", name
    ):
        return "application/vnd.openxmlformats-package.relationships+xml"
    if name in _PART_TYPES:
        return _PART_TYPES[name]
    if re.fullmatch(r"xl/worksheets/sheet[1-9][0-9]*\.xml", name):
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"
    if re.fullmatch(r"xl/theme/theme[1-9][0-9]*\.xml", name):
        return "application/vnd.openxmlformats-officedocument.theme+xml"
    if re.fullmatch(r"xl/tables/table[1-9][0-9]*\.xml", name):
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.table+xml"
    if re.fullmatch(r"xl/comments/comment[1-9][0-9]*\.xml", name):
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.comments+xml"
    if re.fullmatch(r"xl/drawings/commentsDrawing[1-9][0-9]*\.vml", name):
        return "application/vnd.openxmlformats-officedocument.vmlDrawing"
    return None


def _xlsx_used_bounds(contents: bytes, sheet_title: str) -> tuple[int, int]:
    del sheet_title  # All worksheets are bounded; openpyxl selects the first safely.
    max_row = 1
    max_column = 1
    cell_count = 0
    with ZipFile(BytesIO(contents)) as archive:
        for info in archive.infolist():
            if not (
                info.filename.startswith("xl/worksheets/")
                and info.filename.endswith(".xml")
            ):
                continue
            root = SafeElementTree.fromstring(archive.read(info.filename))
            for cell in root.iter():
                if not cell.tag.endswith("}c"):
                    continue
                reference = cell.attrib.get("r", "")
                letters = "".join(character for character in reference if character.isalpha())
                digits = "".join(character for character in reference if character.isdigit())
                if not letters or not digits:
                    raise _invalid_file()
                row_number = int(digits)
                column_number = column_index_from_string(letters)
                cell_count += 1
                if (
                    cell_count > _MAX_XLSX_CELL_RECORDS
                    or row_number > _MAX_XLSX_SOURCE_ROW
                    or column_number > _MAX_XLSX_USED_COLUMNS
                ):
                    raise _workbook_too_complex()
                max_row = max(max_row, row_number)
                max_column = max(max_column, column_number)
    return max_row, max_column


def _sparse_worksheet_rows(workbook, worksheet):
    """Yield stored cells only so a far coordinate cannot force dense allocation."""
    with worksheet._get_source() as source:
        parser = WorkSheetParser(
            source,
            worksheet._shared_strings,
            data_only=workbook.data_only,
            epoch=workbook.epoch,
            date_formats=workbook._date_formats,
            timedelta_formats=workbook._timedelta_formats,
        )
        for row_number, cells in parser.parse():
            yield SparseSourceRow(
                row_number,
                {
                    cell["column"] - 1: cell["value"]
                    for cell in cells
                    if not _blank(cell["value"])
                },
            )


def _bounded_rows(rows) -> list[tuple[int, dict[int, object]]]:
    table: list[tuple[int, dict[int, object]]] = []
    data_rows = 0
    for source_row, raw_row in enumerate(rows, start=1):
        if isinstance(raw_row, SparseSourceRow):
            source_row = raw_row.row
            row = raw_row.values
        else:
            row = {
                index: value
                for index, value in enumerate(raw_row)
                if not _blank(value)
            }
        if not row:
            continue
        if table:
            data_rows += 1
            if data_rows > MAX_IMPORT_ROWS:
                raise ImportServiceError(
                    "research_import_too_many_rows",
                    "The subject import exceeds the 5000 row limit.",
                    413,
                )
        table.append((source_row, row))
    return table


def _build_preview(
    table: list[tuple[int, dict[int, object]]], subject_type: str
) -> ImportPreview:
    if not table:
        return ImportPreview(
            subject_type,
            (),
            (
                ImportRowError(
                    1, "名称", "empty_import", "The import contains no subject rows."
                ),
            ),
        )
    header_row, raw_headers = table[0]
    headers, header_errors = _headers(raw_headers, header_row)
    used_columns = sorted({index for _, values in table for index in values})
    for index in used_columns:
        raw_header = raw_headers.get(index)
        if not _blank(raw_header):
            continue
        for row_number, values in table[1:]:
            value = values.get(index)
            if not _blank(value):
                header_errors.append(
                    ImportRowError(
                        row_number,
                        get_column_letter(index + 1),
                        "missing_header_for_data",
                        "A non-empty cell must have a supported column header.",
                    )
                )
                break
    if header_errors:
        exact = tuple(header_errors)
        return ImportPreview(subject_type, (), exact, header_errors=exact)
    rows: list[ImportRow] = []
    errors: list[ImportRowError] = []
    for row_number, values in table[1:]:
        by_field = {
            field: values.get(index)
            for index, field in headers.items()
        }
        row = ImportRow(
            row_number,
            _normalized_text(by_field.get("name")),
            _normalized_optional_text(by_field.get("description"), empty=""),
            _normalized_optional_text(by_field.get("parent_name"), empty=None),
        )
        rows.append(row)
        errors.extend(_validate_import_row(row, subject_type))
    if not rows:
        errors.append(
            ImportRowError(
                1, "名称", "empty_import", "The import contains no subject rows."
            )
        )
    return ImportPreview(subject_type, tuple(rows), tuple(errors), header_errors=())


def _headers(
    raw_headers: dict[int, object], row_number: int
) -> tuple[dict[int, str], list[ImportRowError]]:
    headers: dict[int, str] = {}
    errors: list[ImportRowError] = []
    seen: set[str] = set()
    for index, raw_header in sorted(raw_headers.items()):
        if _blank(raw_header):
            continue
        if not isinstance(raw_header, str):
            errors.append(
                ImportRowError(
                    row_number,
                    f"列{index + 1}",
                    "unknown_column",
                    "The import column is not supported.",
                )
            )
            continue
        header = raw_header.strip()
        if header not in IMPORT_COLUMNS:
            errors.append(
                ImportRowError(
                    row_number,
                    header,
                    "unknown_column",
                    "The import column is not supported.",
                )
            )
            continue
        if header in seen:
            errors.append(
                ImportRowError(
                    row_number,
                    header,
                    "duplicate_column",
                    "The import column appears more than once.",
                )
            )
            continue
        seen.add(header)
        headers[index] = IMPORT_COLUMNS[header]
    if "名称" not in seen:
        errors.append(
            ImportRowError(
                row_number,
                "名称",
                "missing_required_column",
                "The import must contain the 名称 column.",
            )
        )
    return headers, errors


def _validate_import_row(
    row: ImportRow, subject_type: str
) -> list[ImportRowError]:
    errors: list[ImportRowError] = []
    if not isinstance(row.name, str) or not row.name or len(row.name) > 160:
        errors.append(
            ImportRowError(
                row.row,
                "名称",
                "invalid_name",
                "A subject name between 1 and 160 characters is required.",
            )
        )
    if not isinstance(row.description, str):
        errors.append(
            ImportRowError(
                row.row,
                "说明",
                "invalid_description",
                "The subject description must be text.",
            )
        )
    if row.parent_name is not None and not isinstance(row.parent_name, str):
        errors.append(
            ImportRowError(
                row.row,
                "父级",
                "invalid_parent",
                "The parent subject name must be text.",
            )
        )
    elif row.parent_name and subject_type in {"department", "opportunity"}:
        errors.append(
            ImportRowError(
                row.row,
                "父级",
                "invalid_parent",
                "This subject type cannot reference an imported parent name.",
            )
        )
    return errors


def _load_parent_candidates(session, project_id, subject_type, rows):
    if subject_type not in {"role", "process"}:
        return {}
    names = sorted(
        {
            row.parent_name
            for row in rows
            if isinstance(row.parent_name, str) and row.parent_name
        }
    )
    if not names:
        return {}
    subjects = list(
        session.scalars(
            select(ProjectResearchSubject)
            .where(
                ProjectResearchSubject.project_id == project_id,
                ProjectResearchSubject.subject_type == "department",
                ProjectResearchSubject.status == "active",
                ProjectResearchSubject.name.in_(names),
            )
            .order_by(ProjectResearchSubject.name, ProjectResearchSubject.id)
            .with_for_update()
        )
    )
    candidates: dict[str, list[ProjectResearchSubject]] = {}
    for subject in subjects:
        candidates.setdefault(subject.name, []).append(subject)
    return candidates


def _resolve_parent(row, parents):
    if not row.parent_name:
        return None
    candidates = parents.get(row.parent_name, [])
    if not candidates:
        raise _invalid_import(
            [
                ImportRowError(
                    row.row,
                    "父级",
                    "parent_not_found",
                    "The parent research subject was not found.",
                )
            ]
        )
    if len(candidates) > 1:
        raise _invalid_import(
            [
                ImportRowError(
                    row.row,
                    "父级",
                    "ambiguous_parent",
                    "More than one parent research subject has this name.",
                )
            ]
        )
    return candidates[0].id


def _normalized_text(value: object) -> object:
    return value.strip() if isinstance(value, str) else value


def _normalized_optional_text(value: object, *, empty):
    if value is None:
        return empty
    if isinstance(value, str):
        value = value.strip()
        return value if value else empty
    return value


def _blank(value: object) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _subject_key(preview_id: str, subject_type: str, row: int) -> str:
    return f"import_{subject_type}_{preview_id}_{row}"


def _serializer() -> URLSafeTimedSerializer:
    settings = current_app.config["SETTINGS"]
    return URLSafeTimedSerializer(settings.jwt_secret, salt=_TOKEN_SALT)


def _token_claims(token: str) -> dict[str, str]:
    if not isinstance(token, str) or not token:
        raise _preview_not_found()
    try:
        claims = _serializer().loads(token, max_age=PREVIEW_TTL_SECONDS)
    except SignatureExpired:
        raise _preview_expired() from None
    except BadSignature:
        raise _preview_not_found() from None
    if (
        not isinstance(claims, dict)
        or not isinstance(claims.get("preview_id"), str)
        or not isinstance(claims.get("project_id"), str)
    ):
        raise _preview_not_found()
    return claims


def _redis() -> Redis:
    return current_app.extensions["fde_api_redis"]


def _preview_key(preview_id: str) -> str:
    return _PREVIEW_KEY_PREFIX + sha256(preview_id.encode()).hexdigest()


def _lock_key(preview_id: str) -> str:
    return _LOCK_KEY_PREFIX + sha256(preview_id.encode()).hexdigest()


def _fence_key(preview_id: str) -> str:
    return _FENCE_KEY_PREFIX + sha256(preview_id.encode()).hexdigest()


def _load_state(key: str) -> dict[str, Any]:
    try:
        raw = _redis().get(key)
    except RedisError:
        raise _storage_unavailable() from None
    if raw is None:
        raise _preview_expired()
    try:
        state = json.loads(raw)
    except (TypeError, UnicodeDecodeError, json.JSONDecodeError):
        raise _preview_not_found() from None
    if not isinstance(state, dict):
        raise _preview_not_found()
    return state


def _acquire_commit_lease(
    state_key: str,
    preview_id: str,
    expected_version: int,
    *,
    owner: str | None = None,
) -> CommitLease:
    owner = owner or token_urlsafe(24)
    lock_key = _lock_key(preview_id)
    fence_key = _fence_key(preview_id)
    script = """
local raw = redis.call('get', KEYS[1])
if not raw then return {-1, 0} end
local ttl = redis.call('pttl', KEYS[1])
if ttl <= 0 then return {-1, 0} end
local state = cjson.decode(raw)
if state['status'] == 'committed' then return {-2, ttl} end
if redis.call('exists', KEYS[2]) == 1 then return {0, ttl} end
local fence = redis.call('incr', KEYS[3])
redis.call('pexpire', KEYS[3], ttl)
local lock_value = ARGV[1] .. ':' .. tostring(fence)
redis.call('set', KEYS[2], lock_value, 'PX', ARGV[3])
state['status'] = 'committing'
state['commit_owner'] = ARGV[1]
state['commit_fence'] = fence
state['expected_version'] = tonumber(ARGV[2])
redis.call('set', KEYS[1], cjson.encode(state), 'PX', ttl)
return {fence, ttl}
"""
    try:
        fence, ttl_ms = _redis().eval(
            script,
            3,
            state_key,
            lock_key,
            fence_key,
            owner,
            expected_version,
            _COMMIT_LOCK_TTL_SECONDS * 1000,
        )
    except RedisError:
        raise _storage_unavailable() from None
    fence = int(fence)
    if fence == -1:
        raise _preview_expired()
    if fence <= 0:
        raise ImportServiceError(
            "research_import_commit_in_progress",
            "This subject import is already being committed.",
            409,
        )
    return CommitLease(state_key, lock_key, owner, fence, max(int(ttl_ms), 1))


def _mark_commit_state(lease: CommitLease, result: dict[str, Any]) -> bool:
    script = """
local raw = redis.call('get', KEYS[1])
if not raw then return 0 end
local state = cjson.decode(raw)
if state['status'] ~= 'committing'
   or state['commit_owner'] ~= ARGV[1]
   or tonumber(state['commit_fence']) ~= tonumber(ARGV[2]) then return 0 end
local ttl = redis.call('pttl', KEYS[1])
if ttl <= 0 then return 0 end
state['status'] = 'committed'
state['result'] = cjson.decode(ARGV[3])
state['commit_owner'] = nil
state['commit_fence'] = nil
redis.call('set', KEYS[1], cjson.encode(state), 'PX', ttl)
local lock_value = ARGV[1] .. ':' .. ARGV[2]
if redis.call('get', KEYS[2]) == lock_value then redis.call('del', KEYS[2]) end
return 1
"""
    try:
        return bool(
            _redis().eval(
                script,
                2,
                lease.state_key,
                lease.lock_key,
                lease.owner,
                lease.fence,
                json.dumps(result, ensure_ascii=False, separators=(",", ":")),
            )
        )
    except RedisError:
        raise _storage_unavailable() from None


def _reset_commit_state(lease: CommitLease) -> bool:
    script = """
local raw = redis.call('get', KEYS[1])
if not raw then return 0 end
local state = cjson.decode(raw)
if state['status'] ~= 'committing'
   or state['commit_owner'] ~= ARGV[1]
   or tonumber(state['commit_fence']) ~= tonumber(ARGV[2]) then return 0 end
local ttl = redis.call('pttl', KEYS[1])
if ttl <= 0 then return 0 end
state['status'] = 'preview'
state['commit_owner'] = nil
state['commit_fence'] = nil
state['expected_version'] = nil
redis.call('set', KEYS[1], cjson.encode(state), 'PX', ttl)
local lock_value = ARGV[1] .. ':' .. ARGV[2]
if redis.call('get', KEYS[2]) == lock_value then redis.call('del', KEYS[2]) end
return 1
"""
    try:
        return bool(
            _redis().eval(
                script,
                2,
                lease.state_key,
                lease.lock_key,
                lease.owner,
                lease.fence,
            )
        )
    except RedisError:
        return False


def _release_commit_lease(lease: CommitLease) -> None:
    script = "if redis.call('get', KEYS[1]) == ARGV[1] then return redis.call('del', KEYS[1]) else return 0 end"
    try:
        _redis().eval(
            script,
            1,
            lease.lock_key,
            f"{lease.owner}:{lease.fence}",
        )
    except RedisError:
        pass


def _result_from_state(state: dict[str, Any], *, replay: bool) -> ImportResult:
    raw = state.get("result")
    if not isinstance(raw, dict):
        raise _preview_not_found()
    subjects = raw.get("subjects")
    version = raw.get("project_version")
    if not isinstance(subjects, list) or type(version) is not int:
        raise _preview_not_found()
    return ImportResult(tuple(subjects), version, idempotent_replay=replay)


def _row_from_dict(value: object) -> ImportRow:
    if not isinstance(value, dict) or type(value.get("row")) is not int:
        raise _preview_not_found()
    return ImportRow(
        value["row"],
        value.get("name"),
        value.get("description"),
        value.get("parent_name"),
    )


def _error_from_dict(value: object) -> ImportRowError:
    if not isinstance(value, dict):
        raise _preview_not_found()
    try:
        return ImportRowError(
            int(value["row"]),
            str(value["column"]),
            str(value["code"]),
            str(value["message"]),
        )
    except (KeyError, TypeError, ValueError):
        raise _preview_not_found() from None


def _invalid_import(errors: list[ImportRowError]) -> ImportServiceError:
    return ImportServiceError(
        "invalid_research_import",
        "The subject import contains invalid rows.",
        422,
        {"errors": [error.to_dict() for error in errors]},
    )


def _invalid_request() -> ImportServiceError:
    return ImportServiceError(
        "invalid_request", "The request body is invalid.", 400
    )


def _invalid_file() -> ImportServiceError:
    return ImportServiceError(
        "invalid_research_import_file",
        "The subject import file is invalid.",
        422,
    )


def _workbook_too_complex() -> ImportServiceError:
    return ImportServiceError(
        "research_import_workbook_too_complex",
        "The XLSX workbook exceeds safe processing limits.",
        413,
    )


def _unsupported_active_workbook() -> ImportServiceError:
    return ImportServiceError(
        "unsupported_import_format",
        "Workbooks containing active or external content are not supported.",
        415,
    )


def _preview_not_found() -> ImportServiceError:
    return ImportServiceError(
        "research_import_preview_not_found",
        "The subject import preview was not found.",
        404,
    )


def _preview_expired() -> ImportServiceError:
    return ImportServiceError(
        "research_import_preview_expired",
        "The subject import preview has expired.",
        410,
    )


def _storage_unavailable() -> ImportServiceError:
    return ImportServiceError(
        "research_import_storage_unavailable",
        "Subject import previews are temporarily unavailable.",
        503,
    )
