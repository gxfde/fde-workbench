from __future__ import annotations

import hashlib
import io
import json
from copy import deepcopy
from datetime import UTC, datetime
from typing import Any
from uuid import uuid4

from docx import Document
from flask import current_app
from sqlalchemy import func, select
from sqlalchemy.orm import selectinload

from fde_api.ai.deepseek import DeepSeekClient, DeepSeekError
from fde_api.auth.models import User
from fde_api.config import Settings
from fde_api.extensions import db, object_storage
from fde_api.documents.catalog import RESOURCE_TEMPLATES_DIR
from fde_api.documents.models import DocumentTemplate, DocumentTemplateVersion
from fde_api.files.models import ProjectFile, ProjectFileVersion
from fde_api.files.names import safe_filename
from fde_api.jobs.outbox import enqueue_outbox
from fde_api.projects.events import record_event
from fde_api.projects.permissions import project_access
from fde_api.research.models import ProjectResearchExport, ProjectResearchForm, ProjectResearchFormRevision, ProjectResearchSubject
from fde_api.workbench.models import Project

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


class ResearchExportError(Exception):
    def __init__(self, code: str, message: str, status: int):
        super().__init__(message); self.code, self.message, self.status = code, message, status


def request_research_export(*, actor: User, project_id: str, form_id: str) -> tuple[dict[str, Any], bool]:
    """Persist an export request and enqueue it without waiting for AI or DOCX work."""
    session = db.session()
    try:
        project = session.scalar(select(Project).options(selectinload(Project.members)).where(Project.id == project_id))
        access = project_access(actor, project) if project else None
        if project is None or access is None or not access.can_view:
            raise ResearchExportError("project_not_found", "项目不存在或无权访问。", 404)
        if not (access.can_manage or access.can_update_assigned_tasks):
            raise ResearchExportError("forbidden", "你没有权限导出当前调研结果。", 403)
        form = session.scalar(
            select(ProjectResearchForm).options(
                selectinload(ProjectResearchForm.subject),
                selectinload(ProjectResearchForm.current_revision).selectinload(ProjectResearchFormRevision.answers),
            ).where(ProjectResearchForm.project_id == project.id, ProjectResearchForm.id == form_id)
        )
        if form is None or form.subject.status != "active" or form.subject.subject_type == "opportunity":
            raise ResearchExportError("research_form_not_found", "调研表不存在或所属对象已停用。", 404)
        context = [_form_context(form)] if form.current_revision else []
        if not context or not context[0]["answers"]:
            raise ResearchExportError("research_answers_required", "当前调研表还没有可导出的调研内容。", 400)

        # Lock the form row so concurrent confirmations reuse the same active request.
        session.scalar(select(ProjectResearchForm.id).where(ProjectResearchForm.id == form.id).with_for_update())
        active = session.scalar(
            select(ProjectResearchExport)
            .where(ProjectResearchExport.form_id == form.id, ProjectResearchExport.status.in_(("queued", "generating")))
            .order_by(ProjectResearchExport.created_at.desc())
            .limit(1)
        )
        if active is not None:
            session.commit()
            return serialize_research_export(active, form.name), False

        display_name = _research_result_name(form.subject.name, form.name)
        project_file = _find_research_result_file(session, project.id, form.id)
        if project_file is None:
            project_file = ProjectFile(
                project_id=project.id,
                category="document",
                business_category="调研",
                display_name=display_name,
                description="由单份调研表答案快照生成",
                tags_json={"research_form_id": form.id},
                created_by_user_id=actor.id,
            )
            session.add(project_file)
            session.flush()
        else:
            display_name = project_file.display_name

        export = ProjectResearchExport(
            project_id=project.id,
            form_id=form.id,
            requested_by_user_id=actor.id,
            status="queued",
            project_file_id=project_file.id,
            display_name=display_name,
        )
        session.add(export)
        session.flush()
        enqueue_outbox(session, "research.export", export.id, {"export_id": export.id})
        session.commit()
        return serialize_research_export(export, form.name), True
    except ResearchExportError:
        session.rollback()
        raise
    finally:
        session.close()


def get_research_export(*, actor: User, project_id: str, export_id: str) -> dict[str, Any]:
    session = db.session()
    try:
        project = session.scalar(select(Project).options(selectinload(Project.members)).where(Project.id == project_id))
        access = project_access(actor, project) if project else None
        if project is None or access is None or not access.can_view:
            raise ResearchExportError("project_not_found", "项目不存在或无权访问。", 404)
        export = session.scalar(
            select(ProjectResearchExport)
            .where(ProjectResearchExport.id == export_id, ProjectResearchExport.project_id == project.id)
        )
        if export is None:
            raise ResearchExportError("research_export_not_found", "调研结果生成任务不存在。", 404)
        form = session.get(ProjectResearchForm, export.form_id)
        return serialize_research_export(export, form.name if form else "调研表")
    finally:
        session.close()


def serialize_research_export(export: ProjectResearchExport, form_name: str) -> dict[str, Any]:
    return {
        "id": export.id,
        "project_id": export.project_id,
        "form_id": export.form_id,
        "form_name": form_name,
        "status": export.status,
        "file_id": export.project_file_id,
        "version_id": export.file_version_id,
        "version_number": export.version_number,
        "display_name": export.display_name,
        "failure_code": export.failure_code,
        "failure_message": export.failure_message,
        "created_at": export.created_at.isoformat() if export.created_at else None,
        "updated_at": export.updated_at.isoformat() if export.updated_at else None,
    }


def export_research_result(*, actor: User, project_id: str, form_id: str) -> dict[str, Any]:
    session = db.session()
    try:
        project = session.scalar(select(Project).options(selectinload(Project.members)).where(Project.id == project_id))
        access = project_access(actor, project) if project else None
        if project is None or access is None or not access.can_view:
            raise ResearchExportError("project_not_found", "项目不存在或无权访问。", 404)
        if not (access.can_manage or access.can_update_assigned_tasks):
            raise ResearchExportError("forbidden", "你没有权限导出当前调研结果。", 403)
        form = session.scalar(
            select(ProjectResearchForm).options(
                selectinload(ProjectResearchForm.subject),
                selectinload(ProjectResearchForm.current_revision).selectinload(ProjectResearchFormRevision.answers),
            ).where(ProjectResearchForm.project_id == project.id, ProjectResearchForm.id == form_id)
        )
        if form is None or form.subject.status != "active" or form.subject.subject_type == "opportunity":
            raise ResearchExportError("research_form_not_found", "调研表不存在或所属对象已停用。", 404)
        subject = form.subject
        context = [_form_context(form)] if form.current_revision else []
        if not context or not context[0]["answers"]:
            raise ResearchExportError("research_answers_required", "当前调研表还没有可导出的调研内容。", 400)
        summary = _ai_summary(project, subject, context)
        data = _build_docx(project, subject, context, summary, session=session)
        display_name = _research_result_name(subject.name, form.name)
        existing = _find_research_result_file(session, project.id, form.id)
        with session.begin_nested():
            if existing is None:
                existing = ProjectFile(project_id=project.id, category="document", business_category="调研", display_name=display_name, description="由单份调研表答案快照生成", tags_json={"research_form_id": form.id}, created_by_user_id=actor.id)
                session.add(existing); session.flush()
            else:
                display_name = existing.display_name
            number = int(session.scalar(select(func.coalesce(func.max(ProjectFileVersion.version_number), 0)).where(ProjectFileVersion.file_id == existing.id)) or 0) + 1
            filename = f"{display_name}-v{number}.docx"
            safe = safe_filename(filename)
            storage_key = f"projects/{project.id}/library/{existing.id}/v{number}/{safe}"
            stored = object_storage.current.put_stream(storage_key, io.BytesIO(data), DOCX_MIME, {"source": "research-export", "subject-id": subject.id, "form-id": form.id})
            version = ProjectFileVersion(
                file_id=existing.id, version_number=number, source="system_generated",
                original_filename=filename, safe_filename=safe, extension=".docx", mime_type=DOCX_MIME,
                bucket=getattr(object_storage.current, "bucket_name", None) or "local", storage_key=storage_key,
                size_bytes=stored.size, etag=stored.etag, sha256=hashlib.sha256(data).hexdigest(),
                scan_status="not_required", preview_status="none", uploaded_by_user_id=actor.id, status="available",
            )
            session.add(version); session.flush(); existing.current_version_id = version.id; existing.version += 1
            for old in session.scalars(select(ProjectFileVersion).where(ProjectFileVersion.file_id == existing.id, ProjectFileVersion.id != version.id, ProjectFileVersion.status == "available")).all():
                old.status = "deprecated"; old.deprecated_by_user_id = actor.id; old.deprecated_at = datetime.now(UTC); old.deprecation_reason = "已生成更新版本"
            try:
                enqueue_outbox(session, "file.preview", version.id, {"version_id": version.id})
            except Exception:  # noqa: BLE001 - preview failure must not block export
                pass
            record_event(session, actor, "research_result_exported", existing, {"file_id": existing.id, "version_id": version.id, "subject_id": subject.id, "form_id": form.id})
        session.commit()
        return {"file_id": existing.id, "version_id": version.id, "version_number": number, "display_name": display_name}
    except ResearchExportError:
        session.rollback(); raise
    finally:
        session.close()


def _research_result_name(subject_name: str, form_name: str) -> str:
    return f"调研结果-{subject_name}-{form_name}"[:255]


def _find_research_result_file(session, project_id: str, form_id: str) -> ProjectFile | None:
    files = session.scalars(
        select(ProjectFile).where(
            ProjectFile.project_id == project_id,
            ProjectFile.category == "document",
        )
    ).all()
    return next(
        (
            item
            for item in files
            if isinstance(item.tags_json, dict) and item.tags_json.get("research_form_id") == form_id
        ),
        None,
    )


def _descendant_ids(subjects, root_id: str) -> set[str]:
    result: set[str] = set(); queue = [root_id]
    while queue:
        parent = queue.pop(0)
        for item in subjects:
            if item.parent_subject_id == parent and item.subject_type != "opportunity" and item.id not in result:
                result.add(item.id); queue.append(item.id)
    return result


def _form_context(form: ProjectResearchForm) -> dict[str, Any]:
    revision = form.current_revision; definition = revision.definition_snapshot if revision else {}
    labels: dict[str, str] = {}
    for section in definition.get("sections", []) if isinstance(definition, dict) else []:
        if isinstance(section, dict):
            for field in section.get("fields", []):
                if isinstance(field, dict) and isinstance(field.get("field_key"), str): labels[field["field_key"]] = str(field.get("name") or field["field_key"])
    answers = [] if revision is None else [{"field_key": answer.field_key, "question": labels.get(answer.field_key, answer.field_key), "value": answer.value_json} for answer in revision.answers if answer.value_json not in (None, "", [], {})]
    return {"form_id": form.id, "form_name": form.name, "subject_name": form.subject.name, "answers": answers}


def _ai_summary(project: Project, subject: ProjectResearchSubject, forms: list[dict[str, Any]]) -> dict[str, Any]:
    messages = [{"role": "system", "content": "你是 FDE 调研报告助手。仅依据提供的答案生成 JSON，字段 executive_summary、key_findings、recommendations；后两项为字符串数组。信息不足要明确写待确认，不得编造。"}, {"role": "user", "content": json.dumps({"enterprise": project.enterprise_name, "subject": subject.name, "forms": forms}, ensure_ascii=False)}]
    try:
        raw = _client().complete_json(messages=messages, request_id=f"research-export-{uuid4()}", max_tokens=4096)
    except DeepSeekError as error:
        status = {"ai_timeout": 504, "ai_output_invalid": 502}.get(error.code, 503)
        raise ResearchExportError(error.code, error.public_message, status) from error
    summary = raw.get("executive_summary"); findings = raw.get("key_findings"); recommendations = raw.get("recommendations")
    if not isinstance(summary, str) or not isinstance(findings, list) or not isinstance(recommendations, list) or not all(isinstance(item, str) for item in findings + recommendations):
        raise ResearchExportError("ai_output_invalid", "AI 返回的调研总结格式不正确，请重试。", 502)
    return {"executive_summary": summary, "key_findings": findings, "recommendations": recommendations}


def _build_docx(
    project: Project,
    subject: ProjectResearchSubject,
    forms: list[dict[str, Any]],
    summary: dict[str, Any],
    *,
    session,
) -> bytes:
    """Render the latest published research template and inject the dynamic body."""
    template_bytes = _published_research_template(session)
    document = Document(io.BytesIO(template_bytes)) if template_bytes else Document()
    values = {
        "企业名称": project.enterprise_name,
        "项目名称": project.name,
        "调研对象": subject.name,
        "生成时间": datetime.now(UTC).astimezone().strftime("%Y-%m-%d %H:%M"),
        "调研摘要": summary["executive_summary"],
        "关键发现": "\n".join(f"• {item}" for item in summary["key_findings"]),
        "建议": "\n".join(f"• {item}" for item in summary["recommendations"]),
    }
    _replace_template_values(document, values)
    if not _insert_research_content(document, forms):
        # A custom template may omit the reserved insertion marker. Preserve its
        # header/footer and append the dynamic answers rather than dropping data.
        document.add_heading("调研记录", level=2)
        _append_research_content(document, forms)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _published_research_template(session) -> bytes | None:
    template = session.scalar(
        select(DocumentTemplate).where(DocumentTemplate.document_type == "research_result")
    )
    if template is not None and template.latest_published_version_number is not None:
        version = session.scalar(
            select(DocumentTemplateVersion).where(
                DocumentTemplateVersion.template_id == template.id,
                DocumentTemplateVersion.version_number == template.latest_published_version_number,
                DocumentTemplateVersion.status == "published",
            )
        )
        if version is not None and version.docx_storage_key:
            stream = object_storage.current.open_stream(version.docx_storage_key)
            try:
                return stream.read()
            finally:
                stream.close()
        if version is not None and version.docx_file_version_id:
            file_version = session.get(ProjectFileVersion, version.docx_file_version_id)
            if file_version is not None and file_version.storage_key:
                stream = object_storage.current.open_stream(file_version.storage_key)
                try:
                    return stream.read()
                finally:
                    stream.close()
    fallback = RESOURCE_TEMPLATES_DIR / "research_result.docx"
    return fallback.read_bytes() if fallback.is_file() else None


def _replace_template_values(document: Document, values: dict[str, str]) -> None:
    def replace_paragraph(paragraph) -> None:
        text = paragraph.text
        replaced = text
        for label, value in values.items():
            replaced = replaced.replace(f"【待填写：{label}】", value)
        if replaced != text:
            if paragraph.runs:
                paragraph.runs[0].text = replaced
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.text = replaced

    def walk_table(table) -> None:
        seen: set[object] = set()
        for row in table.rows:
            for cell in row.cells:
                # Keep the XML element alive while traversing. Using
                # ``id(cell._tc)`` allows Python to reuse ids for temporary
                # cell wrappers, which can skip unrelated cells and leave
                # template placeholders unreplaced.
                marker = cell._tc
                if marker in seen:
                    continue
                seen.add(marker)
                for paragraph in cell.paragraphs:
                    replace_paragraph(paragraph)
                for nested in cell.tables:
                    walk_table(nested)

    for paragraph in document.paragraphs:
        replace_paragraph(paragraph)
    for table in document.tables:
        walk_table(table)
    for section in document.sections:
        for container in (section.header, section.footer):
            for paragraph in container.paragraphs:
                replace_paragraph(paragraph)
            for table in container.tables:
                walk_table(table)


def _insert_research_content(document: Document, forms: list[dict[str, Any]]) -> bool:
    marker = next((p for p in document.paragraphs if "【待插入：调研内容】" in p.text), None)
    if marker is None:
        return False
    fragment = Document()
    _append_research_content(fragment, forms)
    for element in list(fragment.element.body):
        if element.tag.endswith("sectPr"):
            continue
        marker._p.addprevious(deepcopy(element))
    marker._element.getparent().remove(marker._element)
    return True


def _append_research_content(document: Document, forms: list[dict[str, Any]]) -> None:
    for form in forms:
        document.add_heading(f'{form["subject_name"]} · {form["form_name"]}', level=3)
        for answer in form["answers"]:
            document.add_paragraph(answer["question"], style="List Bullet")
            document.add_paragraph(_stringify(answer["value"]))


def _stringify(value: Any) -> str:
    if isinstance(value, list): return "、".join(_stringify(item) for item in value)
    if isinstance(value, dict): return "；".join(f"{key}：{_stringify(item)}" for key, item in value.items())
    return str(value)


def _client() -> DeepSeekClient:
    injected = current_app.extensions.get("fde_api_research_export_deepseek")
    if injected is not None: return injected
    settings: Settings = current_app.config["SETTINGS"]
    return DeepSeekClient(base_url=settings.deepseek_base_url, api_key=settings.deepseek_api_key, model=settings.deepseek_research_export_model, timeout_seconds=settings.deepseek_timeout_seconds)
