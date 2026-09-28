from __future__ import annotations

from dataclasses import dataclass
from typing import BinaryIO

from docx import Document


class PresurveyParseError(ValueError):
    pass


@dataclass(frozen=True, slots=True)
class ExtractedBlock:
    source_ref: str
    text: str


@dataclass(frozen=True, slots=True)
class ExtractedPresurvey:
    blocks: tuple[ExtractedBlock, ...]

    @property
    def source_refs(self) -> set[str]:
        return {block.source_ref for block in self.blocks}


def extract_docx(stream: BinaryIO) -> ExtractedPresurvey:
    try:
        document = Document(stream)
    except Exception as error:  # python-docx exposes several parser exceptions
        raise PresurveyParseError("无法读取 DOCX 预调研表。") from error
    blocks: list[ExtractedBlock] = []
    for index, paragraph in enumerate(document.paragraphs, 1):
        text = _normalize(paragraph.text)
        if text:
            blocks.append(ExtractedBlock(f"paragraph-{index}", text))
    for table_index, table in enumerate(document.tables, 1):
        for row_index, row in enumerate(table.rows, 1):
            for cell_index, cell in enumerate(row.cells, 1):
                text = _normalize(cell.text)
                if text:
                    blocks.append(ExtractedBlock(
                        f"table-{table_index}-row-{row_index}-cell-{cell_index}", text
                    ))
    if not blocks:
        raise PresurveyParseError("DOCX 中没有可分析的文字内容。")
    return ExtractedPresurvey(tuple(blocks))


def _normalize(value: str) -> str:
    return " ".join(value.replace("\u3000", " ").split())
