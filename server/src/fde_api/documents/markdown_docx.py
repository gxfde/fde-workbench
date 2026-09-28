"""Offline Markdown blocks to native Word content at a template placeholder."""
from __future__ import annotations

import io
import os
from pathlib import Path
import subprocess
import sys
import tempfile

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from markdown_it import MarkdownIt


def diagram_png(source: str) -> bytes:
    with tempfile.TemporaryDirectory(prefix="fde-mermaid-") as directory:
        src, png = Path(directory) / "source.txt", Path(directory) / "diagram.png"
        src.write_text(source, encoding="utf-8")
        result = subprocess.run(
            [sys.executable, "-m", "fde_api.documents.mermaid_render", str(src), str(png)],
            capture_output=True, timeout=45, check=False,
        )
        if result.returncode or not png.is_file():
            raise ValueError("流程图无法导出，请检查 Mermaid 语法或服务器渲染配置。")
        return png.read_bytes()


def inline(paragraph, children):
    bold = italic = False
    for token in children or []:
        if token.type == "strong_open": bold = True
        elif token.type == "strong_close": bold = False
        elif token.type == "em_open": italic = True
        elif token.type == "em_close": italic = False
        elif token.type in ("softbreak", "hardbreak"):
            paragraph.add_run().add_break()
        elif token.type in ("text", "code_inline", "html_inline", "image"):
            content = token.content
            if token.type == "html_inline" and content.lower().replace(" ", "") in ("<br>", "<br/>"):
                paragraph.add_run().add_break(); continue
            run = paragraph.add_run(content)
            run.bold, run.italic = bold, italic
            if token.type == "code_inline": run.font.name = "Consolas"


def insert_markdown(document, anchor, source: str):
    tokens = MarkdownIt("commonmark", {"html": False}).enable("table").parse(source)
    section = document.sections[0]
    width = section.page_width - section.left_margin - section.right_margin
    height = section.page_height - section.top_margin - section.bottom_margin - Inches(1.1)
    lists = []; item_pending = False; current = None; table = row = cell = None
    in_header = False
    for token in tokens:
        kind = token.type
        if kind in ("bullet_list_open", "ordered_list_open"):
            lists.append([kind, int(token.attrGet("start") or 1)])
        elif kind in ("bullet_list_close", "ordered_list_close"): lists.pop()
        elif kind == "list_item_open": item_pending = True
        elif kind == "list_item_close": item_pending = False
        elif kind in ("paragraph_open", "heading_open"):
            current = anchor.insert_paragraph_before(style=anchor.style)
            fmt = current.paragraph_format
            fmt.space_before = Pt(0); fmt.space_after = Pt(6)
            fmt.line_spacing = 1.15; fmt.keep_with_next = kind == "heading_open"
            if lists:
                fmt.left_indent = Inches(.18 * len(lists))
                if item_pending:
                    ordered = lists[-1][0] == "ordered_list_open"
                    current.add_run(f"{lists[-1][1]}. " if ordered else "• ")
                    if ordered: lists[-1][1] += 1
                    item_pending = False
            if kind == "heading_open":
                current.style = document.styles[f"Heading {min(int(token.tag[1:]), 4)}"]
        elif kind == "inline":
            if cell is not None: inline(cell.paragraphs[0], token.children)
            elif current is not None: inline(current, token.children)
        elif kind == "table_open":
            table = document.add_table(rows=0, cols=0)
            anchor._p.addprevious(table._tbl)
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                e = OxmlElement(f"w:{edge}")
                for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
                    e.set(qn(f"w:{key}"), value)
                borders.append(e)
            table._tbl.tblPr.append(borders)
        elif kind == "thead_open": in_header = True
        elif kind == "thead_close": in_header = False
        elif kind == "tr_open":
            row = table.add_row(); column_index = 0
            if in_header: row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        elif kind in ("th_open", "td_open"):
            if column_index >= len(table.columns): table.add_column(Inches(1))
            cell = row.cells[column_index]; column_index += 1
            if in_header:
                shading = OxmlElement("w:shd"); shading.set(qn("w:fill"), "EAF0F7")
                cell._tc.get_or_add_tcPr().append(shading)
            fmt = cell.paragraphs[0].paragraph_format
            fmt.space_after = Pt(5); fmt.space_before = Pt(5); fmt.line_spacing = 1.1
        elif kind in ("th_close", "td_close"):
            for run in cell.paragraphs[0].runs:
                run.font.size = Pt(10)
                if in_header: run.bold = True
            cell = None
        elif kind == "table_close":
            # Allocate more width to verbose columns while keeping every column readable.
            weights = [max(6, min(35, max(len(r.cells[i].text) for r in table.rows)**.6)) for i in range(len(table.columns))]
            for i, weight in enumerate(weights):
                size = int(width * weight / sum(weights)); table.columns[i].width = size
                for r in table.rows: r.cells[i].width = size
            table.autofit = False; table = None
        elif kind in ("fence", "code_block"):
            current = anchor.insert_paragraph_before(style=anchor.style)
            current.paragraph_format.space_after = Pt(8)
            current.paragraph_format.keep_with_next = False
            if token.info.strip().lower() == "mermaid":
                image = diagram_png(token.content)
                from docx.image.image import Image
                info = Image.from_blob(image)
                scale = min(width / info.width, height / info.height)
                current.alignment = WD_ALIGN_PARAGRAPH.CENTER
                current.add_run().add_picture(io.BytesIO(image), width=int(info.width * scale), height=int(info.height * scale))
            else: current.add_run(token.content).font.name = "Consolas"
    anchor._p.getparent().remove(anchor._p)
