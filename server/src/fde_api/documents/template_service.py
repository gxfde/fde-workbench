"""Admin-only document template service.

Template versions are immutable system-level artifacts composed of a source DOCX
plus the mapping/section metadata
that the document generator consumes. All entry points here are admin-gated by
the routes (``require_minimum_role("admin")``); this module focuses on the
persistence rules and validation, raising :class:`DocumentTemplateServiceError`
for every user-addressable failure.
"""

from __future__ import annotations

import hashlib
import io
import re
from collections.abc import Mapping
from datetime import UTC, datetime
from typing import Any

from sqlalchemy import func, select
from sqlalchemy.orm import Session, selectinload

from fde_api.auth.models import ROLE_ADMIN, User
from fde_api.documents.catalog import (
    DOCUMENT_TYPES,
    RESOURCE_TEMPLATES_DIR,
    load_document_catalog,
)
from fde_api.documents.models import (
    DocumentTemplate,
    DocumentTemplateVersion,
)
from fde_api.documents.template_parser import scan_docx_placeholders
from fde_api.extensions import db, object_storage
from fde_api.files.names import safe_filename

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
DOCX_EXTENSION = ".docx"
DOWNLOAD_EXPIRES_SECONDS = 300

SYSTEM_USERNAME = "system"
SYSTEM_INDUSTRY = "通用"

# Scan issues that make a template untrustworthy or unrenderable.
BLOCKING_ISSUE_CODES = frozenset(
    {
        "unsafe_template_expression",
        "unknown_template_root",
        "invalid_template_expression",
        "unbalanced_loop",
    }
)

# Residual Jinja tokens that should never survive generation.
_RESIDUAL_TOKEN_RE = re.compile(r"\{\{|\{%")


class DocumentTemplateServiceError(Exception):
    """A stable, user-facing document template failure."""

    def __init__(self, code: str, message: str, status: int):
        super().__init__(message)
        self.code = code
        self.message = message
        self.status = status


def list_document_templates(*, actor: User) -> list[dict[str, Any]]:
    """Return every document template with its full version chain (admin only)."""
    del actor  # admin enforcement lives in the route decorator
    session = db.session()
    try:
        templates = session.scalars(
            select(DocumentTemplate)
            .options(
                selectinload(DocumentTemplate.versions)
                .selectinload(DocumentTemplateVersion.created_by),
            )
            .order_by(DocumentTemplate.document_type, DocumentTemplate.id)
        )
        return [_serialize_template(template) for template in templates]
    finally:
        session.close()


def create_document_template_version(
    *,
    actor: User,
    document_type: str,
    upload: Any,
) -> dict[str, Any]:
    """Create a new draft version for a document type from an uploaded DOCX.

    ``upload`` may be raw ``bytes``, a readable binary stream, or an existing
    ``UploadSession``-like object exposing ``storage_key``.
    """
    entry = DOCUMENT_TYPES.get(document_type)
    if entry is None:
        raise DocumentTemplateServiceError(
            "unknown_document_type",
            f"Unknown document type '{document_type}'.",
            400,
        )

    docx_bytes = _read_upload_bytes(upload)
    if not docx_bytes:
        raise DocumentTemplateServiceError(
            "invalid_template_file", "The template file is empty.", 400
        )

    _reject_unsafe_placeholders(docx_bytes)

    session = db.session()
    try:
        with session.begin():
            template = _get_or_create_template(session, document_type, entry)
            version_number = _next_template_version_number(session, template.id)
            stored_docx = _store_template_docx(
                document_type, template.id, version_number, docx_bytes
            )
            version = DocumentTemplateVersion(
                template_id=template.id,
                version_number=version_number,
                status="draft",
                docx_file_version_id=None,
                **stored_docx,
                mapping_json=default_mapping(document_type, entry),
                sections_json=default_sections(entry),
                table_loops_json={},
                required_data_json={},
                source_sha256=sha256_hex(docx_bytes),
                created_by_user_id=actor.id,
            )
            session.add(version)
            session.flush()
            session.refresh(version)
        return _serialize_version(version)
    except DocumentTemplateServiceError:
        raise
    finally:
        session.close()


def test_generate_template_version(
    *, actor: User, version_id: str
) -> dict[str, Any]:
    """Validate and render sample data before checking for residual tokens."""
    del actor
    session = db.session()
    try:
        version = _load_version(session, version_id)
        data = _read_version_docx(version)
        text = _test_render_template(data)
        residual = _RESIDUAL_TOKEN_RE.findall(text)
        if residual:
            raise DocumentTemplateServiceError(
                "template_generation_failed",
                "Template still contains render placeholders; report failed.",
                400,
            )
        return {
            "ok": True,
            "generated_preview": text[:400],
            "residual_tokens": [],
        }
    finally:
        session.close()


def _test_render_template(data: bytes) -> str:
    # Never test by searching the unrendered template: its tokens are intentional.
    # Use only synthetic data, with no AI calls, project reads or file writes.
    from fde_api.documents.generator import render_docx
    from jinja2 import TemplateError

    _reject_unsafe_placeholders(data)
    context = {
        "project": {"name": "测试项目", "project_code": "TEST-001"},
        "enterprise": {"name": "测试企业"},
        "document": {"business_code": "SOW-测试方案", "version_number": 1,
                     "document_type": "SOW"},
        "research": {}, "subjects": [], "tasks": [], "members": [],
    }
    try:
        output = render_docx(io.BytesIO(data), context, [])
    except TemplateError as error:
        raise DocumentTemplateServiceError(
            "template_generation_failed", "Template sample rendering failed.", 400
        ) from error
    return _extract_docx_text(output.read())


def publish_template_version(*, actor: User, version_id: str) -> dict[str, Any]:
    """Publish a draft template version, freezing it and advancing the pointer."""
    session = db.session()
    try:
        with session.begin():
            version = _load_version(session, version_id)
            if version.status != "draft":
                raise DocumentTemplateServiceError(
                    "invalid_version_state",
                    "Only a draft version can be published.",
                    409,
                )
            version.status = "published"
            version.published_by_user_id = actor.id
            version.published_at = datetime.now(UTC)
            template = version.template
            template.latest_published_version_number = version.version_number
            session.flush()
            session.refresh(version)
        return _serialize_version(version)
    except DocumentTemplateServiceError:
        raise
    finally:
        session.close()


def copy_template_version(*, actor: User, template_id: str) -> dict[str, Any]:
    """Create a new draft version copying mapping/sections and source DOCX."""
    session = db.session()
    try:
        with session.begin():
            template = session.get(DocumentTemplate, template_id)
            if template is None:
                raise DocumentTemplateServiceError(
                    "template_not_found", "Document template was not found.", 404
                )
            versions = session.scalars(
                select(DocumentTemplateVersion)
                .where(DocumentTemplateVersion.template_id == template.id)
                .order_by(DocumentTemplateVersion.version_number.desc())
            ).all()
            if not versions:
                raise DocumentTemplateServiceError(
                    "invalid_version_state",
                    "Template has no version to copy.",
                    409,
                )
            source = versions[0]
            version = DocumentTemplateVersion(
                template_id=template.id,
                version_number=source.version_number + 1,
                status="draft",
                docx_file_version_id=source.docx_file_version_id,
                docx_original_filename=source.docx_original_filename,
                docx_mime_type=source.docx_mime_type,
                docx_bucket=source.docx_bucket,
                docx_storage_key=source.docx_storage_key,
                docx_size_bytes=source.docx_size_bytes,
                docx_etag=source.docx_etag,
                mapping_json=dict(source.mapping_json),
                sections_json=dict(source.sections_json),
                table_loops_json=dict(source.table_loops_json),
                required_data_json=dict(source.required_data_json),
                source_sha256=source.source_sha256,
                created_by_user_id=actor.id,
            )
            session.add(version)
            session.flush()
            session.refresh(version)
        return _serialize_version(version)
    except DocumentTemplateServiceError:
        raise
    finally:
        session.close()


def deactivate_template_version(*, actor: User, version_id: str) -> dict[str, Any]:
    """Mark a template version inactive (soft-delete)."""
    del actor
    session = db.session()
    try:
        with session.begin():
            version = _load_version(session, version_id)
            version.status = "inactive"
            session.flush()
            session.refresh(version)
        return _serialize_version(version)
    except DocumentTemplateServiceError:
        raise
    finally:
        session.close()


def get_template_version_download_url(*, actor: User, version_id: str) -> str:
    """Return a short-lived download URL for a template source DOCX.

    Older installations may still have catalog-seeded versions whose source
    only exists in the bundled resources directory. Materialize those bytes in
    object storage on first download so the same signed-download path works for
    both legacy and newly uploaded templates.
    """
    del actor
    session = db.session()
    try:
        with session.begin():
            version = _load_version(session, version_id)
            filename = (
                version.docx_original_filename
                or f"{version.template.document_type}-v{version.version_number}{DOCX_EXTENSION}"
            )
            storage_key = version.docx_storage_key
            if not storage_key and version.docx_file is not None:
                storage_key = version.docx_file.storage_key
                filename = version.docx_file.original_filename or filename
            if not storage_key:
                stored_docx = _store_template_docx(
                    version.template.document_type,
                    version.template_id,
                    version.version_number,
                    _read_version_docx(version),
                )
                for field, value in stored_docx.items():
                    setattr(version, field, value)
                storage_key = version.docx_storage_key
                filename = version.docx_original_filename or filename
                session.flush()
            if not storage_key:
                raise DocumentTemplateServiceError(
                    "template_docx_unavailable", "模板文件暂不可用。", 409
                )
            return object_storage.current.sign_download(
                storage_key,
                filename,
                expires_seconds=DOWNLOAD_EXPIRES_SECONDS,
            )
    except DocumentTemplateServiceError:
        raise
    finally:
        session.close()


def seed_document_templates(session: Session) -> dict[str, Any]:
    """Idempotently seed a published v1 version for every catalog key.

    Reads the built/copied DOCX from ``resources/document-templates`` for the
    ``source_sha256`` and the catalog metadata. Missing DOCX skip the key rather
    than fail the whole seed. Already-seeded (published v1) keys are untouched.
    """
    system_user = _get_or_create_system_user(session)
    created: list[str] = []
    skipped: list[str] = []
    missing: list[str] = []

    catalog = load_document_catalog()
    for document_type, entry in catalog.items():
        existing_template = session.scalar(
            select(DocumentTemplate).where(
                DocumentTemplate.document_type == document_type
            )
        )
        if existing_template is not None:
            published_v1 = session.scalar(
                select(DocumentTemplateVersion).where(
                    DocumentTemplateVersion.template_id == existing_template.id,
                    DocumentTemplateVersion.version_number == 1,
                    DocumentTemplateVersion.status == "published",
                )
            )
            if published_v1 is not None:
                skipped.append(document_type)
                continue

        resource_path = RESOURCE_TEMPLATES_DIR / f"{document_type}{DOCX_EXTENSION}"
        if not resource_path.is_file():
            missing.append(document_type)
            continue

        docx_bytes = resource_path.read_bytes()
        template = existing_template or DocumentTemplate(
            document_type=document_type,
            name=entry["name"],
            description="",
            industry_name=SYSTEM_INDUSTRY,
            status="active",
        )
        if template.id is None:
            session.add(template)
            session.flush()

        version = DocumentTemplateVersion(
            template_id=template.id,
            version_number=1,
            status="published",
            docx_file_version_id=None,
            mapping_json=default_mapping(document_type, entry),
            sections_json=default_sections(entry),
            table_loops_json={},
            required_data_json={},
            source_sha256=sha256_hex(docx_bytes),
            created_by_user_id=system_user.id,
            published_by_user_id=system_user.id,
            published_at=datetime.now(UTC),
        )
        session.add(version)
        template.latest_published_version_number = 1
        created.append(document_type)

    session.flush()
    return {"created": created, "skipped": skipped, "missing": missing}


# ---------------------------------------------------------------------------
# Serialization
# ---------------------------------------------------------------------------


def _serialize_template(template: DocumentTemplate) -> dict[str, Any]:
    return {
        "id": template.id,
        "document_type": template.document_type,
        "name": template.name,
        "description": template.description,
        "industry_name": template.industry_name,
        "status": template.status,
        "latest_published_version_number": template.latest_published_version_number,
        "versions": [
            _serialize_version(version)
            for version in sorted(template.versions, key=lambda item: item.version_number, reverse=True)
        ],
    }


def _serialize_version(version: DocumentTemplateVersion) -> dict[str, Any]:
    return {
        "id": version.id,
        "template_id": version.template_id,
        "version_number": version.version_number,
        "status": version.status,
        "docx_file_version_id": version.docx_file_version_id,
        "docx_original_filename": version.docx_original_filename,
        "docx_size_bytes": version.docx_size_bytes,
        "mapping_json": version.mapping_json,
        "sections_json": version.sections_json,
        "table_loops_json": version.table_loops_json,
        "required_data_json": version.required_data_json,
        "source_sha256": version.source_sha256,
        "created_by_user_id": version.created_by_user_id,
        "published_by_user_id": version.published_by_user_id,
        "published_at": (
            version.published_at.isoformat() if version.published_at else None
        ),
        "created_at": version.created_at.isoformat() if version.created_at else None,
    }


# ---------------------------------------------------------------------------
# Catalog metadata builders
# ---------------------------------------------------------------------------


def default_mapping(document_type: str, entry: Mapping[str, Any]) -> dict[str, Any]:
    return {
        "document_type": document_type,
        "name": entry["name"],
        "modules": list(entry.get("modules", [])),
        "field_map": {},
        "research_keys": [],
    }


def default_sections(entry: Mapping[str, Any]) -> dict[str, Any]:
    return {"sections": list(entry.get("required_sections", []) or [])}


# ---------------------------------------------------------------------------
# DOCX storage helpers
# ---------------------------------------------------------------------------


def _get_or_create_template(
    session: Session, document_type: str, entry: Mapping[str, Any]
) -> DocumentTemplate:
    existing = session.scalar(
        select(DocumentTemplate).where(DocumentTemplate.document_type == document_type)
    )
    if existing is not None:
        return existing
    template = DocumentTemplate(
        document_type=document_type,
        name=entry["name"],
        description="",
        industry_name=SYSTEM_INDUSTRY,
        status="active",
    )
    session.add(template)
    session.flush()
    return template


def _next_template_version_number(session: Session, template_id: str) -> int:
    current = session.scalar(
        select(func.max(DocumentTemplateVersion.version_number)).where(
            DocumentTemplateVersion.template_id == template_id
        )
    )
    return int(current or 0) + 1


def _store_template_docx(
    document_type: str,
    template_id: str,
    version_number: int,
    docx_bytes: bytes,
) -> dict[str, Any]:
    filename = f"{document_type}-v{version_number}{DOCX_EXTENSION}"
    safe = safe_filename(filename)
    storage_key = (
        f"document-templates/{template_id}/{document_type}"
        f"/versions/{version_number}/{safe}"
    )
    stored = object_storage.current.put_stream(
        storage_key,
        io.BytesIO(docx_bytes),
        DOCX_MIME,
        {},
    )
    return {
        "docx_original_filename": filename,
        "docx_mime_type": DOCX_MIME,
        "docx_bucket": _storage_bucket(),
        "docx_storage_key": storage_key,
        "docx_size_bytes": stored.size,
        "docx_etag": stored.etag,
    }


def _get_or_create_system_user(session: Session) -> User:
    """Return the persistent system seed user, creating it on first use."""
    existing = session.scalar(select(User).where(User.username == SYSTEM_USERNAME))
    if existing is not None:
        return existing
    user = User(
        username=SYSTEM_USERNAME,
        display_name="System",
        role=ROLE_ADMIN,
        password_hash="",
        must_change_password=True,
        is_active=True,
    )
    session.add(user)
    session.flush()
    return user


# ---------------------------------------------------------------------------
# Shared helpers
# ---------------------------------------------------------------------------


def _load_version(session: Session, version_id: str) -> DocumentTemplateVersion:
    version = session.scalar(
        select(DocumentTemplateVersion)
        .options(
            selectinload(DocumentTemplateVersion.template),
            selectinload(DocumentTemplateVersion.docx_file),
        )
        .where(DocumentTemplateVersion.id == version_id)
    )
    if version is None:
        raise DocumentTemplateServiceError(
            "version_not_found", "Document template version was not found.", 404
        )
    return version


def _read_upload_bytes(upload: Any) -> bytes:
    if isinstance(upload, bytes):
        return upload
    if hasattr(upload, "read"):
        data = upload.read()
        return data if isinstance(data, bytes) else bytes(data)
    storage_key = getattr(upload, "storage_key", None)
    if isinstance(storage_key, str) and storage_key:
        stream = object_storage.current.open_stream(storage_key)
        try:
            return stream.read()
        finally:
            stream.close()
    raise DocumentTemplateServiceError(
        "invalid_template_file", "The template file is missing.", 400
    )


def _read_version_docx(version: DocumentTemplateVersion) -> bytes:
    if version.docx_storage_key:
        stream = object_storage.current.open_stream(version.docx_storage_key)
        try:
            return stream.read()
        finally:
            stream.close()

    # Rolling-upgrade fallback for installations that have not migrated the
    # legacy project-file reference yet.
    docx_file = version.docx_file
    if docx_file is not None and docx_file.storage_key:
        stream = object_storage.current.open_stream(docx_file.storage_key)
        try:
            return stream.read()
        finally:
            stream.close()

    resource_path = RESOURCE_TEMPLATES_DIR / f"{version.template.document_type}{DOCX_EXTENSION}"
    if resource_path.is_file():
        return resource_path.read_bytes()
    raise DocumentTemplateServiceError(
        "template_docx_unavailable", "The template DOCX is not available.", 400
    )


def _extract_docx_text(data: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(data))
    texts: list[str] = []

    def _collect_paragraphs(paragraphs: object) -> None:
        for paragraph in paragraphs:
            if paragraph.text:
                texts.append(paragraph.text)

    def _collect_table(table: object, seen_cells: set[object]) -> None:
        for row in table.rows:
            for cell in row.cells:
                cell_key = cell._tc if hasattr(cell, "_tc") else cell
                if cell_key in seen_cells:
                    continue
                seen_cells.add(cell_key)
                _collect_paragraphs(cell.paragraphs)
                for nested in cell.tables:
                    _collect_table(nested, seen_cells)

    _collect_paragraphs(document.paragraphs)
    seen: set[object] = set()
    for table in document.tables:
        _collect_table(table, seen)
    for section in document.sections:
        if section.header is not None:
            _collect_paragraphs(section.header.paragraphs)
        if section.footer is not None:
            _collect_paragraphs(section.footer.paragraphs)
    return "\n".join(texts)


def _reject_unsafe_placeholders(docx_bytes: bytes) -> None:
    scan = scan_docx_placeholders(io.BytesIO(docx_bytes))
    blocking = [issue for issue in scan.issues if issue.code in BLOCKING_ISSUE_CODES]
    if blocking:
        code = "unsafe_template_expression"
        message = "The template contains unsafe or unknown template placeholders."
        first = blocking[0]
        if first.code == "unknown_template_root":
            code = "unknown_template_root"
            message = first.message
        elif first.code == "invalid_template_expression":
            code = "invalid_template_expression"
            message = first.message
        elif first.code == "unbalanced_loop":
            code = "unbalanced_loop"
            message = first.message
        raise DocumentTemplateServiceError(code, message, 400)


def _storage_bucket() -> str:
    storage = object_storage.current
    bucket = getattr(storage, "bucket_name", None)
    return bucket or "local"


def sha256_hex(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()
