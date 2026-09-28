"""Idempotent DOCX generation for project documents.

``render_docx`` renders a DOCX template (already carrying Jinja placeholders) with
a plain dict/list context under a closed sandbox, appends sanitized rich-text
blocks, and returns the rendered document as a ``BytesIO`` — never touching the
template source.

``handle_document_generation`` is the body of the ``document.generate`` worker
handler. It is idempotent: a version that already generated a DOCX (or was
confirmed/archived) returns its current state without re-rendering. A QA failure
records ``generation_error`` and marks the generation job failed without creating
a file version; success writes one immutable DOCX file version and enqueues a
best-effort preview.
"""

from __future__ import annotations

import hashlib
import io
import json
import re
from datetime import date, datetime, timedelta
from typing import Any, BinaryIO
from zoneinfo import ZoneInfo

import docxtpl
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from jinja2.sandbox import SandboxedEnvironment
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from fde_api.documents.catalog import RESOURCE_TEMPLATES_DIR, load_document_catalog
from fde_api.documents.ai_content import generate_document_fields
from fde_api.documents.models import (
    DocumentGenerationJob,
    DocumentTemplateVersion,
    ProjectDocument,
    ProjectDocumentVersion,
)
from fde_api.documents.qa import inspect_generated_docx
from fde_api.documents.rich_text import RichTextDocument, render_rich_text
from fde_api.documents.snapshots import build_document_snapshot
from fde_api.extensions import db, object_storage
from fde_api.files.models import ProjectFile, ProjectFileVersion
from fde_api.files.names import document_key, extension_of, safe_filename
from fde_api.jobs.outbox import enqueue_outbox

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
DOCX_EXTENSION = ".docx"

# The document states considered terminal success for idempotency purposes.
_SUCCESS_STATES = frozenset({"confirmed", "archived"})


class DocumentGenerationError(Exception):
    """A stable generation failure that should surface to the caller."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


class ClosedSandbox(SandboxedEnvironment):
    """A sandbox that denies every attribute traversal and callable.

    Returning ``False`` from ``is_safe_attribute`` blocks attribute chains such as
    ``cycler.__init__.__globals__.os.system``. Item lookups (``project.name`` on a
    dict) still resolve through the container-item path, so benign placeholders
    keep working while the attribute-escape surface is closed.
    """

    def is_safe_attribute(self, obj: object, attr: str, value: object) -> bool:
        return False

    def is_safe_callable(self, obj: object) -> bool:
        return False


def render_docx(template_stream: BinaryIO, context: dict, rich_text) -> BinaryIO:
    """Render ``template_stream`` with ``context`` and return a ``BytesIO``.

    ``context`` must be plain dict/list primitives (never ORM objects). ``rich_text``
    may be a :class:`RichTextDocument`, a list of sanitized block dicts, a single
    block dict, or empty. The rendered content is appended at the end of the body.
    """
    template = docxtpl.DocxTemplate(template_stream)
    template.render(context, jinja_env=ClosedSandbox())

    document = template.docx
    _replace_solution_markdown(document, context)
    _replace_chinese_placeholders(document, context)
    _replace_project_plan_markers(document, context)
    _append_rich_text(document, rich_text)
    _constrain_tables_to_page(document)

    output = io.BytesIO()
    template.save(output)
    output.seek(0)
    return output


def _constrain_tables_to_page(document) -> None:
    """Cap malformed table grids to the printable page width.

    Some customer-authored templates contain ``w:gridCol`` and ``w:tcW`` values
    written in EMU even though Word expects twips. LibreOffice silently shrinks
    those tables while Word lets them run beyond the right page edge. Preserve
    the existing column proportions, but scale any over-wide table to the
    printable width and force fixed layout so long generated text wraps.
    """
    if not document.sections:
        return
    section = document.sections[0]
    usable_width = max(
        1,
        int(
            (
                section.page_width
                - section.left_margin
                - section.right_margin
            )
            / 635
        ),
    )

    def constrain(table, seen: set[object]) -> None:
        key = table._tbl
        if key in seen:
            return
        seen.add(key)

        table_properties = table._tbl.tblPr
        indent = table_properties.find(qn("w:tblInd"))
        indent_width = 0
        if indent is not None:
            try:
                indent_width = max(0, int(indent.get(qn("w:w")) or 0))
            except (TypeError, ValueError):
                indent_width = 0
        target_width = max(1, usable_width - indent_width)

        grid = table._tbl.tblGrid
        columns = list(grid.gridCol_lst) if grid is not None else []
        widths: list[int] = []
        for column in columns:
            try:
                widths.append(max(0, int(column.get(qn("w:w")) or 0)))
            except (TypeError, ValueError):
                widths.append(0)
        grid_width = sum(widths)
        scale = target_width / grid_width if grid_width > target_width else 1.0

        if scale < 1.0:
            scaled_widths = _scaled_widths(widths, target_width)
            for column, width in zip(columns, scaled_widths, strict=False):
                column.set(qn("w:w"), str(width))
            for cell_width in table._tbl.iter(qn("w:tcW")):
                try:
                    width = int(cell_width.get(qn("w:w")) or 0)
                except (TypeError, ValueError):
                    continue
                if width > 0:
                    cell_width.set(qn("w:w"), str(max(1, round(width * scale))))

        table_width = table_properties.find(qn("w:tblW"))
        if table_width is None:
            table_width = OxmlElement("w:tblW")
            table_properties.insert(0, table_width)
        table_width.set(qn("w:type"), "dxa")
        table_width.set(qn("w:w"), str(min(grid_width or target_width, target_width)))

        layout = table_properties.find(qn("w:tblLayout"))
        if layout is None:
            layout = OxmlElement("w:tblLayout")
            table_properties.append(layout)
        layout.set(qn("w:type"), "fixed")

        for row in table.rows:
            for cell in row.cells:
                for nested in cell.tables:
                    constrain(nested, seen)

    seen: set[object] = set()
    for table in document.tables:
        constrain(table, seen)
    for section in document.sections:
        for part in (section.header, section.footer):
            for table in part.tables:
                constrain(table, seen)


def _scaled_widths(widths: list[int], target: int) -> list[int]:
    if not widths or sum(widths) <= 0:
        return widths
    scaled = [max(1, round(width * target / sum(widths))) for width in widths]
    scaled[-1] += target - sum(scaled)
    return scaled


def handle_document_generation(*, version_id: str) -> dict[str, Any]:
    """Generate a DOCX for ``version_id`` exactly once; never raises.

    Returns a result dict like ``{"status": "succeeded", "sha256": ...,
    "docx_file_version_id": ...}`` or ``{"status": "failed",
    "generation_error": ...}``.
    """
    session = db.session()
    try:
        with session.begin():
            version = _load_version(session, version_id)
            if version is None:
                return {
                    "status": "failed",
                    "generation_error": "version_not_found",
                }

            if _is_already_generated(version):
                return {
                    "status": "succeeded",
                    "sha256": version.sha256,
                    "docx_file_version_id": version.docx_file_version_id,
                }

            document = version.document
            template_bytes = _load_template_docx(
                session, document.document_type, version.source_template_version
            )
            context = _build_context(session, version, document, template_bytes)
            rich_text = _content_rich_text(version)

            rendered = render_docx(io.BytesIO(template_bytes), context, rich_text)
            required_sections = _required_sections(document.document_type, version.source_template_version)
            qa = inspect_generated_docx(
                io.BytesIO(_stream_bytes(rendered)), required_sections
            )
            if qa.errors:
                message = "; ".join(f"{e.code}: {e.message}" for e in qa.errors)
                version.generation_error = message
                _mark_job(session, version.id, "failed", message)
                session.flush()
                return {"status": "failed", "generation_error": message}

            rendered_bytes = _stream_bytes(rendered)
            sha256 = hashlib.sha256(rendered_bytes).hexdigest()
            file_version = _ensure_docx_file_version(
                session, version, document, rendered_bytes, sha256
            )
            version.docx_file_version_id = file_version.id
            version.sha256 = sha256
            version.generation_error = ""
            _mark_job(session, version.id, "succeeded", None)

            try:
                enqueue_outbox(
                    session, "file.preview", file_version.id, {"version_id": file_version.id}
                )
            except Exception:  # noqa: BLE001 - preview is best-effort, never fails generation
                pass

            session.flush()
            return {
                "status": "succeeded",
                "sha256": sha256,
                "docx_file_version_id": file_version.id,
            }
    except Exception as exc:  # noqa: BLE001 - a generation failure must never escape
        message = getattr(exc, "public_message", None) or str(exc) or "文档生成失败，请重试。"
        session.rollback()
        try:
            with session.begin():
                failed_version = session.get(ProjectDocumentVersion, version_id)
                if failed_version is not None:
                    failed_version.generation_error = message
                    _mark_job(session, version_id, "failed", message)
        except Exception:  # noqa: BLE001 - best effort failure recording
            session.rollback()
        return {"status": "failed", "generation_error": message}
    finally:
        session.close()


def register_document_generation_handler() -> None:
    """Register the ``document.generate`` worker handler (idempotent)."""
    from fde_api.jobs.handlers import register_handler

    register_handler("document.generate", _document_generate_handler)


def _document_generate_handler(job) -> None:
    """Background job body: generate ``job.target_id``; never raises."""
    try:
        handle_document_generation(version_id=job.target_id)
    except Exception:  # noqa: BLE001 - a job body must not re-raise to the worker
        pass


# ---------------------------------------------------------------------------
# Rendering helpers
# ---------------------------------------------------------------------------


def _append_rich_text(document, rich_text) -> None:
    for block in _normalize_blocks(rich_text):
        block_type = block.get("type")
        content = block.get("content")
        if block_type == "heading":
            document.add_heading(str(content), level=1)
        elif block_type in ("ordered_list", "bullet_list"):
            style = "List Number" if block_type == "ordered_list" else "List Bullet"
            items = content if isinstance(content, list) else [content]
            for item in items:
                if item is None or str(item).strip() == "":
                    continue
                document.add_paragraph(str(item), style=style)
        elif block_type == "table":
            _append_table(document, content if isinstance(content, list) else [])
        else:
            # paragraph and any unknown block type default to a plain paragraph.
            text = "" if content is None else str(content)
            document.add_paragraph(text)


def _normalize_blocks(rich_text) -> list[dict]:
    if isinstance(rich_text, RichTextDocument):
        return render_rich_text(rich_text)
    if isinstance(rich_text, dict):
        if "type" in rich_text:
            return [rich_text]
        return []
    if isinstance(rich_text, list):
        return [block for block in rich_text if isinstance(block, dict)]
    return []


def _append_table(document, rows: list) -> None:
    normalized = [row for row in rows if isinstance(row, list)]
    if not normalized:
        return
    column_count = max((len(row) for row in normalized), default=0)
    if column_count == 0:
        return
    table = document.add_table(rows=len(normalized), cols=column_count)
    try:
        table.style = "Table Grid"
    except Exception:  # noqa: BLE001 - style availability is not guaranteed
        pass
    for row_index, row in enumerate(normalized):
        for column_index in range(column_count):
            value = row[column_index] if column_index < len(row) else ""
            table.cell(row_index, column_index).text = "" if value is None else str(value)


_PROJECT_PLAN_MARKERS = {
    "【待插入：企业与项目信息】": "enterprise",
    "【待插入：客户目标与项目指引】": "guidance",
    "【待插入：任务甘特图】": "gantt",
    "【待插入：任务清单】": "tasks",
}


def _replace_project_plan_markers(document, context: dict[str, Any]) -> None:
    """Replace project-plan markers with data-driven tables at their position."""
    for paragraph in list(document.paragraphs):
        marker = paragraph.text.strip()
        section = _PROJECT_PLAN_MARKERS.get(marker)
        if section is None:
            continue
        elements = _project_plan_elements(document, section, context)
        cursor = paragraph._p
        for element in elements:
            cursor.addnext(element)
            cursor = element
        parent = paragraph._p.getparent()
        if parent is not None:
            parent.remove(paragraph._p)


def _project_plan_elements(document, section: str, context: dict[str, Any]) -> list[Any]:
    project = context.get("project") if isinstance(context.get("project"), dict) else {}
    enterprise = context.get("enterprise") if isinstance(context.get("enterprise"), dict) else {}
    guidance = context.get("guidance") if isinstance(context.get("guidance"), dict) else {}
    tasks = context.get("tasks") if isinstance(context.get("tasks"), list) else []
    if section == "enterprise":
        rows = [
            ["项目名称", project.get("name", ""), "项目编号", project.get("project_code", "")],
            ["企业名称", enterprise.get("name", ""), "行业", project.get("industry_name", "")],
            ["联系人", enterprise.get("contact_name", ""), "联系方式", enterprise.get("contact_phone", "")],
            ["企业地址", enterprise.get("address", ""), "项目负责人", project.get("leader_name", "")],
            ["计划周期", _project_period(project), "当前状态", _project_status_label(project.get("status"))],
            ["项目背景", project.get("background", ""), "备注", project.get("notes", "")],
        ]
        return [_new_table(document, rows, header=False, widths=[1.15, 3.35, 1.15, 4.15])._tbl]
    if section == "guidance":
        if not guidance:
            return [_new_paragraph(document, "暂无已确认的客户目标与项目指引。")]
        rows: list[list[Any]] = [
            ["客户长期愿景", guidance.get("customer_vision", "")],
            ["当前阶段目标", guidance.get("current_phase_objective", "")],
        ]
        for key, label in (
            ("key_business_problems", "优先关注"),
            ("priority_departments", "优先部门"),
            ("priority_roles", "优先岗位"),
            ("priority_processes", "优先流程"),
            ("success_criteria", "成功标准"),
            ("next_actions", "下一步行动"),
        ):
            value = guidance.get(key)
            rows.append([label, "\n".join(f"• {item}" for item in value) if isinstance(value, list) else ""])
        return [_new_table(document, rows, header=False, widths=[1.5, 8.3])._tbl]
    if section == "gantt":
        return [_gantt_table(document, tasks)._tbl]
    rows = [["任务", "模块", "状态", "负责人", "进度", "计划", "前置任务"]]
    for task in tasks:
        if not isinstance(task, dict):
            continue
        assignee = task.get("assignee") if isinstance(task.get("assignee"), dict) else {}
        dependencies = task.get("dependencies") if isinstance(task.get("dependencies"), list) else []
        rows.append([
            task.get("name", ""), task.get("module_name", ""), _task_status_label(task.get("status")),
            assignee.get("display_name", "待分配"), f"{task.get('progress', 0)}%",
            f"{task.get('planned_start_date', '')} 至 {task.get('planned_end_date', '')}", "、".join(map(str, dependencies)),
        ])
    return [_new_table(document, rows, header=True, widths=[2.2, 1.15, 0.85, 1.0, 0.65, 1.65, 2.3], font_size=8)._tbl]


def _new_table(
    document,
    rows: list[list[Any]],
    *,
    header: bool,
    widths: list[float] | None = None,
    font_size: float = 9,
):
    column_count = max((len(row) for row in rows), default=1)
    table = document.add_table(rows=max(1, len(rows)), cols=column_count)
    table.style = "Table Grid"
    table.autofit = False
    if widths and len(widths) == column_count:
        for index, width in enumerate(widths):
            table.columns[index].width = Inches(width)
            for cell in table.columns[index].cells:
                cell.width = Inches(width)
    _set_table_geometry(table, widths)
    for row_index, row in enumerate(rows):
        for column_index in range(column_count):
            cell = table.cell(row_index, column_index)
            cell.text = str(row[column_index] if column_index < len(row) and row[column_index] is not None else "")
            _format_table_cell(cell, font_size=font_size, fill="E8EEF5" if header and row_index == 0 else None)
    if header and rows:
        for cell in table.rows[0].cells:
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.color.rgb = RGBColor.from_string("0B2545")
        header_properties = table.rows[0]._tr.get_or_add_trPr()
        repeat = OxmlElement("w:tblHeader")
        repeat.set(qn("w:val"), "true")
        header_properties.append(repeat)
    return table


def _set_table_geometry(table, widths: list[float] | None) -> None:
    properties = table._tbl.tblPr
    indent = properties.find(qn("w:tblInd"))
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), "120")
    layout = properties.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")
    if widths:
        grid = table._tbl.tblGrid
        for column, width in zip(grid.gridCol_lst, widths, strict=False):
            column.set(qn("w:w"), str(round(width * 1440)))


def _format_table_cell(cell, *, font_size: float, fill: str | None) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.find(qn("w:tcMar"))
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", 80), ("bottom", 80), ("start", 120), ("end", 120)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")
    if fill:
        shading = properties.find(qn("w:shd"))
        if shading is None:
            shading = OxmlElement("w:shd")
            properties.append(shading)
        shading.set(qn("w:fill"), fill)
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.08
        for run in paragraph.runs:
            run.font.name = "PingFang SC"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
            run.font.size = Pt(font_size)


def _new_paragraph(document, text: str):
    return document.add_paragraph(text)._p


def _gantt_table(document, raw_tasks: list[Any]):
    tasks = [task for task in raw_tasks if isinstance(task, dict)]
    starts = [_parse_date(task.get("planned_start_date")) for task in tasks]
    ends = [_parse_date(task.get("planned_end_date")) for task in tasks]
    valid_starts = [value for value in starts if value is not None]
    valid_ends = [value for value in ends if value is not None]
    if not valid_starts or not valid_ends:
        return _new_table(document, [["任务", "计划"], ["暂无已排期任务", ""]], header=True, widths=[3.0, 6.8])
    first, last = min(valid_starts), max(valid_ends)
    span = max(1, (last - first).days + 1)
    bucket_count = min(12, span)
    bucket_days = max(1, (span + bucket_count - 1) // bucket_count)
    buckets = [first + timedelta(days=index * bucket_days) for index in range(bucket_count)]
    rows: list[list[Any]] = [["任务"] + [value.strftime("%m-%d") for value in buckets]]
    for task in tasks:
        start, end = _parse_date(task.get("planned_start_date")), _parse_date(task.get("planned_end_date"))
        timeline = []
        for bucket in buckets:
            bucket_end = min(last, bucket + timedelta(days=bucket_days - 1))
            timeline.append("■" if start is not None and end is not None and start <= bucket_end and end >= bucket else "")
        rows.append([task.get("name", "")] + timeline)
    timeline_width = 7.15 / bucket_count
    return _new_table(document, rows, header=True, widths=[2.65] + [timeline_width] * bucket_count, font_size=7.5)


def _parse_date(value: Any) -> date | None:
    try:
        return date.fromisoformat(str(value))
    except (TypeError, ValueError):
        return None


def _project_period(project: dict[str, Any]) -> str:
    return f"{project.get('planned_start_date', '')} 至 {project.get('planned_end_date') or '待排期'}"


def _project_status_label(value: Any) -> str:
    return {"draft": "草稿", "active": "进行中", "paused": "已暂停", "completed": "已完成", "cancelled": "已取消"}.get(str(value), str(value or ""))


def _task_status_label(value: Any) -> str:
    return {"not_started": "未开始", "in_progress": "进行中", "blocked": "阻塞", "completed": "已完成", "cancelled": "已取消"}.get(str(value), str(value or ""))


_CHINESE_PLACEHOLDER = re.compile(r"【待填写：\s*([^】]+?)\s*】")
_CHINESE_FIELD_ALIASES = {
    "项目名称": "project.name",
    "项目编号": "project.project_code",
    "企业名称": "enterprise.name",
    "联系人": "enterprise.contact_name",
    "联系方式": "enterprise.contact_phone",
    "生成日期": "generated_date",
    "AI机会名称": "opportunity_name",
    "AI 机会名称": "opportunity_name",
    "现状描述": "current_situation",
    "业务目标": "business_goal",
    "目标使用对象": "target_audience",
    "方案范围": "solution_scope",
    "交付物": "deliverables",
    "成功指标": "success_metrics",
    "验收标准": "acceptance_criteria",
    "时间计划": "timeline",
    "不包含项": "exclusions",
    "前提假设": "assumptions",
    "数据要求": "data_requirements",
    "安全要求": "security_requirements",
    "风险控制": "risk_controls",
    "下一步行动": "next_action",
    "方案名称": "solution_name",
    "关联AI机会": "associated_opportunities",
    "方案设计": "design_markdown",
    "预期指标与验收标准": "acceptance_criteria",
    "所需数据与系统": "data_systems",
    "计划排期": "schedule",
    "风险与依赖": "risks_dependencies",
    "来源方案版本": "solution_version",
    "待确认商务与责任事项": "commercial_terms",
}


def _replace_solution_markdown(document, context: dict[str, Any]) -> None:
    """Render common Markdown blocks at the plan placeholder, without HTML execution.

    The canonical Markdown stays unchanged in the immutable source snapshot.
    Unsupported notation is retained as readable literal text rather than lost.
    """
    markdown = context.get("design_markdown")
    if not isinstance(markdown, str) or not markdown.strip():
        return

    def add_inline(paragraph, text: str):
        for token in re.split(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^\)]+\))", text):
            if not token:
                continue
            if token.startswith("**") and token.endswith("**"):
                paragraph.add_run(token[2:-2]).bold = True
            elif token.startswith("`") and token.endswith("`"):
                paragraph.add_run(token[1:-1]).font.name = "Consolas"
            else:
                link = re.fullmatch(r"\[([^\]]+)\]\(([^\)]+)\)", token)
                paragraph.add_run(f"{link[1]}（{link[2]}）" if link else token)

    def replace(paragraph):
        if paragraph.text.strip() != "【待填写：方案设计】":
            return
        from fde_api.documents.markdown_docx import insert_markdown
        insert_markdown(document, paragraph, markdown)

    for paragraph in list(document.paragraphs):
        replace(paragraph)
    seen: set[object] = set()

    def table_rows(table):
        for row in table.rows:
            for cell in row.cells:
                if cell._tc in seen:
                    continue
                seen.add(cell._tc)
                for paragraph in list(cell.paragraphs):
                    replace(paragraph)
                for nested in cell.tables:
                    table_rows(nested)

    for table in document.tables:
        table_rows(table)


def _replace_chinese_placeholders(document, context: dict[str, Any]) -> None:
    def replace_paragraph(paragraph) -> None:
        original = paragraph.text
        if "【待填写：" not in original:
            return
        replaced = _CHINESE_PLACEHOLDER.sub(lambda match: _placeholder_value(context, match.group(1), match.group(0)), original)
        if replaced == original:
            return
        if paragraph.runs:
            paragraph.runs[0].text = replaced
            for run in paragraph.runs[1:]:
                run.text = ""
        else:
            paragraph.add_run(replaced)

    def replace_table(table, seen: set[object]) -> None:
        for row in table.rows:
            for cell in row.cells:
                # Keep the XML element itself alive while traversing. Using
                # ``id(cell._tc)`` lets Python reuse ids for later temporary
                # cell wrappers, which can make unrelated cells look like
                # duplicate merged cells and leave placeholders unreplaced.
                key = cell._tc
                if key in seen:
                    continue
                seen.add(key)
                for paragraph in cell.paragraphs:
                    replace_paragraph(paragraph)
                for nested in cell.tables:
                    replace_table(nested, seen)

    for paragraph in document.paragraphs:
        replace_paragraph(paragraph)
    seen: set[object] = set()
    for table in document.tables:
        replace_table(table, seen)
    for section in document.sections:
        for part in (section.header, section.footer):
            for paragraph in part.paragraphs:
                replace_paragraph(paragraph)
            for table in part.tables:
                replace_table(table, seen)


def _placeholder_value(context: dict[str, Any], label: str, fallback: str) -> str:
    path = _CHINESE_FIELD_ALIASES.get(label.strip(), label.strip())
    value: Any = context
    for segment in path.split("."):
        if not isinstance(value, dict) or segment not in value:
            return fallback
        value = value[segment]
    if value in (None, ""):
        return fallback
    if isinstance(value, list):
        return "、".join(str(item) for item in value)
    if isinstance(value, dict):
        return json.dumps(value, ensure_ascii=False)
    return str(value)


# ---------------------------------------------------------------------------
# Context / template loading
# ---------------------------------------------------------------------------


def _load_version(session: Session, version_id: str) -> ProjectDocumentVersion | None:
    return session.scalar(
        select(ProjectDocumentVersion)
        .options(
            selectinload(ProjectDocumentVersion.document),
            selectinload(ProjectDocumentVersion.source_template_version),
            selectinload(ProjectDocumentVersion.docx_file),
        )
        .where(ProjectDocumentVersion.id == version_id)
    )


def _is_already_generated(version: ProjectDocumentVersion) -> bool:
    return version.status in _SUCCESS_STATES or bool(version.sha256)


def _load_template_docx(
    session: Session,
    document_type: str,
    template_version: DocumentTemplateVersion | None,
) -> bytes:
    if template_version is not None and template_version.docx_storage_key:
        stream = object_storage.current.open_stream(template_version.docx_storage_key)
        try:
            return stream.read()
        finally:
            stream.close()

    # Legacy rolling-upgrade fallback. New template uploads are not project files.
    if template_version is not None and template_version.docx_file_version_id:
        file_version = session.get(
            ProjectFileVersion, template_version.docx_file_version_id
        )
        if file_version is not None and file_version.storage_key:
            stream = object_storage.current.open_stream(file_version.storage_key)
            try:
                return stream.read()
            finally:
                stream.close()

    # Fallback to the on-disk built template if no DOCX file version is bound.
    resource_path = RESOURCE_TEMPLATES_DIR / f"{document_type}{DOCX_EXTENSION}"
    if resource_path.is_file():
        return resource_path.read_bytes()

    raise DocumentGenerationError(
        "template_docx_unavailable", "The template DOCX is not available."
    )


def _build_context(
    session: Session,
    version: ProjectDocumentVersion,
    document: ProjectDocument,
    template_bytes: bytes,
) -> dict[str, Any]:
    snapshot = version.source_snapshot_json
    if not isinstance(snapshot, dict) or not snapshot:
        mapping = (
            version.source_template_version.mapping_json
            if version.source_template_version is not None
            else {}
        )
        snapshot = build_document_snapshot(session, document.project_id, mapping)
    context = dict(snapshot)
    content_snapshot = version.content_snapshot_json if isinstance(version.content_snapshot_json, dict) else {}
    manual_fields = content_snapshot.get("field_overrides") if isinstance(content_snapshot.get("field_overrides"), dict) else {}
    requested_fields = [
        _CHINESE_FIELD_ALIASES.get(label, label)
        for label in _extract_chinese_placeholder_labels(template_bytes)
    ]
    # A solution export is a faithful rendering of an explicitly saved revision,
    # never a second AI design pass. Read the immutable version, not live links.
    has_saved_solution = isinstance(content_snapshot.get("solution"), dict) or isinstance(snapshot.get("solution"), dict)
    ai_fields = {} if has_saved_solution else generate_document_fields(session, document, requested_fields)
    merged_fields = dict(ai_fields)
    merged_fields.update(manual_fields)
    if merged_fields.get("generated_date") in (None, "", [], {}):
        merged_fields["generated_date"] = datetime.now(
            ZoneInfo("Asia/Shanghai")
        ).strftime("%Y年%m月%d日")
    for field_path in requested_fields:
        if field_path in merged_fields and merged_fields[field_path] not in (None, "", [], {}):
            continue
        if _context_value(context, field_path) not in (None, "", [], {}):
            continue
        # AI providers occasionally omit one requested key even when instructed
        # to return the complete schema.  Freeze an explicit, reviewable value
        # instead of leaving a template token behind and failing the whole job.
        merged_fields[field_path] = "待确认"
    if merged_fields:
        frozen_content = dict(content_snapshot)
        frozen_content["field_overrides"] = merged_fields
        if document.opportunity_scope_json and not has_saved_solution:
            frozen_content["opportunity_scope"] = document.opportunity_scope_json
        version.content_snapshot_json = frozen_content
    context["field_overrides"] = merged_fields
    for key, value in merged_fields.items():
        _set_context_default(context, key, value)
    context["document"] = {
        "business_code": document.business_code,
        "version_number": version.version_number,
        "document_type": document.document_type,
    }
    context["version_number"] = version.version_number
    context["document_type"] = document.document_type
    context["business_code"] = document.business_code
    return context


def _context_value(context: dict[str, Any], path: str) -> Any:
    value: Any = context
    for segment in path.split("."):
        if not isinstance(value, dict) or segment not in value:
            return None
        value = value[segment]
    return value


def _set_context_default(context: dict[str, Any], path: str, value: Any) -> None:
    segments = path.split(".")
    target = context
    for segment in segments[:-1]:
        child = target.get(segment)
        if not isinstance(child, dict):
            child = {}
            target[segment] = child
        target = child
    final = segments[-1]
    if target.get(final) in (None, "", [], {}):
        target[final] = value


def _extract_chinese_placeholder_labels(template_bytes: bytes) -> list[str]:
    """Collect custom ``【待填写：...】`` labels from body/header/footer text."""
    document = Document(io.BytesIO(template_bytes))
    labels: list[str] = []

    def collect_paragraphs(paragraphs) -> None:
        for paragraph in paragraphs:
            for match in _CHINESE_PLACEHOLDER.finditer(paragraph.text):
                label = match.group(1).strip()
                if label and label not in labels:
                    labels.append(label)

    def collect_table(table, seen: set[object]) -> None:
        for row in table.rows:
            for cell in row.cells:
                key = cell._tc
                if key in seen:
                    continue
                seen.add(key)
                collect_paragraphs(cell.paragraphs)
                for nested in cell.tables:
                    collect_table(nested, seen)

    collect_paragraphs(document.paragraphs)
    seen: set[object] = set()
    for table in document.tables:
        collect_table(table, seen)
    for section in document.sections:
        for part in (section.header, section.footer):
            collect_paragraphs(part.paragraphs)
            for table in part.tables:
                collect_table(table, seen)
    return labels


def _content_rich_text(version: ProjectDocumentVersion):
    content_snapshot = version.content_snapshot_json
    if isinstance(content_snapshot, dict):
        return content_snapshot.get("rich_text") or {}
    return {}


def _required_sections(document_type: str, template_version=None) -> list[str]:
    # Published SOW versions retain their own chapter contract. Updating the
    # catalog for solution designs must not invalidate historical templates.
    if document_type == "sow" and template_version is not None:
        sections = (template_version.sections_json or {}).get("sections", [])
        return [item for item in sections if isinstance(item, str)] if isinstance(sections, list) else []
    entry = load_document_catalog().get(document_type)
    if not entry:
        return []
    return list(entry.get("required_sections", []) or [])


# ---------------------------------------------------------------------------
# File version promotion
# ---------------------------------------------------------------------------


def _ensure_docx_file_version(
    session: Session,
    version: ProjectDocumentVersion,
    document: ProjectDocument,
    rendered_bytes: bytes,
    sha256: str,
) -> ProjectFileVersion:
    if version.docx_file_version_id:
        existing = session.get(ProjectFileVersion, version.docx_file_version_id)
        if existing is not None:
            return _write_docx_bytes(existing, rendered_bytes, sha256)

    filename = f"{document.business_code}-v{version.version_number}{DOCX_EXTENSION}"
    safe = safe_filename(filename)
    storage_key = document_key(document.project_id, document.id, version.id, filename)
    stored = object_storage.current.put_stream(
        storage_key, io.BytesIO(rendered_bytes), DOCX_MIME, {}
    )
    project_file = ProjectFile(
        project_id=document.project_id,
        category="document",
        display_name=filename,
        description="",
        created_by_user_id=version.generated_by_user_id,
    )
    session.add(project_file)
    session.flush()
    file_version = ProjectFileVersion(
        file_id=project_file.id,
        version_number=1,
        source="system_generated",
        original_filename=filename,
        safe_filename=safe,
        extension=extension_of(filename),
        mime_type=DOCX_MIME,
        bucket=_storage_bucket(),
        storage_key=storage_key,
        size_bytes=stored.size,
        etag=stored.etag,
        sha256=sha256,
        scan_status="not_required",
        preview_status="none",
        uploaded_by_user_id=version.generated_by_user_id,
        status="available",
    )
    session.add(file_version)
    session.flush()
    project_file.current_version_id = file_version.id
    return file_version


def _write_docx_bytes(
    file_version: ProjectFileVersion, rendered_bytes: bytes, sha256: str
) -> ProjectFileVersion:
    stored = object_storage.current.put_stream(
        file_version.storage_key, io.BytesIO(rendered_bytes), DOCX_MIME, {}
    )
    file_version.size_bytes = stored.size
    file_version.etag = stored.etag
    file_version.sha256 = sha256
    file_version.status = "available"
    file_version.scan_status = "not_required"
    file_version.source = "system_generated"
    return file_version


def _mark_job(
    session: Session, version_id: str, status: str, error: str | None
) -> None:
    job = session.scalar(
        select(DocumentGenerationJob)
        .where(DocumentGenerationJob.document_version_id == version_id, DocumentGenerationJob.status.in_(("queued", "running")))
        .order_by(DocumentGenerationJob.created_at.desc(), DocumentGenerationJob.id.desc())
        .limit(1)
    )
    if job is None:
        job = session.scalar(select(DocumentGenerationJob).where(DocumentGenerationJob.document_version_id == version_id)
            .order_by(DocumentGenerationJob.created_at.desc(), DocumentGenerationJob.id.desc()).limit(1))
    if job is not None:
        job.status = status
        job.last_error = error or ""


def _storage_bucket() -> str:
    storage = object_storage.current
    bucket = getattr(storage, "bucket_name", None)
    return bucket or "local"


def _stream_bytes(stream: BinaryIO) -> bytes:
    if hasattr(stream, "seek"):
        try:
            stream.seek(0)
        except Exception:  # noqa: BLE001 - seek is best-effort
            pass
    data = stream.read()
    return data if isinstance(data, bytes) else bytes(data)
