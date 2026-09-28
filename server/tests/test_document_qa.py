from __future__ import annotations

import io

from docx import Document

from fde_api.documents.qa import extract_docx_xml_text, inspect_generated_docx


def _docx_bytes(*paragraphs: str) -> io.BytesIO:
    doc = Document()
    for paragraph in paragraphs:
        doc.add_paragraph(paragraph)
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def _error_codes(result) -> list[str]:
    return [error.code for error in result.errors]


def test_qa_rejects_residual_placeholder():
    result = inspect_generated_docx(
        _docx_bytes("客户：{{ enterprise.name }}"), required_sections=()
    )

    assert _error_codes(result) == ["template_placeholder_remaining"]


def test_qa_rejects_missing_required_section():
    result = inspect_generated_docx(
        _docx_bytes("这是一份缺少关键章节的文档"), required_sections=["交付范围"]
    )

    assert "missing_required_section" in _error_codes(result)
    assert result.sections == []


def test_qa_accepts_clean_docx():
    result = inspect_generated_docx(
        _docx_bytes("交付范围", "验收标准"), required_sections=["交付范围"]
    )

    assert result.errors == []
    assert result.sections == ["交付范围"]
    assert "交付范围" in result.text


def test_qa_rejects_non_ooxml():
    result = inspect_generated_docx(io.BytesIO(b"not a zip"), required_sections=())

    assert "invalid_docx" in _error_codes(result)


def test_extract_docx_xml_text_includes_tables_and_headers():
    doc = Document()
    doc.add_paragraph("正文段落")
    table = doc.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "表格内容"
    section = doc.sections[0]
    section.header.paragraphs[0].text = "页眉内容"
    section.footer.paragraphs[0].text = "页脚内容"
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    text = extract_docx_xml_text(buffer)

    assert "正文段落" in text
    assert "表格内容" in text
    assert "页眉内容" in text
    assert "页脚内容" in text
