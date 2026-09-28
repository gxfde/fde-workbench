from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO

import pytest
from sqlalchemy import func, select

from fde_api.auth.models import User
from fde_api.auth.passwords import hash_password
from fde_api.auth.tokens import issue_access_token
from fde_api.documents.models import DocumentTemplate, DocumentTemplateVersion
from fde_api.documents.template_service import seed_document_templates


INITIAL_PASSWORD = "InitialPass!234"


def _build_docx_bytes(*, text: str = "{{ project.name }}") -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("测试模板", level=1)
    document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@dataclass(frozen=True)
class AuthorizedClient:
    client: object
    headers: dict[str, str]

    def get(self, *args, **kwargs):
        return self.client.get(*args, headers=self.headers, **kwargs)

    def post(self, *args, **kwargs):
        return self.client.post(*args, headers=self.headers, **kwargs)


@pytest.fixture
def authorized_client_factory(client, db_session, settings):
    def create(role: str) -> AuthorizedClient:
        user = User(
            username=f"{role}.doc-template",
            display_name=role,
            role=role,
            password_hash=hash_password(INITIAL_PASSWORD),
            must_change_password=False,
            is_active=True,
        )
        db_session.add(user)
        db_session.commit()
        return AuthorizedClient(
            client,
            {"Authorization": f"Bearer {issue_access_token(user, settings)}"},
        )

    return create


@pytest.fixture
def admin_client(authorized_client_factory):
    return authorized_client_factory("admin")


@pytest.fixture
def project_lead_client(authorized_client_factory):
    return authorized_client_factory("project_lead")


def assert_error(response, status: int, code: str) -> None:
    assert response.status_code == status, response.get_data(as_text=True)
    assert response.json["data"] is None
    assert response.json["error"]["code"] == code


def test_only_admin_can_create_template(admin_client, project_lead_client):
    url = "/api/v1/document-templates/prediagnosis/versions"
    docx = _build_docx_bytes()

    response = project_lead_client.post(
        url,
        data={"file": (BytesIO(docx), "prediagnosis.docx")},
        content_type="multipart/form-data",
    )
    assert_error(response, 403, "forbidden")

    admin_response = admin_client.post(
        url,
        data={"file": (BytesIO(docx), "prediagnosis.docx")},
        content_type="multipart/form-data",
    )
    assert admin_response.status_code == 201, admin_response.get_data(as_text=True)
    assert admin_response.json["data"]["version_number"] == 1
    assert admin_response.json["data"]["status"] == "draft"


def test_admin_can_download_uploaded_template(admin_client, project_lead_client):
    docx = _build_docx_bytes(text="模板下载测试")
    created = admin_client.post(
        "/api/v1/document-templates/sow/versions",
        data={"file": (BytesIO(docx), "sow.docx")},
        content_type="multipart/form-data",
    )
    assert created.status_code == 201, created.get_data(as_text=True)
    version_id = created.json["data"]["id"]

    forbidden = project_lead_client.get(
        f"/api/v1/document-templates/{version_id}/download-url"
    )
    assert_error(forbidden, 403, "forbidden")

    response = admin_client.get(
        f"/api/v1/document-templates/{version_id}/download-url"
    )
    assert response.status_code == 200, response.get_data(as_text=True)
    assert response.json["data"]["url"]
    assert response.json["data"]["expires_seconds"] == 300


def test_create_version_with_unsafe_placeholder_is_rejected(admin_client):
    docx = _build_docx_bytes(
        text="{{ cycler.__init__.__globals__.os.system('id') }}"
    )
    response = admin_client.post(
        "/api/v1/document-templates/sow/versions",
        data={"file": (BytesIO(docx), "sow.docx")},
        content_type="multipart/form-data",
    )
    assert_error(response, 400, "unsafe_template_expression")


def test_publish_freezes_version_and_edit_creates_next(admin_client):
    docx = _build_docx_bytes()
    create = admin_client.post(
        "/api/v1/document-templates/sow/versions",
        data={"file": (BytesIO(docx), "sow.docx")},
        content_type="multipart/form-data",
    )
    assert create.status_code == 201, create.get_data(as_text=True)
    version_id = create.json["data"]["id"]
    template_id = create.json["data"]["template_id"]
    assert create.json["data"]["version_number"] == 1
    assert create.json["data"]["status"] == "draft"

    publish = admin_client.post(f"/api/v1/document-templates/{version_id}/publish")
    assert publish.status_code == 200, publish.get_data(as_text=True)
    assert publish.json["data"]["status"] == "published"
    assert publish.json["data"]["version_number"] == 1

    copy = admin_client.post(f"/api/v1/document-templates/{template_id}/versions")
    assert copy.status_code == 201, copy.get_data(as_text=True)
    assert copy.json["data"]["version_number"] == 2
    assert copy.json["data"]["status"] == "draft"
    # A published version can no longer be published again.
    republish = admin_client.post(f"/api/v1/document-templates/{version_id}/publish")
    assert_error(republish, 409, "invalid_version_state")


def test_seed_document_templates_is_idempotent(db_session):
    first = seed_document_templates(db_session)
    assert len(first["created"]) == 17
    assert first["missing"] == []

    first_template_count = int(
        db_session.scalar(select(func.count()).select_from(DocumentTemplate))
    )
    first_version_count = int(
        db_session.scalar(
            select(func.count()).select_from(DocumentTemplateVersion)
        )
    )
    assert first_template_count == 17
    assert first_version_count == 17

    second = seed_document_templates(db_session)
    assert second["created"] == []
    assert len(second["skipped"]) == 17
    assert second["missing"] == []

    second_template_count = int(
        db_session.scalar(select(func.count()).select_from(DocumentTemplate))
    )
    second_version_count = int(
        db_session.scalar(
            select(func.count()).select_from(DocumentTemplateVersion)
        )
    )
    assert first_template_count == second_template_count
    assert first_version_count == second_version_count
