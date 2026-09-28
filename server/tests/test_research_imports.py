from __future__ import annotations

import csv
from concurrent.futures import ThreadPoolExecutor
from dataclasses import dataclass
from io import BytesIO, StringIO
import json
import re
import subprocess
import sys
import threading
import time
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from freezegun import freeze_time
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.worksheet.table import Table, TableStyleInfo
from sqlalchemy import event, func, select
from sqlalchemy.exc import SQLAlchemyError

from fde_api.auth.models import User
from fde_api.auth.passwords import hash_password
from fde_api.auth.tokens import issue_access_token
from fde_api.extensions import db
from fde_api.research import subject_service
from fde_api.research import imports as imports_service
from fde_api.research import project_routes
from fde_api.research.imports import (
    MAX_IMPORT_BYTES,
    MAX_IMPORT_ROWS,
    ImportServiceError,
    commit_subject_import,
    preview_subject_import,
    store_subject_import_preview,
)
from fde_api.research.models import ProjectResearchSubject
from fde_api.workbench.models import IndustryTemplateVersion, OperationEvent, ProjectMember
from fde_api.workbench.seed import seed_workbench


@dataclass(frozen=True)
class AuthorizedClient:
    client: object
    user: User
    headers: dict[str, str]

    def get(self, *args, **kwargs):
        return self.client.get(*args, headers=self.headers, **kwargs)

    def post(self, *args, **kwargs):
        return self.client.post(*args, headers=self.headers, **kwargs)


@pytest.fixture
def authorized_client_factory(client, db_session, settings):
    def create(role: str) -> AuthorizedClient:
        user = User(
            username=f"research.import.{role}.{uuid4().hex}",
            display_name=role,
            role=role,
            password_hash=hash_password("InitialPass!234"),
            must_change_password=False,
            is_active=True,
        )
        db_session.add(user)
        db_session.commit()
        return AuthorizedClient(
            client,
            user,
            {"Authorization": f"Bearer {issue_access_token(user, settings)}"},
        )

    return create


@pytest.fixture
def admin_client(authorized_client_factory):
    return authorized_client_factory("admin")


@pytest.fixture
def leader_client(authorized_client_factory):
    return authorized_client_factory("project_lead")


@pytest.fixture
def engineer_client(authorized_client_factory):
    return authorized_client_factory("fde_engineer")


@pytest.fixture
def viewer_client(authorized_client_factory):
    return authorized_client_factory("viewer")


@pytest.fixture
def published_template(db_session):
    seed_workbench(db_session)
    db_session.commit()
    return db_session.scalar(
        select(IndustryTemplateVersion).where(
            IndustryTemplateVersion.status == "published"
        )
    )


@pytest.fixture(autouse=True)
def clear_import_previews(app):
    with app.app_context():
        redis = app.extensions["fde_api_redis"]
        keys = list(redis.scan_iter(match="fde:research-import:*"))
        if keys:
            redis.delete(*keys)
    yield
    with app.app_context():
        redis = app.extensions["fde_api_redis"]
        keys = list(redis.scan_iter(match="fde:research-import:*"))
        if keys:
            redis.delete(*keys)


def _project_payload(template, leader, *, name="Research import project"):
    return {
        "name": name,
        "enterprise_name": "星河制造",
        "template_version_id": template.id,
        "leader_user_id": leader.id,
        "planned_start_date": "2026-08-22",
        "module_keys": ["pre_diagnosis"],
    }


def _create_project(client, template, leader, *, name="Research import project"):
    response = client.post(
        "/api/v1/projects", json=_project_payload(template, leader, name=name)
    )
    assert response.status_code == 201, response.get_data(as_text=True)
    return response.json["data"]


def _csv_bytes(rows, *, encoding="utf-8-sig"):
    output = StringIO(newline="")
    writer = csv.writer(output, lineterminator="\n")
    writer.writerows(rows)
    return output.getvalue().encode(encoding)


def _xlsx_bytes(rows):
    workbook = Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _xlsx_with_table_and_comment():
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["名称", "说明"])
    sheet.append(["销售部", "客户增长"])
    table = Table(displayName="SubjectsTable", ref="A1:B2")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    sheet["B2"].comment = Comment("普通备注", "测试用户")
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _macro_disguised_as_xlsx(rows):
    source = BytesIO(_xlsx_bytes(rows))
    output = BytesIO()
    with ZipFile(source) as original, ZipFile(output, "w", ZIP_DEFLATED) as disguised:
        for item in original.infolist():
            disguised.writestr(item, original.read(item.filename))
        disguised.writestr("xl/vbaProject.bin", b"macro-body")
    return output.getvalue()


def _rewrite_xlsx(contents, *, replacements=None, additions=None):
    replacements = replacements or {}
    additions = additions or {}
    output = BytesIO()
    with ZipFile(BytesIO(contents)) as original, ZipFile(
        output, "w", ZIP_DEFLATED
    ) as rewritten:
        for item in original.infolist():
            if item.filename in replacements:
                rewritten.writestr(item.filename, replacements[item.filename])
            else:
                rewritten.writestr(item, original.read(item.filename))
        for name, value in additions.items():
            rewritten.writestr(name, value)
    return output.getvalue()


def _rewrite_comment_vml(contents, transform):
    with ZipFile(BytesIO(contents)) as archive:
        vml_name = next(name for name in archive.namelist() if name.endswith(".vml"))
        original_vml = archive.read(vml_name)
    rewritten_vml = transform(original_vml)
    assert rewritten_vml != original_vml
    return _rewrite_xlsx(contents, replacements={vml_name: rewritten_vml})


_VML_NS_URI = b"urn:schemas-microsoft-com:vml"
_OFFICE_NS_URI = b"urn:schemas-microsoft-com:office:office"
_EXCEL_NS_URI = b"urn:schemas-microsoft-com:office:excel"
_PFX_RE = rb"([A-Za-z_][A-Za-z0-9_.-]*)"


def _comment_vml_shape_context(contents):
    """Return (vml, office, excel) prefixes actually used in the comment VML.

    openpyxl uses lxml when it is installed (as the document engine requires),
    which binds namespaces per element: the ``<shape>`` element binds its own
    ``ns0`` to the VML namespace and ``ns1`` to the office namespace, while the
    ``<ClientData>`` child binds ``ns2`` to the excel namespace. The security
    tests smuggle active content into the shape, so resolve the prefixes from
    the tags themselves rather than assuming a stable ``ns1``/``ns2`` scheme.
    """
    with ZipFile(BytesIO(contents)) as archive:
        vml_name = next(name for name in archive.namelist() if name.endswith(".vml"))
        vml = archive.read(vml_name)
    shape_match = re.search(rb"<" + _PFX_RE + rb":shape[\s>/]", vml)
    if shape_match is None:
        raise AssertionError("comment VML has no shape element")
    p_vml = shape_match.group(1)
    shape_open = vml[shape_match.start() : shape_match.start() + 2048]
    office_match = re.search(
        rb"xmlns:" + _PFX_RE + rb'\s*=\s*"' + re.escape(_OFFICE_NS_URI) + rb'"',
        shape_open,
    )
    p_office = office_match.group(1) if office_match else b"ns"
    client_match = re.search(rb"<" + _PFX_RE + rb":ClientData[\s>/]", vml)
    p_excel = client_match.group(1) if client_match else b"ns"
    return p_vml, p_office, p_excel


def _xlsx_with_active_part(kind):
    contents = _xlsx_bytes([["名称"], ["销售部"]])
    with ZipFile(BytesIO(contents)) as archive:
        content_types = archive.read("[Content_Types].xml").decode()
        relationships = archive.read("xl/_rels/workbook.xml.rels").decode()

    if kind == "xlm":
        part = "xl/macrosheets/sheet1.xml"
        content_type = "application/vnd.ms-excel.macrosheet+xml"
        relationship_type = (
            "http://schemas.microsoft.com/office/2006/relationships/xlMacrosheet"
        )
        target = "/xl/macrosheets/sheet1.xml"
    elif kind == "external":
        part = "xl/externalLinks/externalLink1.xml"
        content_type = (
            "application/vnd.openxmlformats-officedocument.spreadsheetml."
            "externalLink+xml"
        )
        relationship_type = (
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
            "externalLink"
        )
        target = "https://example.invalid/workbook.xlsx"
    else:
        part = "xl/embeddings/oleObject1.bin"
        content_type = "application/vnd.openxmlformats-officedocument.oleObject"
        relationship_type = (
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
            "oleObject"
        )
        target = "/xl/embeddings/oleObject1.bin"

    content_types = content_types.replace(
        "</Types>",
        f'<Override PartName="/{part}" ContentType="{content_type}" /></Types>',
    )
    external_mode = ' TargetMode="External"' if kind == "external" else ""
    relationships = relationships.replace(
        "</Relationships>",
        f'<Relationship Type="{relationship_type}" Target="{target}" '
        f'Id="rIdActive"{external_mode} /></Relationships>',
    )
    return _rewrite_xlsx(
        contents,
        replacements={
            "[Content_Types].xml": content_types,
            "xl/_rels/workbook.xml.rels": relationships,
        },
        additions={part: b"active-content"},
    )


def _xlsx_with_dimension(rows, dimension):
    contents = _xlsx_bytes(rows)
    with ZipFile(BytesIO(contents)) as archive:
        sheet = archive.read("xl/worksheets/sheet1.xml").decode()
    start = sheet.index("<dimension")
    end = sheet.index("/>", start) + 2
    sheet = sheet[:start] + f'<dimension ref="{dimension}"/>' + sheet[end:]
    return _rewrite_xlsx(
        contents, replacements={"xl/worksheets/sheet1.xml": sheet}
    )


def _preview(client, project_id, rows, subject_type="department", filename="subjects.csv"):
    contents = _xlsx_bytes(rows) if filename.endswith(".xlsx") else _csv_bytes(rows)
    return client.post(
        f"/api/v1/projects/{project_id}/research/subjects/imports/preview",
        data={
            "subject_type": subject_type,
            "file": (BytesIO(contents), filename),
        },
        content_type="multipart/form-data",
    )


def _commit(client, project_id, token, version):
    return client.post(
        f"/api/v1/projects/{project_id}/research/subjects/imports/commit",
        json={"preview_token": token, "version": version},
    )


def _assert_error(response, status, code):
    assert response.status_code == status
    assert response.json["data"] is None
    assert response.json["error"]["code"] == code


def test_csv_preview_maps_chinese_columns_and_decodes_utf8_bom_and_gb18030():
    """Breaking either supported CSV encoding would reject files exported by common Chinese tools."""
    expected = [
        {"row": 2, "name": "销售部", "description": "负责销售", "parent_name": None},
    ]

    utf8 = preview_subject_import(
        BytesIO(_csv_bytes([["名称", "说明", "父级"], ["销售部", "负责销售", ""]])),
        "departments.csv",
        "department",
    )
    gb18030 = preview_subject_import(
        BytesIO(
            _csv_bytes(
                [["名称", "说明", "父级"], ["销售部", "负责销售", ""]],
                encoding="gb18030",
            )
        ),
        "departments.csv",
        "department",
    )

    assert [row.to_dict() for row in utf8.rows] == expected
    assert [row.to_dict() for row in gb18030.rows] == expected
    assert utf8.errors == ()
    assert gb18030.errors == ()


def test_xlsx_preview_reports_exact_row_and_column_and_uses_data_only_values():
    """Losing worksheet coordinates would make import correction impractical."""
    preview = preview_subject_import(
        BytesIO(_xlsx_bytes([["名称", "父级"], ["销售", "不存在"]])),
        "departments.xlsx",
        "department",
    )

    assert preview.errors[0].row == 2
    assert preview.errors[0].column == "父级"
    assert preview.errors[0].code == "invalid_parent"

    formula_preview = preview_subject_import(
        BytesIO(_xlsx_bytes([["名称", "说明"], ["=1+1", "formula cell"]])),
        "departments.xlsx",
        "department",
    )
    assert formula_preview.errors[0].row == 2
    assert formula_preview.errors[0].column == "名称"


@pytest.mark.parametrize("filename", ["subjects.xls", "subjects.xlsm", "subjects.ods", "subjects.txt"])
def test_preview_rejects_macros_and_unsupported_formats(filename):
    """Accepting macro-enabled or legacy containers would violate the bounded parser contract."""
    with pytest.raises(ImportServiceError) as raised:
        preview_subject_import(BytesIO(b"not a supported workbook"), filename, "department")

    assert raised.value.code == "unsupported_import_format"
    assert raised.value.status == 415


def test_preview_rejects_macro_payload_renamed_to_xlsx_and_invalid_csv_encoding():
    """Checking only filename suffixes would allow renamed macro workbooks through."""
    with pytest.raises(ImportServiceError) as macro:
        preview_subject_import(
            BytesIO(_macro_disguised_as_xlsx([["名称"], ["销售部"]])),
            "renamed.xlsx",
            "department",
        )
    with pytest.raises(ImportServiceError) as encoding:
        preview_subject_import(BytesIO(b"\xff"), "subjects.csv", "department")

    assert macro.value.code == "unsupported_import_format"
    assert encoding.value.code == "unsupported_csv_encoding"


def test_xlsx_container_rejects_small_compressed_large_xml_before_reading_parts():
    """A small ZIP that expands into a large XML member must fail before OOXML reads."""
    bomb = _rewrite_xlsx(
        _xlsx_bytes([["名称"], ["销售部"]]),
        additions={"xl/worksheets/bomb.xml": b"<x>" + b"A" * (2 * 1024 * 1024) + b"</x>"},
    )

    with pytest.raises(ImportServiceError) as raised:
        preview_subject_import(BytesIO(bomb), "bomb.xlsx", "department")

    assert raised.value.code == "research_import_workbook_too_complex"
    assert raised.value.status == 413


@pytest.mark.parametrize("kind", ["xlm", "external", "ole"])
def test_xlsx_allowlist_rejects_xlm_external_and_embedded_active_content(kind):
    """An OOXML active part or relationship must be rejected even without VBA strings."""
    with pytest.raises(ImportServiceError) as raised:
        preview_subject_import(
            BytesIO(_xlsx_with_active_part(kind)), "active.xlsx", "department"
        )

    assert raised.value.code == "unsupported_import_format"
    assert raised.value.status == 415


def test_xlsx_xml_entity_payload_is_rejected_with_defused_parser():
    """Workbook XML entity declarations must never be expanded by openpyxl."""
    contents = _xlsx_bytes([["名称"], ["销售部"]])
    with ZipFile(BytesIO(contents)) as archive:
        workbook_xml = archive.read("xl/workbook.xml").decode()
    entity_xml = (
        '<!DOCTYPE workbook [<!ENTITY x "EXPANDED">]>'
        + workbook_xml.replace("工作表1", "&x;")
    )

    with pytest.raises(ImportServiceError) as raised:
        preview_subject_import(
            BytesIO(
                _rewrite_xlsx(
                    contents, replacements={"xl/workbook.xml": entity_xml}
                )
            ),
            "entity.xlsx",
            "department",
        )

    assert raised.value.code == "invalid_research_import_file"
    assert imports_service.openpyxl.DEFUSEDXML is True


def test_xlsx_accepts_standard_table_and_comment_passive_parts():
    """Normal Excel table/comment metadata must not be mistaken for active content."""
    preview = preview_subject_import(
        BytesIO(_xlsx_with_table_and_comment()),
        "passive-parts.xlsx",
        "department",
    )

    assert preview.errors == ()
    assert [row.name for row in preview.rows] == ["销售部"]
    assert preview.rows[0].description == "客户增长"


def test_xlsx_rejects_control_disguised_as_comment_vml():
    """The narrow comment VML exception must not admit an active control shape."""
    contents = _xlsx_with_table_and_comment()
    with ZipFile(BytesIO(contents)) as archive:
        vml_name = next(name for name in archive.namelist() if name.endswith(".vml"))
        active_vml = archive.read(vml_name).replace(
            b'ObjectType="Note"', b'ObjectType="Button"'
        )
    disguised = _rewrite_xlsx(
        contents,
        replacements={vml_name: active_vml},
    )

    with pytest.raises(ImportServiceError) as rejected:
        preview_subject_import(
            BytesIO(disguised), "active-control.xlsx", "department"
        )

    assert rejected.value.code == "unsupported_import_format"
    assert rejected.value.status == 415


@pytest.mark.parametrize(
    "scenario_id",
    [
        "javascript-href",
        "external-source",
        "relationship-id",
        "ole-object",
        "control-element",
        "script-element",
    ],
)
def test_xlsx_rejects_active_content_in_standard_comment_vml(scenario_id):
    """A legal Note shape must not smuggle active attributes or child elements."""
    contents = _xlsx_with_table_and_comment()
    p_vml, p_office, _ = _comment_vml_shape_context(contents)

    shape_open = b"<" + p_vml + b":shape "
    shape_close = b"</" + p_vml + b":shape>"
    injections = {
        "javascript-href": (
            shape_open,
            shape_open + b'href="javascript:alert(1)" ',
        ),
        "external-source": (
            shape_open,
            shape_open + b'src="https://example.invalid/payload" ',
        ),
        "relationship-id": (
            shape_open,
            shape_open + p_office + b':relid="rIdExternal" ',
        ),
        "ole-object": (
            shape_close,
            b'<x__ole:OLEObject xmlns:x__ole="urn:schemas-microsoft-com:office:office" '
            b'Type="Embed" ProgID="Package" />'
            + shape_close,
        ),
        "control-element": (
            shape_close,
            b'<x__ctl:Control xmlns:x__ctl="urn:schemas-microsoft-com:office:excel" '
            b'ObjectType="Button" />'
            + shape_close,
        ),
        "script-element": (
            shape_close,
            b'<script xmlns="http://www.w3.org/1999/xhtml">alert(1)</script>' + shape_close,
        ),
    }
    needle, injection = injections[scenario_id]

    def inject(vml, needle=needle, injection=injection):
        assert needle in vml
        return vml.replace(needle, injection, 1)

    disguised = _rewrite_comment_vml(contents, inject)

    with pytest.raises(ImportServiceError) as rejected:
        preview_subject_import(
            BytesIO(disguised), "active-comment-vml.xlsx", "department"
        )

    assert rejected.value.code == "unsupported_import_format"
    assert rejected.value.status == 415


def test_xlsx_requires_one_note_client_data_per_comment_shape():
    """Global Note counts must not let one shape borrow another shape's ClientData."""
    contents = _xlsx_with_table_and_comment()
    p_vml, _, p_excel = _comment_vml_shape_context(contents)
    shape_open = b"<" + p_vml + b":shape "
    shape_close = b"</" + p_vml + b":shape>"
    client_open = b"<" + p_excel + b":ClientData "
    client_close = b"</" + p_excel + b":ClientData>"

    def unbalance_client_data(vml):
        shape_start = vml.index(shape_open)
        shape_end = vml.index(shape_close, shape_start) + len(shape_close)
        shape = vml[shape_start:shape_end]
        client_start = shape.index(client_open)
        client_end = shape.index(client_close, client_start) + len(client_close)
        client_data = shape[client_start:client_end]
        two_notes = shape.replace(shape_close, client_data + shape_close, 1)
        no_note = shape.replace(client_data, b"", 1)
        return vml[:shape_start] + two_notes + no_note + vml[shape_end:]

    disguised = _rewrite_comment_vml(contents, unbalance_client_data)

    with pytest.raises(ImportServiceError) as rejected:
        preview_subject_import(
            BytesIO(disguised), "unbalanced-comment-vml.xlsx", "department"
        )

    assert rejected.value.code == "unsupported_import_format"
    assert rejected.value.status == 415


def test_preview_rejects_unknown_or_missing_headers_with_header_coordinates():
    """Silently ignoring columns would make a user believe data was imported when it was dropped."""
    unknown = preview_subject_import(
        BytesIO(_csv_bytes([["名称", "秘密列"], ["销售部", "x"]])),
        "departments.csv",
        "department",
    )
    missing = preview_subject_import(
        BytesIO(_csv_bytes([["说明"], ["没有名称"]])),
        "departments.csv",
        "department",
    )

    assert [error.to_dict() for error in unknown.errors] == [
        {
            "row": 1,
            "column": "秘密列",
            "code": "unknown_column",
            "message": "The import column is not supported.",
        }
    ]
    assert missing.errors[0].row == 1
    assert missing.errors[0].column == "名称"
    assert missing.errors[0].code == "missing_required_column"


def test_xlsx_scans_unknown_headers_after_column_d_and_data_under_empty_headers():
    """A fixed max_col would silently discard used columns after D."""
    unknown = preview_subject_import(
        BytesIO(
            _xlsx_bytes(
                [["名称", "说明", "父级", None, "秘密列"], ["销售部", "", "", None, "x"]]
            )
        ),
        "unknown.xlsx",
        "department",
    )
    missing_header = preview_subject_import(
        BytesIO(
            _xlsx_bytes(
                [["名称", "说明", "父级", None, None], ["销售部", "", "", None, "hidden"]]
            )
        ),
        "missing-header.xlsx",
        "department",
    )

    assert unknown.errors[0].row == 1
    assert unknown.errors[0].column == "秘密列"
    assert unknown.errors[0].code == "unknown_column"
    assert missing_header.errors[0].row == 2
    assert missing_header.errors[0].column == "E"
    assert missing_header.errors[0].code == "missing_header_for_data"


def test_xlsx_reports_far_sparse_unknown_header_without_dense_column_expansion():
    """Even XFD header errors stay exact; actual-cell count, not dense width, bounds work."""
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "名称"
    sheet["XFD1"] = "远端列"
    sheet["A2"] = "销售部"
    output = BytesIO()
    workbook.save(output)
    workbook.close()

    preview = preview_subject_import(
        BytesIO(output.getvalue()), "far-column.xlsx", "department"
    )

    assert preview.errors[0].code == "unknown_column"
    assert preview.errors[0].row == 1
    assert preview.errors[0].column == "远端列"


def test_xlsx_ignores_forged_dimension_without_iterating_fake_used_range(monkeypatch):
    """A forged XFD dimension must not make parsing iterate every implied row or column."""
    contents = _xlsx_with_dimension([["名称"], ["销售部"]], "A1:XFD500")
    real_bounded_rows = imports_service._bounded_rows
    observed_rows = 0

    def bounded_rows(rows):
        nonlocal observed_rows
        materialized = list(rows)
        observed_rows = len(materialized)
        return real_bounded_rows(materialized)

    monkeypatch.setattr(imports_service, "_bounded_rows", bounded_rows)
    preview = preview_subject_import(
        BytesIO(contents), "forged-dimension.xlsx", "department"
    )

    assert preview.errors == ()
    assert len(preview.rows) == 1
    assert observed_rows <= 3


def test_preview_enforces_exact_file_and_row_boundaries():
    """Off-by-one checks must accept the documented maxima and reject the next byte or row."""
    valid_prefix = _csv_bytes([["名称"], ["销售部"]], encoding="utf-8")
    exact_size = valid_prefix + b" " * (MAX_IMPORT_BYTES - len(valid_prefix))
    accepted = preview_subject_import(
        BytesIO(exact_size), "departments.csv", "department"
    )
    assert len(accepted.rows) == 1

    with pytest.raises(ImportServiceError) as too_large:
        preview_subject_import(
            BytesIO(exact_size + b"\n"), "departments.csv", "department"
        )
    assert too_large.value.code == "research_import_too_large"
    assert too_large.value.status == 413

    rows_at_limit = [["名称"]] + [[f"部门{i}"] for i in range(MAX_IMPORT_ROWS)]
    accepted_rows = preview_subject_import(
        BytesIO(_csv_bytes(rows_at_limit)), "departments.csv", "department"
    )
    assert len(accepted_rows.rows) == MAX_IMPORT_ROWS

    with pytest.raises(ImportServiceError) as too_many:
        preview_subject_import(
            BytesIO(_csv_bytes(rows_at_limit + [["超出边界"]])),
            "departments.csv",
            "department",
        )
    assert too_many.value.code == "research_import_too_many_rows"

    coordinate_preview = preview_subject_import(
        BytesIO(_csv_bytes([["名称", "父级"], ["", ""], ["销售部", "不存在"]])),
        "departments.csv",
        "department",
    )
    assert coordinate_preview.errors[0].row == 3


def test_csv_field_limit_change_is_serialized_and_restored_across_threads(monkeypatch):
    """Process-global csv.field_size_limit must never overlap between parser threads."""
    real_field_size_limit = imports_service.csv.field_size_limit
    guard = threading.Lock()
    active_changes = 0
    overlap = False

    def observed_field_size_limit(limit=None):
        nonlocal active_changes, overlap
        if limit is None:
            return real_field_size_limit()
        if limit == MAX_IMPORT_BYTES:
            with guard:
                active_changes += 1
                overlap = overlap or active_changes > 1
            time.sleep(0.05)
            result = real_field_size_limit(limit)
            with guard:
                active_changes -= 1
            return result
        return real_field_size_limit(limit)

    monkeypatch.setattr(
        imports_service.csv, "field_size_limit", observed_field_size_limit
    )
    payload = _csv_bytes([["名称"], ["销售部"]])
    with ThreadPoolExecutor(max_workers=2) as executor:
        previews = list(
            executor.map(
                lambda _: preview_subject_import(
                    BytesIO(payload), "departments.csv", "department"
                ),
                range(2),
            )
        )

    assert all(preview.errors == () for preview in previews)
    assert overlap is False
    assert real_field_size_limit() == 131072


def test_import_contract_loads_in_a_fresh_python_process():
    """The public import functions must not depend on Flask's module import order."""
    completed = subprocess.run(
        [
            sys.executable,
            "-c",
            (
                "from fde_api.research.imports import "
                "preview_subject_import, commit_subject_import"
            ),
        ],
        capture_output=True,
        text=True,
        check=False,
    )

    assert completed.returncode == 0, completed.stderr


def test_commit_import_is_atomic_and_reports_the_invalid_row(
    admin_client, published_template, leader_client, db_session
):
    """An invalid later row must roll back an earlier row inserted in the same transaction."""
    project = _create_project(admin_client, published_template, leader_client.user)
    preview = _preview(
        admin_client,
        project["id"],
        [["名称", "说明"], ["销售部", "valid"], ["", "invalid"]],
    )
    assert preview.status_code == 200

    response = _commit(
        admin_client,
        project["id"],
        preview.json["data"]["preview_token"],
        project["version"],
    )

    _assert_error(response, 422, "invalid_research_import")
    assert response.json["error"]["details"]["errors"][0]["row"] == 3
    assert response.json["error"]["details"]["errors"][0]["column"] == "名称"
    assert db_session.scalar(
        select(func.count()).select_from(ProjectResearchSubject).where(
            ProjectResearchSubject.project_id == project["id"],
            ProjectResearchSubject.name == "销售部",
        )
    ) == 0


def test_leading_blank_header_error_is_preserved_at_commit(
    admin_client, published_template, leader_client
):
    """Commit must use explicit preview header metadata, not assume the header is row 1."""
    project = _create_project(admin_client, published_template, leader_client.user)
    response = admin_client.post(
        f"/api/v1/projects/{project['id']}/research/subjects/imports/preview",
        data={
            "subject_type": "department",
            "file": (BytesIO(_csv_bytes([[""], [""], ["说明"], ["没有名称"]])), "leading.csv"),
        },
        content_type="multipart/form-data",
    )
    assert response.status_code == 200
    preview_error = response.json["data"]["errors"][0]
    assert preview_error["row"] == 3
    assert preview_error["column"] == "名称"

    committed = _commit(
        admin_client,
        project["id"],
        response.json["data"]["preview_token"],
        project["version"],
    )

    _assert_error(committed, 422, "invalid_research_import")
    assert committed.json["error"]["details"]["errors"] == [preview_error]


def test_csv_commit_creates_one_atomic_batch_and_duplicate_commit_is_idempotent(
    admin_client, published_template, leader_client, db_session
):
    """Replaying a preview token must return the first result without duplicating rows or versions."""
    project = _create_project(admin_client, published_template, leader_client.user)
    preview = _preview(
        admin_client,
        project["id"],
        [["名称", "说明"], ["销售部", "A"], ["产品部", "B"]],
    )
    token = preview.json["data"]["preview_token"]

    first = _commit(admin_client, project["id"], token, project["version"])
    replay = _commit(admin_client, project["id"], token, project["version"])

    assert first.status_code == 201
    assert replay.status_code == 200
    assert replay.json["data"]["idempotent_replay"] is True
    assert replay.json["data"]["subjects"] == first.json["data"]["subjects"]
    assert first.json["data"]["project_version"] == project["version"] + 1
    assert db_session.scalar(
        select(func.count()).select_from(ProjectResearchSubject).where(
            ProjectResearchSubject.project_id == project["id"],
            ProjectResearchSubject.subject_type == "department",
        )
    ) == 2


def test_xlsx_commit_resolves_existing_department_parent_and_preserves_row_order(
    admin_client, published_template, leader_client
):
    """Role imports must resolve a unique in-project department before inserting rows in file order."""
    project = _create_project(admin_client, published_template, leader_client.user)
    department = admin_client.post(
        f"/api/v1/projects/{project['id']}/research/subjects",
        json={
            "version": project["version"],
            "subject_type": "department",
            "subject_key": "sales_department",
            "name": "销售部",
            "description": "",
        },
    )
    assert department.status_code == 201
    preview = _preview(
        admin_client,
        project["id"],
        [["名称", "父级"], ["销售经理", "销售部"], ["销售专员", "销售部"]],
        subject_type="role",
        filename="roles.xlsx",
    )

    committed = _commit(
        admin_client,
        project["id"],
        preview.json["data"]["preview_token"],
        department.json["data"]["project_version"],
    )

    assert committed.status_code == 201
    subjects = committed.json["data"]["subjects"]
    assert [subject["name"] for subject in subjects] == ["销售经理", "销售专员"]
    assert {subject["parent_subject_id"] for subject in subjects} == {
        department.json["data"]["id"]
    }
    assert [subject["tracking_code"] for subject in subjects] == ["JOB-0001", "JOB-0002"]


def test_commit_revalidates_parent_and_rolls_back_all_rows(
    admin_client, published_template, leader_client, db_session
):
    """A foreign-project parent must look missing, identify its source cell, and write nothing."""
    project = _create_project(admin_client, published_template, leader_client.user)
    foreign_project = _create_project(
        admin_client, published_template, leader_client.user, name="Foreign parent project"
    )
    foreign_parent = admin_client.post(
        f"/api/v1/projects/{foreign_project['id']}/research/subjects",
        json={
            "version": foreign_project["version"],
            "subject_type": "department",
            "subject_key": "foreign_department",
            "name": "不存在",
            "description": "",
        },
    )
    assert foreign_parent.status_code == 201
    preview = _preview(
        admin_client,
        project["id"],
        [["名称", "父级"], ["无父级岗位", ""], ["销售经理", "不存在"]],
        subject_type="role",
    )

    response = _commit(
        admin_client,
        project["id"],
        preview.json["data"]["preview_token"],
        project["version"],
    )

    _assert_error(response, 422, "invalid_research_import")
    assert response.json["error"]["details"]["errors"][0] == {
        "row": 3,
        "column": "父级",
        "code": "parent_not_found",
        "message": "The parent research subject was not found.",
    }
    assert db_session.scalar(
        select(func.count()).select_from(ProjectResearchSubject).where(
            ProjectResearchSubject.project_id == project["id"],
            ProjectResearchSubject.subject_type == "role",
        )
    ) == 0


def test_preview_token_is_project_bound_expires_and_cannot_be_tampered(
    app, admin_client, published_template, leader_client
):
    """A bearer preview must not cross project scope, survive state expiry, or accept signature edits."""
    first = _create_project(admin_client, published_template, leader_client.user, name="First")
    second = _create_project(admin_client, published_template, leader_client.user, name="Second")
    preview = _preview(admin_client, first["id"], [["名称"], ["销售部"]])
    token = preview.json["data"]["preview_token"]

    cross_project = _commit(admin_client, second["id"], token, second["version"])
    tampered = _commit(admin_client, first["id"], token + "x", first["version"])
    _assert_error(cross_project, 404, "project_not_found")
    _assert_error(tampered, 404, "research_import_preview_not_found")

    with app.app_context():
        redis = app.extensions["fde_api_redis"]
        state_key = next(redis.scan_iter(match="fde:research-import:preview:*"))
        assert 1_700 <= redis.ttl(state_key) <= 1_800
        redis.delete(state_key)
    expired = _commit(admin_client, first["id"], token, first["version"])
    _assert_error(expired, 410, "research_import_preview_expired")


def test_signed_preview_token_expires_after_thirty_minutes(
    app, admin_client, published_template, leader_client
):
    """A Redis key that outlives clock skew must not bypass the signed 30-minute lifetime."""
    project = _create_project(admin_client, published_template, leader_client.user)
    parsed = preview_subject_import(
        BytesIO(_csv_bytes([["名称"], ["销售部"]])),
        "departments.csv",
        "department",
    )
    with app.app_context():
        with freeze_time("2026-08-22 08:00:00"):
            preview = store_subject_import_preview(
                actor=admin_client.user, project_id=project["id"], preview=parsed
            )
        with freeze_time("2026-08-22 08:30:01"):
            with pytest.raises(ImportServiceError) as expired:
                commit_subject_import(
                    project["id"],
                    preview.preview_token,
                    project["version"],
                    admin_client.user,
                )

    assert expired.value.status == 410
    assert expired.value.code == "research_import_preview_expired"


def test_stale_commit_does_not_consume_preview_and_can_retry_current_version(
    admin_client, published_template, leader_client
):
    """A version conflict must leave the preview available for an explicit refreshed retry."""
    project = _create_project(admin_client, published_template, leader_client.user)
    preview = _preview(admin_client, project["id"], [["名称"], ["销售部"]])
    token = preview.json["data"]["preview_token"]
    intervening = admin_client.post(
        f"/api/v1/projects/{project['id']}/research/subjects",
        json={
            "version": project["version"],
            "subject_type": "department",
            "subject_key": "intervening",
            "name": "中间更新",
            "description": "",
        },
    )

    stale = _commit(admin_client, project["id"], token, project["version"])
    retry = _commit(
        admin_client,
        project["id"],
        token,
        intervening.json["data"]["project_version"],
    )

    _assert_error(stale, 409, "stale_version")
    assert retry.status_code == 201


def test_import_permissions_match_subject_writes(
    admin_client,
    engineer_client,
    viewer_client,
    published_template,
    leader_client,
    db_session,
):
    """Import routes must not bypass project membership or viewer read-only access."""
    project = _create_project(admin_client, published_template, leader_client.user)
    db_session.add(
        ProjectMember(
            project_id=project["id"], user_id=viewer_client.user.id, role="viewer"
        )
    )
    db_session.commit()

    unassigned = _preview(
        engineer_client, project["id"], [["名称"], ["销售部"]]
    )
    viewer = _preview(viewer_client, project["id"], [["名称"], ["销售部"]])
    admin_preview = _preview(
        admin_client, project["id"], [["名称"], ["销售部"]]
    )
    token = admin_preview.json["data"]["preview_token"]
    unassigned_commit = _commit(
        engineer_client, project["id"], token, project["version"]
    )
    viewer_commit = _commit(viewer_client, project["id"], token, project["version"])

    _assert_error(unassigned, 404, "project_not_found")
    _assert_error(viewer, 403, "forbidden")
    _assert_error(unassigned_commit, 404, "project_not_found")
    _assert_error(viewer_commit, 403, "forbidden")


def test_unauthorized_preview_is_rejected_before_file_parsing(
    admin_client,
    engineer_client,
    published_template,
    leader_client,
    monkeypatch,
):
    """An unassigned caller must not make the server parse an otherwise valid upload."""
    project = _create_project(admin_client, published_template, leader_client.user)

    def parsing_would_be_a_bug(*_args, **_kwargs):
        raise AssertionError("unauthorized upload reached the parser")

    monkeypatch.setattr(
        project_routes, "preview_subject_import", parsing_would_be_a_bug
    )
    response = _preview(
        engineer_client, project["id"], [["名称"], ["销售部"]]
    )

    _assert_error(response, 404, "project_not_found")


def test_retry_recovers_committed_batch_when_redis_result_write_fails(
    app,
    admin_client,
    published_template,
    leader_client,
    db_session,
    monkeypatch,
):
    """A crash-window result-store failure must replay the committed deterministic batch."""
    project = _create_project(admin_client, published_template, leader_client.user)
    preview = _preview(
        admin_client,
        project["id"],
        [["名称"], ["销售部"], ["产品部"]],
    )
    token = preview.json["data"]["preview_token"]
    real_mark_commit_state = imports_service._mark_commit_state
    mark_calls = 0

    def fail_committed_result_store(lease, result):
        nonlocal mark_calls
        mark_calls += 1
        if mark_calls == 1:
            raise ImportServiceError(
                "research_import_storage_unavailable",
                "Subject import previews are temporarily unavailable.",
                503,
            )
        return real_mark_commit_state(lease, result)

    monkeypatch.setattr(
        imports_service, "_mark_commit_state", fail_committed_result_store
    )
    first = _commit(admin_client, project["id"], token, project["version"])
    monkeypatch.undo()
    retry = _commit(admin_client, project["id"], token, project["version"])

    _assert_error(first, 503, "research_import_storage_unavailable")
    assert retry.status_code == 200
    assert retry.json["data"]["idempotent_replay"] is True
    assert retry.json["data"]["project_version"] == project["version"] + 1
    assert db_session.scalar(
        select(func.count()).select_from(ProjectResearchSubject).where(
            ProjectResearchSubject.project_id == project["id"],
            ProjectResearchSubject.subject_type == "department",
        )
    ) == 2

    from fde_api.research.models import ProjectResearchImportBatch

    with app.app_context():
        preview_id = imports_service._token_claims(token)["preview_id"]
    receipt = db_session.get(ProjectResearchImportBatch, preview_id)
    assert receipt is not None
    assert receipt.committed_project_version == project["version"] + 1
    assert receipt.result_json["subjects"] == retry.json["data"]["subjects"]


def test_committed_receipt_replays_after_redis_preview_is_completely_lost(
    app, admin_client, published_template, leader_client
):
    """The atomic database receipt, not transient Redis state, owns committed replay."""
    project = _create_project(admin_client, published_template, leader_client.user)
    preview = _preview(
        admin_client,
        project["id"],
        [["名称"], ["销售部"], ["产品部"]],
    )
    token = preview.json["data"]["preview_token"]
    first = _commit(admin_client, project["id"], token, project["version"])
    assert first.status_code == 201

    with app.app_context():
        redis = app.extensions["fde_api_redis"]
        keys = list(redis.scan_iter(match="fde:research-import:*"))
        assert keys
        redis.delete(*keys)

    replay = _commit(admin_client, project["id"], token, project["version"])

    assert replay.status_code == 200
    assert replay.json["data"]["idempotent_replay"] is True
    assert replay.json["data"]["project_version"] == project["version"] + 1
    assert replay.json["data"]["subjects"] == first.json["data"]["subjects"]


def test_receipt_lookup_does_not_leak_across_project_or_permission_boundaries(
    app,
    admin_client,
    viewer_client,
    published_template,
    leader_client,
    db_session,
):
    """Project binding and write authorization must precede any receipt lookup."""
    project = _create_project(admin_client, published_template, leader_client.user)
    other = _create_project(
        admin_client,
        published_template,
        leader_client.user,
        name="Other receipt scope",
    )
    db_session.add(
        ProjectMember(
            project_id=project["id"], user_id=viewer_client.user.id, role="viewer"
        )
    )
    db_session.commit()
    preview = _preview(admin_client, project["id"], [["名称"], ["销售部"]])
    token = preview.json["data"]["preview_token"]
    committed = _commit(admin_client, project["id"], token, project["version"])
    assert committed.status_code == 201
    with app.app_context():
        redis = app.extensions["fde_api_redis"]
        keys = list(redis.scan_iter(match="fde:research-import:*"))
        redis.delete(*keys)

    statements: list[str] = []

    def capture(_connection, _cursor, statement, _parameters, _context, _many):
        statements.append(" ".join(statement.lower().split()))

    with app.app_context():
        event.listen(db.engine, "before_cursor_execute", capture)
        try:
            viewer = _commit(
                viewer_client, project["id"], token, project["version"]
            )
        finally:
            event.remove(db.engine, "before_cursor_execute", capture)
    wrong_project = _commit(admin_client, other["id"], token, other["version"])

    _assert_error(viewer, 403, "forbidden")
    _assert_error(wrong_project, 404, "project_not_found")
    assert not any("project_research_import_batches" in sql for sql in statements)


def test_predictable_subject_keys_cannot_forge_an_idempotent_import_receipt(
    app, admin_client, published_template, leader_client
):
    """Ordinary subject writes matching preview-derived keys must not bypass stale_version."""
    project = _create_project(admin_client, published_template, leader_client.user)
    preview = _preview(
        admin_client,
        project["id"],
        [["名称"], ["销售部"], ["产品部"]],
    )
    token = preview.json["data"]["preview_token"]
    with app.app_context():
        preview_id = imports_service._token_claims(token)["preview_id"]

    version = project["version"]
    for row_number, name in ((2, "销售部"), (3, "产品部")):
        created = admin_client.post(
            f"/api/v1/projects/{project['id']}/research/subjects",
            json={
                "version": version,
                "subject_type": "department",
                "subject_key": f"import_department_{preview_id}_{row_number}",
                "name": name,
                "description": "",
                "sort_order": row_number - 2,
            },
        )
        assert created.status_code == 201
        version = created.json["data"]["project_version"]

    forged = _commit(admin_client, project["id"], token, project["version"])

    _assert_error(forged, 409, "stale_version")


def test_lost_redis_commit_owner_cannot_reset_or_overwrite_newer_fence(
    app, admin_client, published_template, leader_client
):
    """Every Redis transition must compare the current owner/fence atomically."""
    project = _create_project(admin_client, published_template, leader_client.user)
    preview = _preview(admin_client, project["id"], [["名称"], ["销售部"]])
    token = preview.json["data"]["preview_token"]
    result = {
        "subjects": [],
        "imported_count": 0,
        "project_version": 2,
        "idempotent_replay": False,
    }

    with app.app_context():
        preview_id = imports_service._token_claims(token)["preview_id"]
        state_key = imports_service._preview_key(preview_id)
        first = imports_service._acquire_commit_lease(
            state_key, preview_id, 1, owner="first-owner"
        )
        app.extensions["fde_api_redis"].delete(first.lock_key)
        second = imports_service._acquire_commit_lease(
            state_key, preview_id, 1, owner="second-owner"
        )
        assert second.fence > first.fence
        assert imports_service._mark_commit_state(second, result) is True
        assert imports_service._mark_commit_state(first, {**result, "project_version": 99}) is False
        assert imports_service._reset_commit_state(first) is False
        state = imports_service._load_state(state_key)

    assert state["status"] == "committed"
    assert state["result"]["project_version"] == 2


def test_large_import_uses_bounded_tracking_and_form_definition_queries(
    app, admin_client, published_template, leader_client
):
    """Batch size must not multiply tracking scans or form-definition lookup queries."""
    project = _create_project(admin_client, published_template, leader_client.user)
    rows = [["名称"]] + [[f"机会{i}"] for i in range(120)]
    preview = _preview(
        admin_client, project["id"], rows, subject_type="opportunity"
    )
    statements: list[str] = []

    def capture(_connection, _cursor, statement, _parameters, _context, _many):
        statements.append(" ".join(statement.lower().split()))

    with app.app_context():
        event.listen(db.engine, "before_cursor_execute", capture)
        try:
            committed = _commit(
                admin_client,
                project["id"],
                preview.json["data"]["preview_token"],
                project["version"],
            )
        finally:
            event.remove(db.engine, "before_cursor_execute", capture)

    assert committed.status_code == 201
    assert committed.json["data"]["imported_count"] == 120
    tracking_scans = [
        statement
        for statement in statements
        if "project_research_subjects.tracking_code" in statement
    ]
    active_module_lookups = [
        statement
        for statement in statements
        if "module_catalog.module_key" in statement
    ]
    existing_form_lookups = [
        statement
        for statement in statements
        if "project_research_forms.form_key" in statement
    ]
    assert len(tracking_scans) <= 1
    assert len(active_module_lookups) <= 1
    assert existing_form_lookups == []


def test_event_failure_rolls_back_batch_and_events_exclude_file_content(
    admin_client,
    published_template,
    leader_client,
    db_session,
    monkeypatch,
):
    """Audit storage must fail atomically and successful events must omit names/descriptions/file bytes."""
    project = _create_project(admin_client, published_template, leader_client.user)
    failed_preview = _preview(
        admin_client,
        project["id"],
        [["名称", "说明"], ["销售部", "RAW-CONTENT-ONE"], ["产品部", "RAW-CONTENT-TWO"]],
    )
    real_record_event = subject_service.record_event
    event_calls = 0

    def fail_second_event(*args, **kwargs):
        nonlocal event_calls
        event_calls += 1
        if event_calls == 2:
            raise SQLAlchemyError("event storage unavailable")
        return real_record_event(*args, **kwargs)

    monkeypatch.setattr(subject_service, "record_event", fail_second_event)
    failed = _commit(
        admin_client,
        project["id"],
        failed_preview.json["data"]["preview_token"],
        project["version"],
    )
    monkeypatch.undo()

    _assert_error(failed, 503, "research_import_failed")
    assert db_session.scalar(
        select(func.count()).select_from(ProjectResearchSubject).where(
            ProjectResearchSubject.project_id == project["id"],
            ProjectResearchSubject.subject_type == "department",
        )
    ) == 0
    db_session.rollback()

    succeeded = _commit(
        admin_client,
        project["id"],
        failed_preview.json["data"]["preview_token"],
        project["version"],
    )
    assert succeeded.status_code == 201
    events = list(
        db_session.scalars(
            select(OperationEvent).where(
                OperationEvent.project_id == project["id"],
                OperationEvent.event_type == "project_research_subject_created",
            )
        )
    )
    serialized_changes = json.dumps(
        [event.changes for event in events], ensure_ascii=False
    )
    assert "销售部" not in serialized_changes
    assert "产品部" not in serialized_changes
    assert "RAW-CONTENT" not in serialized_changes
