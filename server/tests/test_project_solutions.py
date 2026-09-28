import io
import json
from pathlib import Path
from unittest.mock import Mock, patch

from docx import Document
from sqlalchemy import func, select

from test_document_drafts_api import (
    INITIAL_PASSWORD, authorized_client_factory, leader_client, viewer_client, outsider_client,
    _add_viewer_member, _persist_doc_template, _persist_project,
)
from fde_api.documents.generator import _build_context, render_docx
from fde_api.documents.catalog import RESOURCE_TEMPLATES_DIR
from fde_api.documents.models import DocumentGenerationJob, ProjectDocument, ProjectDocumentVersion
from fde_api.research.models import ProjectAIOpportunityProfile, ProjectResearchAnswer, ProjectResearchForm, ProjectResearchFormRevision, ProjectResearchSubject
from fde_api.solutions.models import ProjectSolution, ProjectSolutionExport
from fde_api.solutions.legacy_migration import migrate_legacy_opportunity_delivery
from fde_api.workbench.models import OperationEvent


def _opportunity(db_session, project, key="one"):
    item = ProjectResearchSubject(project_id=project.id, subject_type="opportunity", subject_key=key,
        name=f"机会-{key}", description=f"需求-{key}", tracking_code=f"OPP-{key}", status="active")
    item.opportunity_profile = ProjectAIOpportunityProfile(opportunity_status="ready")
    db_session.add(item)
    db_session.commit()
    return item


def _create(client, project, opportunities, **overrides):
    return client.post(f"/api/v1/projects/{project.id}/solutions", json={
        "name": "统一方案", "opportunity_ids": [item.id for item in opportunities],
        "design_markdown": "## 业务流程\n接收订单 → 审核\n- **人工确认** 后提交。",
        "deliverables": "交付应用和使用说明", "acceptance_criteria": "与客户约定后验收",
        "schedule": "待确认", "data_systems": "客户 ERP", "risks_dependencies": "接口授权待确认", **overrides})


def _form(db_session, project, subject, answers):
    form = ProjectResearchForm(project_id=project.id, subject_id=subject.id, form_key="evidence", name="识别依据")
    db_session.add(form)
    db_session.flush()
    revision = ProjectResearchFormRevision(form_id=form.id, revision_number=1, definition_snapshot={"sections": []})
    db_session.add(revision)
    db_session.flush()
    for key, value in answers.items():
        db_session.add(ProjectResearchAnswer(revision_id=revision.id, field_key=key, value_json=value))
    form.current_revision_id = revision.id
    db_session.commit()
    return form, revision


def test_solution_membership_permissions_and_validation(db_session, leader_client, viewer_client, outsider_client):
    project = _persist_project(db_session, leader_client.user, code="SOLUTION-ACCESS")
    other = _persist_project(db_session, leader_client.user, code="SOLUTION-OTHER")
    opportunity = _opportunity(db_session, project)
    foreign = _opportunity(db_session, other)
    _add_viewer_member(db_session, project, viewer_client.user)
    assert _create(viewer_client, project, [opportunity]).status_code == 403
    assert _create(outsider_client, project, [opportunity]).status_code == 403
    assert _create(leader_client, project, [foreign]).status_code == 400
    assert _create(leader_client, project, []).status_code == 400
    assert _create(leader_client, project, [opportunity, opportunity]).status_code == 400
    created = _create(leader_client, project, [opportunity])
    assert created.status_code == 201, created.json
    item = created.json["data"]
    assert item["opportunity_ids"] == [opportunity.id]
    assert item["opportunities"][0]["tracking_code"] == "OPP-one"
    url = f"/api/v1/projects/{project.id}/solutions"
    assert viewer_client.get(url).json["data"]["items"][0]["id"] == item["id"]
    assert outsider_client.get(url).status_code == 403
    assert leader_client.get(f"/api/v1/projects/{other.id}/solutions/{item['id']}").status_code == 404
    db_session.refresh(opportunity)
    assert opportunity.opportunity_profile.opportunity_status == "ready"


def test_solution_update_stale_and_password_archive_preserve_opportunity(db_session, leader_client):
    project = _persist_project(db_session, leader_client.user, code="SOLUTION-LOCK")
    opportunity = _opportunity(db_session, project)
    data = _create(leader_client, project, [opportunity]).json["data"]
    url = f"/api/v1/projects/{project.id}/solutions/{data['id']}"
    saved = leader_client.patch(url, json={"version": 1, "design_markdown": "新内容"})
    assert saved.status_code == 200, saved.json
    assert saved.json["data"]["version"] == 2
    assert leader_client.patch(url, json={"version": 1, "name": "过期"}).status_code == 409
    assert leader_client.post(url + "/archive", json={"version": 2}).status_code == 400
    assert leader_client.post(url + "/archive", json={"version": 2, "password": "wrong"}).status_code == 403
    archived = leader_client.post(url + "/archive", json={"version": 2, "password": INITIAL_PASSWORD})
    assert archived.status_code == 200
    assert archived.json["data"]["status"] == "archived"
    assert leader_client.patch(url, json={"version": 3, "name": "不能修改"}).status_code == 409
    assert leader_client.post(url + "/export-sow", json={"version": 3}).status_code == 409
    db_session.refresh(opportunity)
    assert opportunity.status == "active"
    events = db_session.scalars(select(OperationEvent).where(OperationEvent.project_id == project.id)).all()
    assert "password" not in str([event.changes for event in events])
    assert INITIAL_PASSWORD not in str([event.changes for event in events])


def test_ai_organizes_selected_scope_without_mutation(app, db_session, leader_client, viewer_client):
    project = _persist_project(db_session, leader_client.user, code="SOLUTION-AI")
    opportunity = _opportunity(db_session, project)
    _opportunity(db_session, project, "unselected")
    form, revision = _form(db_session, project, opportunity, {"current_state": "订单逐张手工录入", "technical_prereqs": "ERP 尚未开放接口"})
    _add_viewer_member(db_session, project, viewer_client.user)
    fake = Mock()
    fake.complete_json.return_value = {"fields": {"name": "方案建议", "design_markdown": "## 建议\n整合订单", "deliverables": "应用", "opportunity_ids": ["evil"], "status": "archived"}}
    app.extensions["fde_api_solution_organizer"] = fake
    url = f"/api/v1/projects/{project.id}/solutions/ai-organize"
    payload = {"name": "", "opportunity_ids": [opportunity.id], "reference": "只处理客户订单"}
    assert viewer_client.post(url, json=payload).status_code == 403
    response = leader_client.post(url, json=payload)
    assert response.status_code == 200, response.json
    assert response.json["data"]["saved"] is False
    assert set(response.json["data"]["fields"]) == {"name", "design_markdown", "deliverables"}
    context = json.loads(fake.complete_json.call_args.kwargs["messages"][1]["content"])
    assert context["reference"] == "只处理客户订单"
    assert len(context["opportunities"]) == 1
    assert "unselected" not in str(context)
    source = next(item for item in context["research_sources"] if item.get("form_id") == form.id)
    assert source["revision_id"] == revision.id
    assert "ERP 尚未开放接口" in str(source)
    assert db_session.scalar(select(func.count()).select_from(ProjectSolution)) == 0
    assert db_session.scalar(select(func.count()).select_from(ProjectDocument)) == 0


def test_sow_export_freezes_saved_solution_without_ai_and_reuses_revision(app, db_session, leader_client):
    project = _persist_project(db_session, leader_client.user, code="SOLUTION-SOW")
    opportunity = _opportunity(db_session, project)
    _persist_doc_template(db_session, leader_client.user)
    item = _create(leader_client, project, [opportunity]).json["data"]
    url = f"/api/v1/projects/{project.id}/solutions/{item['id']}"
    response = leader_client.post(url + "/export-sow", json={"version": 1})
    assert response.status_code == 200, response.json
    export = response.json["data"]
    assert export["document"]["display_name"] == f"SOW-{item['name']}"
    repeated = leader_client.post(url + "/export-sow", json={"version": 1})
    assert repeated.status_code == 200, repeated.json
    assert repeated.json["data"]["document_version_id"] == export["document_version_id"]
    assert repeated.json["data"]["reused"] is True
    legacy_version = db_session.get(ProjectDocumentVersion, export["document_version_id"])
    legacy_version.content_snapshot_json = {k: v for k, v in legacy_version.content_snapshot_json.items() if k != "sow_format_version"}
    db_session.commit()
    upgraded = leader_client.post(url + "/export-sow", json={"version": 1})
    assert upgraded.status_code == 200, upgraded.json
    assert upgraded.json["data"]["document_version_id"] != export["document_version_id"]
    assert upgraded.json["data"]["document_id"] == export["document_id"]
    assert leader_client.patch(url, json={"version": 1, "design_markdown": "后来改成二期方案"}).status_code == 200
    opportunity.name = "机会也改名了"
    db_session.commit()
    next_export = leader_client.post(url + "/export-sow", json={"version": 2})
    assert next_export.status_code == 200, next_export.json
    assert next_export.json["data"]["document_id"] == export["document_id"]
    assert next_export.json["data"]["document_version_id"] != export["document_version_id"]
    version = db_session.get(ProjectDocumentVersion, export["document_version_id"])
    document = db_session.get(ProjectDocument, export["document_id"])
    assert version.source_snapshot_json["solution"]["version"] == 1
    assert version.source_snapshot_json["solution"]["opportunities"][0]["name"] == "机会-one"
    assert db_session.scalar(select(func.count()).select_from(ProjectSolutionExport)) == 2
    template_path = RESOURCE_TEMPLATES_DIR / "sow.docx"
    template_bytes = template_path.read_bytes()
    with patch("fde_api.documents.generator.generate_document_fields", side_effect=AssertionError("Saved solution must not call AI")):
        context = _build_context(db_session, version, document, template_bytes)
    assert context["design_markdown"] == item["design_markdown"]
    assert "后来" not in str(context["solution"])
    rendered = Document(render_docx(io.BytesIO(template_bytes), context, {}))
    text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)
    text += "\n".join(cell.text for table in rendered.tables for row in table.rows for cell in row.cells)
    assert "【待填写：" not in text
    assert "业务流程" in text and "人工确认" in text
    assert "**人工确认**" not in text and "## 业务流程" not in text
    assert "后来改成" not in text
    assert "客户 ERP" in text
    document_url = f"/api/v1/projects/{project.id}/documents/{export['document_id']}"
    assert leader_client.post(document_url + "/generate", json={"version": document.version}).json["error"]["code"] == "solution_source_readonly"
    assert leader_client.patch(document_url + "/draft", json={"version": document.draft.version, "field_overrides": {"design_markdown": "绕过方案"}}).status_code == 409
    forged = leader_client.post(document_url + "/revise", json={"version": document.version, "draft_changes": {"solution": {"id": "forged"}}})
    assert forged.status_code == 400


def test_archived_opportunity_retained_but_cannot_be_newly_linked(db_session, leader_client):
    project = _persist_project(db_session, leader_client.user, code="SOLUTION-ARCH-OPP")
    opportunity = _opportunity(db_session, project)
    item = _create(leader_client, project, [opportunity]).json["data"]
    opportunity.status = "archived"
    db_session.commit()
    assert _create(leader_client, project, [opportunity]).status_code == 400
    response = leader_client.patch(f"/api/v1/projects/{project.id}/solutions/{item['id']}", json={"version": 1, "name": "保留历史关联", "opportunity_ids": [opportunity.id]})
    assert response.status_code == 200, response.json


def test_legacy_copy_is_explicit_idempotent_and_keeps_original_answers(app, db_session, leader_client):
    project = _persist_project(db_session, leader_client.user, code="SOLUTION-LEGACY")
    opportunity = _opportunity(db_session, project)
    form = ProjectResearchForm(project_id=project.id, subject_id=opportunity.id, form_key="legacy", name="旧交付说明")
    db_session.add(form)
    db_session.flush()
    revision = ProjectResearchFormRevision(form_id=form.id, revision_number=1, definition_snapshot={"sections": []})
    db_session.add(revision)
    db_session.flush()
    answer = ProjectResearchAnswer(revision_id=revision.id, field_key="delivery_scope", value_json="原交付内容")
    db_session.add(answer)
    form.current_revision_id = revision.id
    db_session.commit()
    preview = migrate_legacy_opportunity_delivery(actor=leader_client.user, project_id=project.id)
    assert preview["dry_run"] is True and preview["count"] == 1
    assert db_session.scalar(select(func.count()).select_from(ProjectSolution)) == 0
    migrated = migrate_legacy_opportunity_delivery(actor=leader_client.user, project_id=project.id, dry_run=False)
    assert migrated["count"] == 1
    assert migrate_legacy_opportunity_delivery(actor=leader_client.user, project_id=project.id, dry_run=False)["count"] == 0
    db_session.rollback()  # New service transactions must be visible under MySQL REPEATABLE READ.
    solution = db_session.get(ProjectSolution, migrated["items"][0]["solution_id"])
    assert solution.design_markdown == "原交付内容"
    assert solution.source_json["forms"][0]["revision_id"] == revision.id
    db_session.refresh(answer)
    assert answer.value_json == "原交付内容"


def test_failed_export_retry_is_same_version_and_archived_sow_is_not_revived(db_session, leader_client):
    project = _persist_project(db_session, leader_client.user, code="SOLUTION-RETRY")
    opportunity = _opportunity(db_session, project)
    _persist_doc_template(db_session, leader_client.user)
    item = _create(leader_client, project, [opportunity]).json["data"]
    url = f"/api/v1/projects/{project.id}/solutions/{item['id']}"
    first = leader_client.post(url + "/export-sow", json={"version": 1}).json["data"]
    job = db_session.scalar(select(DocumentGenerationJob).where(DocumentGenerationJob.document_version_id == first["document_version_id"]))
    job.status = "failed"
    job.last_error = "临时存储不可用"
    db_session.commit()
    retry = leader_client.post(url + "/export-sow", json={"version": 1})
    assert retry.status_code == 200, retry.json
    assert retry.json["data"]["document_version_id"] == first["document_version_id"]
    assert retry.json["data"]["generation"]["generation_status"] == "queued"
    versions = leader_client.get(f"/api/v1/projects/{project.id}/documents/{first['document_id']}/versions")
    assert versions.json["data"]["items"][0]["generation_status"] == "queued"
    doc = db_session.get(ProjectDocument, first["document_id"])
    doc.status = "archived"
    db_session.commit()
    assert leader_client.post(url + "/export-sow", json={"version": 1}).status_code == 409
    assert leader_client.patch(url, json={"version": 1, "name": "再次评审的方案"}).status_code == 200
    second = leader_client.post(url + "/export-sow", json={"version": 2})
    assert second.status_code == 200, second.json
    assert second.json["data"]["document_id"] != first["document_id"]
    db_session.refresh(doc)
    assert doc.status == "archived"


def test_export_refuses_old_incompatible_template_without_writing_document(db_session, leader_client):
    project = _persist_project(db_session, leader_client.user, code="SOLUTION-TEMPLATE")
    opportunity = _opportunity(db_session, project)
    _persist_doc_template(db_session, leader_client.user)
    item = _create(leader_client, project, [opportunity]).json["data"]
    old = Document()
    old.add_paragraph("旧模板只说明交付范围")
    content = io.BytesIO()
    old.save(content)
    with patch("fde_api.documents.generator._load_template_docx", return_value=content.getvalue()):
        response = leader_client.post(f"/api/v1/projects/{project.id}/solutions/{item['id']}/export-sow", json={"version": 1})
    assert response.status_code == 409, response.json
    assert response.json["error"]["code"] == "solution_template_upgrade_required"
    assert db_session.scalar(select(func.count()).select_from(ProjectDocument)) == 0


def test_export_refuses_archived_export_version_even_when_document_has_active_revision(db_session, leader_client):
    project = _persist_project(db_session, leader_client.user, code="SOLUTION-ARCH-VERSION")
    opportunity = _opportunity(db_session, project)
    _persist_doc_template(db_session, leader_client.user)
    item = _create(leader_client, project, [opportunity]).json["data"]
    solution_url = f"/api/v1/projects/{project.id}/solutions/{item['id']}"
    exported = leader_client.post(solution_url + "/export-sow", json={"version": 1}).json["data"]
    document_url = f"/api/v1/projects/{project.id}/documents/{exported['document_id']}"
    revised = leader_client.post(document_url + "/revise", json={"version": exported["document"]["version"], "draft_changes": {"field_overrides": {"deliverables": "人工审阅修订"}}})
    assert revised.status_code == 200, revised.json
    archived = leader_client.post(f"/api/v1/projects/{project.id}/documents/{exported['document_version_id']}/archive", json={"version": revised.json["data"]["document_version"]})
    assert archived.status_code == 200, archived.json
    document = leader_client.get(document_url).json["data"]
    assert document["status"] == "draft"
    assert document["current_version_id"] == revised.json["data"]["id"]
    repeated = leader_client.post(solution_url + "/export-sow", json={"version": 1})
    assert repeated.status_code == 409, repeated.json
    assert repeated.json["error"]["code"] == "solution_export_archived"
    assert db_session.get(ProjectDocumentVersion, exported["document_version_id"]).status == "archived"


def test_solution_design_stores_long_chinese_text(db_session, leader_client):
    project = _persist_project(db_session, leader_client.user, code="SOLUTION-LONG")
    opportunity = _opportunity(db_session, project)
    response = _create(leader_client, project, [opportunity], design_markdown="方案" * 15000)
    assert response.status_code == 201, response.json
    assert len(response.json["data"]["design_markdown"]) == 30000


def test_legacy_custom_delivery_field_is_copied_even_without_standard_fields(db_session, leader_client):
    project = _persist_project(db_session, leader_client.user, code="SOLUTION-LEGACY-CUSTOM")
    opportunity = _opportunity(db_session, project)
    form, revision = _form(db_session, project, opportunity, {"custom_delivery": "必须支持断网后继续处理", "current_state": "不应误归交付字段"})
    revision.definition_snapshot = {"sections": [{"section_key": "delivery", "name": "交付说明", "fields": [{"field_key": "custom_delivery", "name": "离线交付约束"}]}]}
    db_session.commit()
    copied = migrate_legacy_opportunity_delivery(actor=leader_client.user, project_id=project.id, dry_run=False)
    assert copied["count"] == 1
    solution = db_session.get(ProjectSolution, copied["items"][0]["solution_id"])
    assert "历史补充交付内容" in solution.design_markdown
    assert "离线交付约束" in solution.design_markdown
    assert "必须支持断网后继续处理" in solution.design_markdown
    assert "不应误归" not in solution.design_markdown
    assert solution.source_json["forms"][0]["answers"] == {"custom_delivery": "必须支持断网后继续处理"}
