from __future__ import annotations

import hashlib
import io
from dataclasses import dataclass
from datetime import date
from urllib.parse import urlparse
from uuid import uuid4

import pytest
from docx import Document
from sqlalchemy import select

from fde_api.auth.models import User
from fde_api.auth.passwords import hash_password
from fde_api.auth.tokens import issue_access_token
from fde_api.documents.models import (
    DocumentGenerationJob,
    DocumentTemplate,
    DocumentTemplateVersion,
    ProjectDocument,
    ProjectDocumentVersion,
)
from fde_api.extensions import object_storage
from fde_api.files.models import ProjectFile, ProjectFileVersion
from fde_api.files.names import extension_of, safe_filename
from fde_api.jobs.models import OutboxEvent
from fde_api.workbench.models import (
    IndustryTemplate,
    IndustryTemplateVersion,
    Project,
    ProjectMember,
)

INITIAL_PASSWORD = "InitialPass!234"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@dataclass(frozen=True)
class AuthorizedClient:
    client: object
    user: User
    headers: dict[str, str]

    def get(self, *args, **kwargs):
        return self.client.get(*args, headers=self.headers, **kwargs)

    def post(self, *args, **kwargs):
        return self.client.post(*args, headers=self.headers, **kwargs)


@pytest.fixture
def authorized_client_factory(client, db_session, settings):
    def create(role: str) -> AuthorizedClient:
        user = User(
            username=f"{role}.{uuid4().hex}",
            display_name=role,
            role=role,
            password_hash=hash_password(INITIAL_PASSWORD),
            must_change_password=False,
            is_active=True,
        )
        db_session.add(user)
        db_session.commit()
        return AuthorizedClient(
            client,
            user,
            {"Authorization": f"Bearer {issue_access_token(user, settings)}"},
        )

    return create


@pytest.fixture
def leader_client(authorized_client_factory):
    return authorized_client_factory("project_lead")


@pytest.fixture
def engineer_client(authorized_client_factory):
    return authorized_client_factory("fde_engineer")


def _docx_bytes(*, text: str = "hello revision") -> bytes:
    doc = Document()
    doc.add_paragraph(text)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def _seed_document(
    db_session,
    leader: User,
    engineer: User,
    *,
    old_docx_bytes: bytes | None = None,
    old_docx_status: str = "uploading",
):
    """Build a project with a draft document holding a generated v1 version."""
    industry_template = IndustryTemplate(name="文档模板", industry_name="制造")
    industry_version = IndustryTemplateVersion(
        template=industry_template,
        name=industry_template.name,
        industry_name=industry_template.industry_name,
        description="",
        version_number=1,
        status="published",
        published_by=leader,
    )
    db_session.add_all([industry_template, industry_version])
    db_session.flush()

    project = Project(
        project_code=f"VER-{uuid4().hex[:8]}",
        name="版本项目",
        enterprise_name="星河制造",
        leader=leader,
        source_template_version=industry_version,
        template_snapshot={},
        planned_start_date=date(2026, 8, 22),
    )
    db_session.add(project)
    db_session.flush()
    db_session.add(ProjectMember(project_id=project.id, user_id=engineer.id, role="member"))
    db_session.flush()

    doc_template = DocumentTemplate(
        document_type="sow",
        name="sow 模板",
        description="",
        industry_name="制造",
    )
    db_session.add(doc_template)
    db_session.flush()
    template_version = DocumentTemplateVersion(
        template=doc_template,
        version_number=1,
        status="published",
        mapping_json={
            "document_type": "sow",
            "name": "sow 模板",
            "modules": ["pre_diagnosis"],
            "field_map": {},
            "research_keys": [],
        },
        sections_json={"sections": []},
        table_loops_json={},
        required_data_json={},
        source_sha256="",
        created_by=leader,
    )
    db_session.add(template_version)
    db_session.flush()

    document = ProjectDocument(
        project=project,
        document_type="sow",
        business_code="SOW-1",
        source_template_version_id=template_version.id,
        status="draft",
        owner=leader,
    )
    db_session.add(document)
    db_session.flush()

    docx_filename = "SOW-1-v1.docx"
    docx_safe = safe_filename(docx_filename)
    docx_key = (
        f"projects/{project.id}/documents/{document.id}/versions/1/{docx_safe}"
    )
    docx_file = ProjectFile(
        project=project,
        category="document",
        display_name=docx_filename,
        description="",
        created_by=leader,
    )
    db_session.add(docx_file)
    db_session.flush()
    docx_file_version = ProjectFileVersion(
        project_file=docx_file,
        version_number=1,
        source="system_generated",
        original_filename=docx_filename,
        safe_filename=docx_safe,
        extension=extension_of(docx_filename),
        mime_type=DOCX_MIME,
        bucket="local",
        storage_key=docx_key,
        size_bytes=0,
        etag="",
        sha256="",
        scan_status="not_required",
        preview_status="none",
        uploaded_by=leader,
        status=old_docx_status,
    )
    db_session.add(docx_file_version)
    db_session.flush()
    docx_file.current_version_id = docx_file_version.id

    if old_docx_bytes is not None:
        stored = object_storage.current.put_stream(
            docx_key, io.BytesIO(old_docx_bytes), DOCX_MIME, {}
        )
        docx_file_version.size_bytes = stored.size
        docx_file_version.etag = stored.etag
        docx_file_version.sha256 = hashlib.sha256(old_docx_bytes).hexdigest()

    snapshot = {
        "project": {
            "id": project.id,
            "name": project.name,
            "enterprise_name": project.enterprise_name,
        },
        "enterprise": {"name": project.enterprise_name},
        "research": {},
        "subjects": {},
        "tasks": [],
        "members": [],
        "captured_at": "2026-08-22T00:00:00+00:00",
    }
    version = ProjectDocumentVersion(
        document=document,
        version_number=1,
        source="generated",
        status="draft",
        parent_version_id=None,
        source_template_version_id=template_version.id,
        source_snapshot_json=snapshot,
        content_snapshot_json={
            "field_overrides": {},
            "rich_text": {},
            "list_selections": {},
        },
        docx_file_version_id=docx_file_version.id,
        sha256="",
        generated_by=leader,
    )
    db_session.add(version)
    db_session.flush()

    document.current_version_id = version.id
    document.version = 2
    db_session.commit()
    return {
        "project": project,
        "document": document,
        "template_version": template_version,
        "v1": version,
        "v1_docx_file": docx_file_version,
        "v1_docx_key": docx_key,
    }


def _manual_upload(client, project_id, document_id, version, docx_bytes, *, note=""):
    return client.post(
        f"/api/v1/projects/{project_id}/documents/{document_id}/revisions/manual",
        data={
            "version": str(version),
            "note": note,
            "file": (io.BytesIO(docx_bytes), "revision.docx"),
        },
    )


def assert_error(response, status: int, code: str) -> None:
    assert response.status_code == status, response.get_data(as_text=True)
    assert response.json["data"] is None
    assert response.json["error"]["code"] == code


def test_manual_upload_creates_new_current_version_without_overwrite(
    db_session, leader_client, engineer_client
):
    leader = leader_client.user
    engineer = engineer_client.user
    old_bytes = _docx_bytes(text="original v1")
    seeded = _seed_document(
        db_session, leader, engineer, old_docx_bytes=old_bytes
    )
    document = seeded["document"]
    old_key = seeded["v1_docx_key"]
    uploaded_bytes = _docx_bytes(text="manual v2")

    response = _manual_upload(
        engineer_client,
        seeded["project"].id,
        document.id,
        document.version,
        uploaded_bytes,
    )
    assert response.status_code == 201, response.get_data(as_text=True)
    payload = response.json["data"]

    assert payload["source"] == "manual_upload"
    assert payload["version_number"] == 2
    assert payload["status"] == "draft"
    assert payload["document_version"] == 3

    # The original OSS object is immutable: bytes and key are unchanged.
    with object_storage.current.open_stream(old_key) as stream:
        assert stream.read() == old_bytes

    # The new manifest points at a brand new file version, never the old one.
    new_version = db_session.get(ProjectDocumentVersion, payload["id"])
    assert new_version is not None
    assert new_version.docx_file_version_id != seeded["v1_docx_file"].id
    assert new_version.sha256 == hashlib.sha256(uploaded_bytes).hexdigest()

    db_session.refresh(document)
    assert document.current_version_id == payload["id"]
    assert document.version == 3


def test_engineer_cannot_confirm_document(db_session, leader_client, engineer_client):
    leader = leader_client.user
    engineer = engineer_client.user
    seeded = _seed_document(db_session, leader, engineer)

    response = engineer_client.post(
        f"/api/v1/projects/{seeded['project'].id}/documents/{seeded['v1'].id}/confirm",
        json={"version": seeded["document"].version},
    )
    assert_error(response, 403, "forbidden")


def test_project_lead_confirms_and_archives_prior(
    db_session, leader_client, engineer_client
):
    leader = leader_client.user
    engineer = engineer_client.user
    seeded = _seed_document(db_session, leader, engineer)
    project_id = seeded["project"].id
    document = seeded["document"]
    v1 = seeded["v1"]

    documents_base = f"/api/v1/projects/{project_id}/documents"

    # 1. Confirm the generated v1 first.
    first_confirm = leader_client.post(
        f"{documents_base}/{v1.id}/confirm",
        json={"version": document.version},
    )
    assert first_confirm.status_code == 200, first_confirm.get_data(as_text=True)
    assert first_confirm.json["data"]["status"] == "confirmed"
    assert first_confirm.json["data"]["confirmed_by"]["id"] == leader.id
    document_version_after_first = first_confirm.json["data"]["document_version"]

    # 2. Add a second version via a manual upload.
    upload = _manual_upload(
        engineer_client,
        project_id,
        document.id,
        document_version_after_first,
        _docx_bytes(text="v2"),
    )
    assert upload.status_code == 201, upload.get_data(as_text=True)
    v2 = upload.json["data"]
    document_version_after_upload = v2["document_version"]

    # 3. Confirm v2 -> v2 confirmed, previously confirmed v1 archived.
    second_confirm = leader_client.post(
        f"{documents_base}/{v2['id']}/confirm",
        json={"version": document_version_after_upload},
    )
    assert second_confirm.status_code == 200, second_confirm.get_data(as_text=True)
    confirmed_data = second_confirm.json["data"]
    assert confirmed_data["status"] == "confirmed"
    assert confirmed_data["version_number"] == 2
    assert confirmed_data["document_version"] == document_version_after_upload + 1

    db_session.refresh(document)
    assert document.status == "confirmed"
    assert document.current_version_id == v2["id"]

    refetched_v1 = db_session.get(ProjectDocumentVersion, v1.id)
    db_session.refresh(refetched_v1)
    assert refetched_v1.status == "archived"
    assert refetched_v1.confirmed_by_user_id == leader.id


def test_project_lead_can_archive_current_confirmed_version(
    db_session, leader_client, engineer_client
):
    leader = leader_client.user
    engineer = engineer_client.user
    seeded = _seed_document(db_session, leader, engineer)
    document = seeded["document"]
    version = seeded["v1"]
    base = f"/api/v1/projects/{seeded['project'].id}/documents"

    confirmed = leader_client.post(
        f"{base}/{version.id}/confirm",
        json={"version": document.version},
    )
    assert confirmed.status_code == 200, confirmed.get_data(as_text=True)

    archived = leader_client.post(
        f"{base}/{version.id}/archive",
        json={"version": confirmed.json["data"]["document_version"]},
    )
    assert archived.status_code == 200, archived.get_data(as_text=True)
    assert archived.json["data"]["status"] == "archived"

    db_session.refresh(document)
    assert document.status == "archived"
    assert document.current_version_id == version.id

    restored = leader_client.post(
        f"{base}/{version.id}/restore",
        json={"version": archived.json["data"]["document_version"]},
    )
    assert restored.status_code == 200, restored.get_data(as_text=True)
    assert restored.json["data"]["status"] == "confirmed"
    assert restored.json["data"]["document_version"] == archived.json["data"]["document_version"] + 1
    db_session.rollback()
    document = db_session.get(ProjectDocument, document.id)
    assert document is not None
    db_session.refresh(document)
    assert document.status == "confirmed"
    assert document.current_version_id == version.id


def test_document_card_archives_all_versions_and_restores_only_latest(
    db_session, leader_client, engineer_client
):
    seeded = _seed_document(db_session, leader_client.user, engineer_client.user)
    project_id = seeded["project"].id
    document = seeded["document"]
    v1 = seeded["v1"]
    base = f"/api/v1/projects/{project_id}/documents"
    confirmed = leader_client.post(
        f"{base}/{v1.id}/confirm", json={"version": document.version}
    )
    upload = _manual_upload(
        engineer_client,
        project_id,
        document.id,
        confirmed.json["data"]["document_version"],
        _docx_bytes(text="v2"),
    )

    archived = leader_client.post(
        f"{base}/items/{document.id}/archive",
        json={"version": upload.json["data"]["document_version"]},
    )

    assert archived.status_code == 200, archived.get_data(as_text=True)
    assert archived.json["data"]["archived_count"] == 2
    db_session.expire_all()
    versions = (
        db_session.query(ProjectDocumentVersion)
        .filter_by(document_id=document.id)
        .order_by(ProjectDocumentVersion.version_number)
        .all()
    )
    assert [item.status for item in versions] == ["archived", "archived"]

    restored = leader_client.post(
        f"{base}/items/{document.id}/restore",
        json={"version": archived.json["data"]["document_version"]},
    )

    assert restored.status_code == 200, restored.get_data(as_text=True)
    assert restored.json["data"]["id"] == upload.json["data"]["id"]
    db_session.rollback()
    versions = (
        db_session.query(ProjectDocumentVersion)
        .filter_by(document_id=document.id)
        .order_by(ProjectDocumentVersion.version_number)
        .all()
    )
    assert [item.status for item in versions] == ["archived", "confirmed"]

def test_online_revise_creates_new_draft(db_session, leader_client, engineer_client):
    leader = leader_client.user
    engineer = engineer_client.user
    seeded = _seed_document(db_session, leader, engineer)
    document = seeded["document"]

    response = engineer_client.post(
        f"/api/v1/projects/{seeded['project'].id}/documents/{document.id}/revise",
        json={
            "version": document.version,
            "draft_changes": {
                "field_overrides": {"title": "修订版"},
                "rich_text": {
                    "type": "doc",
                    "content": [{"type": "paragraph", "content": "正文"}],
                },
            },
        },
    )
    assert response.status_code == 200, response.get_data(as_text=True)
    payload = response.json["data"]
    assert payload["source"] == "online_revised"
    assert payload["status"] == "draft"
    assert payload["version_number"] == 2
    assert payload["parent_version_id"] == seeded["v1"].id
    assert payload["content_snapshot_json"]["field_overrides"] == {"title": "修订版"}

    job = db_session.scalar(
        select(DocumentGenerationJob).where(
            DocumentGenerationJob.document_version_id == payload["id"]
        )
    )
    assert job is not None
    assert job.status == "queued"

    outbox = list(
        db_session.scalars(
            select(OutboxEvent).where(OutboxEvent.topic == "document.generate")
        )
    )
    assert len(outbox) == 1
    assert outbox[0].aggregate_id == payload["id"]


def test_download_url_requires_available(db_session, leader_client, engineer_client):
    leader = leader_client.user
    engineer = engineer_client.user
    seeded = _seed_document(db_session, leader, engineer)
    project_id = seeded["project"].id
    document = seeded["document"]
    v1 = seeded["v1"]
    documents_base = f"/api/v1/projects/{project_id}/documents"

    # v1's DOCX is a placeholder (uploading) and therefore not available.
    unavailable = engineer_client.get(
        f"{documents_base}/{v1.id}/download-url"
    )
    assert_error(unavailable, 409, "file_not_available")

    # A manual upload promotes an available, downloadable version.
    upload = _manual_upload(
        engineer_client,
        project_id,
        document.id,
        document.version,
        _docx_bytes(text="available"),
    )
    assert upload.status_code == 201, upload.get_data(as_text=True)
    v2 = upload.json["data"]

    available = engineer_client.get(
        f"{documents_base}/{v2['id']}/download-url"
    )
    assert available.status_code == 200, available.get_data(as_text=True)
    data = available.json["data"]
    assert isinstance(data["url"], str)
    assert urlparse(data["url"]).scheme == "http"
    assert data["expires_seconds"] == 300


def test_version_history_is_listed(db_session, leader_client, engineer_client):
    leader = leader_client.user
    engineer = engineer_client.user
    seeded = _seed_document(db_session, leader, engineer)
    project_id = seeded["project"].id
    document = seeded["document"]
    documents_base = f"/api/v1/projects/{project_id}/documents"

    upload = _manual_upload(
        engineer_client,
        project_id,
        document.id,
        document.version,
        _docx_bytes(text="v2"),
    )
    assert upload.status_code == 201, upload.get_data(as_text=True)

    response = engineer_client.get(f"{documents_base}/{document.id}/versions")
    assert response.status_code == 200, response.get_data(as_text=True)
    items = response.json["data"]["items"]
    assert [item["version_number"] for item in items] == [2, 1]
    assert items[0]["source"] == "manual_upload"
    assert items[0]["status"] == "draft"
    assert items[0]["docx_file_version_id"]
    assert items[0]["download_url"]
    assert items[1]["source"] == "generated"
    assert items[1]["download_url"] is None
    assert items[1]["preview_status"] == "none"
