import io

import pytest
from docx import Document
from fde_api.documents.template_service import (
    _test_render_template, DocumentTemplateServiceError,
)


def template(text):
    doc = Document()
    doc.add_table(rows=1, cols=1).cell(0, 0).text = text
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def test_valid_sow_tokens_are_rendered_before_residual_check():
    text = _test_render_template(template(
        '{{ project.name }} {{ enterprise.name }} '
        '{{ document.business_code }} {{ document.version_number }}'))
    assert '测试项目' in text and 'SOW-测试方案' in text
    assert '{{' not in text


@pytest.mark.parametrize('text', [
    '{{ project.name', '{{ unknown.name }}', '{{ project.__class__ }}',
])
def test_invalid_or_unsafe_tokens_still_fail(text):
    with pytest.raises(DocumentTemplateServiceError):
        _test_render_template(template(text))
