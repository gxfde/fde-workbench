from io import BytesIO

from docx import Document

from fde_api.guidance.docx_extract import extract_docx


def test_extracts_paragraphs_and_tables_with_stable_refs():
    document = Document()
    document.add_paragraph("客户目标")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "部门"
    table.cell(0, 1).text = "研发部"
    stream = BytesIO()
    document.save(stream)
    stream.seek(0)

    result = extract_docx(stream)

    assert result.blocks[0].source_ref == "paragraph-1"
    assert result.blocks[0].text == "客户目标"
    assert any(block.source_ref == "table-1-row-1-cell-1" for block in result.blocks)
