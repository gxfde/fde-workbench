import pytest
from docx import Document
from fde_api.documents.markdown_docx import insert_markdown
from fde_api.documents.mermaid_render import prepare


def test_native_table_and_inline_formatting():
    doc = Document()
    anchor = doc.add_paragraph('placeholder')
    insert_markdown(doc, anchor, '## 方案\n\n- **用途**：分析\n\n  | 模块 | 功能 |\n  | --- | --- |\n  | BOM | 比对 |\n')
    assert len(doc.tables) == 1
    assert doc.tables[0].cell(1, 1).text == '比对'
    assert any(r.bold and r.text == '用途' for p in doc.paragraphs for r in p.runs)
    assert 'placeholder' not in '\n'.join(p.text for p in doc.paragraphs)


def test_mermaid_configuration_is_never_executed():
    src = '%%{init: {"theme": "base"}}%%\ngraph TD\nA["上传<br/>资料"]-->B\nclassDef neutral fill:#fff,stroke:#000;'
    result = prepare(src)
    assert 'init' not in result and 'classDef' not in result
    assert '上传<br/>资料' in result
    with pytest.raises(ValueError): prepare('graph TD\nA-->B;click A "https://example.com"')
