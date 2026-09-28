from io import BytesIO

from docx import Document

from fde_api.documents.template_parser import (
    ALLOWED_BLOCK_TYPES,
    ALLOWED_MARKS,
    ALLOWED_ROOTS,
    scan_docx_placeholders,
)


def _build_docx(
    *,
    paragraphs=None,
    table_cells=None,
    header=None,
    footer=None,
):
    paragraphs = paragraphs or []
    table_cells = table_cells or []
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    for cell_text in table_cells:
        table = document.add_table(rows=1, cols=1)
        table.cell(0, 0).text = cell_text
    section = document.sections[0]
    if header is not None:
        header_paragraph = section.header.paragraphs[0]
        header_paragraph.text = header
    if footer is not None:
        footer_paragraph = section.footer.paragraphs[0]
        footer_paragraph.text = footer
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer


def test_parser_rejects_function_calls_and_attribute_escape():
    buffer = _build_docx(paragraphs=["{{ cycler.__init__.__globals__.os.system('id') }}"])
    scan = scan_docx_placeholders(buffer)
    assert scan.placeholders == ["{{ cycler.__init__.__globals__.os.system('id') }}"]
    assert scan.issues, "expected at least one validation issue"
    assert scan.issues[0].code == "unsafe_template_expression"


def test_parser_accepts_allowed_root_scalar():
    buffer = _build_docx(paragraphs=["{{ project.name }}"])
    scan = scan_docx_placeholders(buffer)
    unsafe_or_unknown = [
        issue
        for issue in scan.issues
        if issue.code in {"unsafe_template_expression", "unknown_template_root"}
    ]
    assert unsafe_or_unknown == [], scan.issues


def test_parser_rejects_unknown_root():
    buffer = _build_docx(paragraphs=["{{ evil.foo }}"])
    scan = scan_docx_placeholders(buffer)
    assert scan.issues
    assert any(issue.code == "unknown_template_root" for issue in scan.issues)


def test_parser_scans_headers_and_tables():
    buffer = _build_docx(
        paragraphs=["body content"],
        table_cells=["{{ project.name }}"],
        header="{{ project.enterprise }}",
    )
    scan = scan_docx_placeholders(buffer)
    assert "{{ project.name }}" in scan.tables
    assert "{{ project.name }}" in scan.placeholders
    assert "{{ project.enterprise }}" in scan.placeholders


def test_parser_records_balanced_loop_context():
    buffer = _build_docx(
        paragraphs=[
            "{% for role in subjects.roles %}{{ role.name }}{% endfor %}",
            "{% for role in subjects.roles %}{{ role.name }} / {{ loop.index }}{% endfor %}",
        ]
    )
    scan = scan_docx_placeholders(buffer)
    assert scan.issues == [], scan.issues
    assert "{{ loop.index }}" in scan.placeholders


def test_module_constants_are_frozen():
    assert ALLOWED_ROOTS == {
        "project",
        "enterprise",
        "research",
        "subjects",
        "tasks",
        "members",
        "document",
    }
    assert "paragraph" in ALLOWED_BLOCK_TYPES
    assert "bold" in ALLOWED_MARKS
